# -*- coding: utf-8 -*-
"""
Agua Bora - DECISION SUPPORT SYSTEM (app web en Streamlit, intento 2)

Prototipo TPM-IIoT de soporte a decisiones para la deteccion, diagnostico,
priorizacion y verificacion de intervenciones de mantenimiento en una linea de
envasado de agua (bidones retornables de 20 L, MYPE de Lima).

Ejecutar:   streamlit run app_web_intento_2.py
Lee la base "Base_OEE_MYPE_Peru_bidones_6_meses_2026.xlsx" (y los casos de prueba
Base_01 ... Base_06) desde esta misma carpeta.

Flujo de la app (menu lateral, un paso por pagina):
  Inicio -> 1. KPIs de entrada -> 2. Datos (empresa o sinteticos) -> 3. Parametros
  (umbrales P10/P90, rangos C1, escalas S/O/D, matriz de criticidad, AMEF, supuestos)
  -> 4. Tablero -> 5. Decision (parte post: alerta, causa AMEF, prioridad C3, decision)
  -> 6. Ejecucion y resultante (verifica con los datos del periodo siguiente)
  -> 7. Validacion (Monte Carlo: sin DSS / DSS parcial / DSS completo) -> 8. Exportacion.

Todo dato que no viene de la base ni de la tesis se marca como "supuesto" o
"(ejemplo)". Los datos de la base son sinteticos, con fines academicos.

Todo dato que no viene de la base ni de la tesis se marca como "supuesto" o
"(ejemplo)". Los datos de la base son sinteticos, con fines academicos.
"""

from __future__ import annotations

import copy
import io
import math
import textwrap
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from plotly.subplots import make_subplots

# =============================================================================
# CONSTANTES GENERALES
# =============================================================================
NOMBRE_BASE = "Base_OEE_MYPE_Peru_bidones_6_meses_2026.xlsx"
RUTA_BASE = Path(__file__).resolve().parent / NOMBRE_BASE
VERSION = "v1.0"

INK = "#0d1117"
MUTED = "#6b7280"
LINEA = "#e5e7eb"
AZUL = "#2f80ed"
VERDE = "#12a150"
GRIS = "#8c96a8"
AMBAR = "#e0a100"
ROJO = "#d92d20"

SIN_PARADA = "Sin parada no planificada"
ESPERA = "Espera de proceso (línea detenida)"
CAUSAS_EQUIPO = [
    "Avería de llenadora-tapadora",
    "Avería de lavadora de bidones",
    "Avería de selladora-etiquetadora",
    "Ajuste o limpieza correctiva",
]
CAUSAS_INSUMO = ["Falta de tapas o precintos", "Falta de bidones vacíos (retorno de reparto)"]

EST_NOM = {1: "E1 · Lavado y sanitizado", 2: "E2 · Llenado y tapado", 3: "E3 · Sellado y etiquetado"}
EST_CORTO = {1: "E1", 2: "E2", 3: "E3"}
MESES = {1: "Ene", 2: "Feb", 3: "Mar", 4: "Abr", 5: "May", 6: "Jun",
         7: "Jul", 8: "Ago", 9: "Sep", 10: "Oct", 11: "Nov", 12: "Dic"}


def mes_lbl(clave):
    """Clave de mes (año*100 + mes) -> etiqueta corta, por ejemplo «Ene-26»."""
    clave = int(clave)
    return f"{MESES[clave % 100]}-{str(clave // 100)[2:]}"

ESCENARIOS_SIM = ["Operación sin DSS", "DSS parcial (solo C1)", "DSS completo (C1–C4)"]
COLOR_ESC = {0: GRIS, 1: AZUL, 2: VERDE}

COLS_REQ = ["Fecha", "Estacion_ID", "Tiempo_Planificado_h", "Parada_No_Planificada_h",
            "Tiempo_Operativo_h", "Causa_Parada_No_Planificada",
            "Capacidad_Nominal_bidones_h", "Produccion_Teorica_bidones",
            "Bidones_Procesados", "Bidones_Conformes", "Bidones_No_Conformes"]
COLS_NUM = [c for c in COLS_REQ if c not in ("Fecha", "Causa_Parada_No_Planificada")]

# Metas de la Tabla 2 (tipo, valor). OEE, D, R, Q, rechazo en fraccion; horas; min.
METAS = {
    "OEE": (">=", 0.65), "D": (">=", 0.92), "R": (">=", 0.72), "Q": (">=", 0.99),
    "rechazo": ("<=", 0.025), "Pa": ("<=", 180.0), "MTBF": (">=", 9.0),
    "MTTR": ("<=", 0.75), "t_resp": ("<=", 15.0),
}
META_TXT = {
    "OEE": "≥ 65 %", "D": "≥ 92 %", "R": "≥ 72 %", "Q": "≥ 99,0 %",
    "rechazo": "≤ 2,5 %", "Pa": "≤ 180 h", "MTBF": "≥ 9 h", "MTTR": "≤ 0,75 h",
    "t_resp": "≤ 15 min",
}
INDICADORES = [  # (clave, etiqueta, tipo de formato)
    ("OEE", "OEE de línea", "pct"), ("D", "Disponibilidad (D)", "pct"),
    ("R", "Rendimiento (R)", "pct"), ("Q", "Calidad (Q)", "pct"),
    ("Pa", "Horas de parada", "h"), ("ev", "Eventos de parada", "int"),
    ("MTBF", "MTBF", "h"), ("MTTR", "MTTR", "h"),
    ("rechazo", "Índice de rechazo", "pct"), ("conformes", "Bidones conformes", "int"),
    ("t_resp", "Tiempo de respuesta", "min"),
]
ETIQ = {k: e for k, e, _ in INDICADORES}
TIPO = {k: t for k, _, t in INDICADORES}

# Supuestos de simulacion (NO provienen de la base ni de la tesis)
SUPUESTOS_DEF = {
    "s_p_det": 0.60,
    "s_r_dur_par": 0.25, "s_r_dur_com": 0.40,
    "s_t_sin": 45.0, "s_t_par": 22.0, "s_t_com": 13.0, "s_disp_resp": 0.15,
    "s_eff_c4": 0.60, "s_tau_c4": 6.0,
    "s_propag": 0.70,
    "s_mej_r_par": 0.02, "s_mej_r_com": 0.06,
    "s_red_def_par": 0.05, "s_red_def_com": 0.25,
}


# =============================================================================
# FORMATO
# =============================================================================
def n(x, d=2):
    """Numero con coma decimal y espacio como separador de miles."""
    try:
        if x is None or not np.isfinite(float(x)):
            return "—"
    except (TypeError, ValueError):
        return "—"
    s = f"{float(x):,.{d}f}"
    return s.replace(",", " ").replace(".", ",")


def fmt(tipo, x):
    if x is None or (isinstance(x, float) and not np.isfinite(x)):
        return "—"
    if tipo == "pct":
        return f"{n(x * 100)} %"
    if tipo == "h":
        return f"{n(x)} h"
    if tipo == "int":
        return n(x, 0)
    if tipo == "min":
        return f"{n(x, 1)} min"
    return n(x)


def cumple(clave, x):
    if clave not in METAS or x is None or not np.isfinite(x):
        return None
    op, v = METAS[clave]
    return bool(x >= v) if op == ">=" else bool(x <= v)


# =============================================================================
# INDICADORES DE LA BASE (razon de totales, como en la tesis)
# =============================================================================
def _parada_ev(df):
    return ((df["Parada_No_Planificada_h"] > 0) & (df["Causa"] != SIN_PARADA)).astype(int)


def tabla_estaciones(df):
    """Por estacion: D, R, Q, OEE, paradas, MTBF, MTTR, no conformes (Tabla 1)."""
    d = df.assign(_ev=_parada_ev(df))
    g = d.groupby("Estacion_ID").agg(
        Tp=("Tiempo_Planificado_h", "sum"), Pa=("Parada_No_Planificada_h", "sum"),
        To=("Tiempo_Operativo_h", "sum"), Teor=("Teorica", "sum"),
        Proc=("Bidones_Procesados", "sum"), Conf=("Bidones_Conformes", "sum"),
        NC=("Bidones_No_Conformes", "sum"), ev=("_ev", "sum"))
    g["D"] = g.To / g.Tp
    g["R"] = g.Proc / g.Teor
    g["Q"] = g.Conf / g.Proc
    g["OEE"] = g.D * g.R * g.Q
    g["MTBF"] = g.To / g.ev.replace(0, np.nan)
    g["MTTR"] = g.Pa / g.ev.replace(0, np.nan)
    g["rech"] = g.NC / g.Proc
    return g


def kpis_base(df):
    g = tabla_estaciones(df)
    t = g[["Tp", "Pa", "To", "Teor", "Proc", "Conf", "NC", "ev"]].sum()
    return dict(
        OEE=float(g.OEE.mean()), D=t.To / t.Tp, R=t.Proc / t.Teor, Q=t.Conf / t.Proc,
        Pa=float(t.Pa), ev=int(t.ev), MTBF=t.To / t.ev if t.ev else np.nan,
        MTTR=t.Pa / t.ev if t.ev else np.nan,
        rechazo=t.NC / g.loc[1, "Proc"], conformes=float(g.loc[3, "Conf"]),
        entrada=float(g.loc[1, "Proc"]), nc=float(t.NC))


def tabla_pareto(df):
    d = df[(df["Causa"] != SIN_PARADA) & (df["Parada_No_Planificada_h"] > 0)]
    p = d.groupby("Causa").agg(Horas=("Parada_No_Planificada_h", "sum"),
                               Eventos=("Parada_No_Planificada_h", "size"))
    p = p.sort_values("Horas", ascending=False)
    p["Pct"] = p.Horas / p.Horas.sum()
    p["Acum"] = p.Pct.cumsum()
    return p.reset_index()


def referencias_p10_p90(df):
    """Umbrales de la parte post (Tabla 7) calculados con los valores diarios."""
    d = df.copy()
    d["oee_d"] = d["Bidones_Conformes"] / (d["Capacidad_Nominal_bidones_h"] * d["Tiempo_Planificado_h"])
    d.loc[d["Tiempo_Operativo_h"] <= 0, "oee_d"] = np.nan      # días de paro total: OEE no calculable
    d["rech_d"] = d["Bidones_No_Conformes"] / d["Bidones_Procesados"].replace(0, np.nan)
    ref = {}
    for e, g in d.groupby("Estacion_ID"):
        ref[e] = dict(
            oee_p10=float(g.oee_d.quantile(0.10)), oee_med=float(g.oee_d.median()),
            rech_p90=float(g.rech_d.quantile(0.90)), rech_med=float(g.rech_d.mean()))
    return ref


# =============================================================================
# MOTOR DE SIMULACION (Monte Carlo con bootstrap de dias por mes)
# =============================================================================
def preparar_sim(df):
    fechas = pd.DatetimeIndex(sorted(df["Fecha"].unique()))

    def piv(col):
        return (df.pivot(index="Fecha", columns="Estacion_ID", values=col)
                .reindex(index=fechas, columns=[1, 2, 3]).to_numpy())
    A = dict(
        Tp=piv("Tiempo_Planificado_h").astype(float), Pa=piv("Parada_No_Planificada_h").astype(float),
        To=piv("Tiempo_Operativo_h").astype(float), cap=piv("Capacidad_Nominal_bidones_h").astype(float),
        Proc=piv("Bidones_Procesados").astype(float), NC=piv("Bidones_No_Conformes").astype(float),
        causa=piv("Causa").astype(object), fechas=fechas,
        mes=(fechas.year * 100 + fechas.month).to_numpy())
    to_s = np.nansum(A["To"], axis=0)
    pr_s = np.nansum(A["Proc"], axis=0)
    nc_s = np.nansum(A["NC"], axis=0)
    A["tasa_est"] = np.where(to_s > 0, pr_s / np.where(to_s > 0, to_s, 1), 0.0)
    A["q_est"] = np.where(pr_s > 0, nc_s / np.where(pr_s > 0, pr_s, 1), 0.0)
    return A


def muestrear_dias(A, rng):
    """Bootstrap de los dias de la base dentro de cada mes (mismo numero de dias)."""
    partes = []
    for m in sorted(set(A["mes"].tolist())):
        ids = np.flatnonzero(A["mes"] == m)
        partes.append(np.sort(rng.choice(ids, size=len(ids), replace=True)))
    return np.concatenate(partes)


def simular_iteracion(A, P, scen, idx, rng, con_detalle=False):
    """Una iteracion de un escenario sobre los dias muestreados `idx`.
    scen: 0 sin DSS, 1 DSS parcial (solo C1), 2 DSS completo (C1-C4).
    Solo las causas de equipo y la espera de proceso cambian con el DSS."""
    Tp, Pa, To, cap = A["Tp"][idx], A["Pa"][idx], A["To"][idx], A["cap"][idx]
    Proc, NC, cau, mes = A["Proc"][idx], A["NC"][idx], A["causa"][idx], A["mes"][idx]
    nd = len(idx)
    es_eq = np.isin(cau, CAUSAS_EQUIPO) & (Pa > 0)
    es_esp = (cau == ESPERA) & (Pa > 0)
    Pa_n = Pa.copy()
    mej_r = red_def = 0.0
    if scen > 0:
        r_dur = P["s_r_dur_par"] if scen == 1 else P["s_r_dur_com"]
        mej_r = P["s_mej_r_par"] if scen == 1 else P["s_mej_r_com"]
        red_def = P["s_red_def_par"] if scen == 1 else P["s_red_def_com"]
        detecta = rng.random(Pa.shape) < P["s_p_det"]          # C1: deteccion anticipada
        u = rng.random(Pa.shape)
        Pa_n = np.where(es_eq & detecta, Pa * (1.0 - r_dur), Pa)
        if scen == 2:                                          # C4: elimina la recurrencia
            cuenta = {}
            for d in range(nd):
                for s in range(3):
                    if es_eq[d, s]:
                        k = (s, cau[d, s])
                        previos = cuenta.get(k, 0)
                        cuenta[k] = previos + 1
                        p_elim = P["s_eff_c4"] * (1.0 - math.exp(-previos / max(P["s_tau_c4"], 1e-9)))
                        if u[d, s] < p_elim:
                            Pa_n[d, s] = 0.0
        base_eq = Pa[es_eq].sum()
        rho = 0.0 if base_eq <= 0 else 1.0 - Pa_n[es_eq].sum() / base_eq
        Pa_n = np.where(es_esp, Pa * (1.0 - P["s_propag"] * rho), Pa_n)   # propagacion a la espera
    To_n = To + (Pa - Pa_n)
    tasa0 = np.where(To > 0, Proc / np.where(To > 0, To, 1.0), A["tasa_est"][None, :])
    tasa_n = np.minimum(tasa0 * (1.0 + mej_r), np.maximum(cap, tasa0))
    Proc_n = tasa_n * To_n
    q0 = np.where(Proc > 0, NC / np.where(Proc > 0, Proc, 1.0), A["q_est"][None, :])
    NC_n = Proc_n * q0 * (1.0 - red_def)
    Conf_n = Proc_n - NC_n
    # tiempo de respuesta a la alerta (eventos de equipo atendidos)
    mu = {0: P["s_t_sin"], 1: P["s_t_par"], 2: P["s_t_com"]}[scen]
    mu *= math.exp(rng.normal(0.0, P["s_disp_resp"]))
    sig = 0.6
    atendido = es_eq & (Pa_n > 1e-9)
    tr = np.where(atendido, rng.lognormal(math.log(mu) - sig ** 2 / 2, sig, Pa.shape), 0.0)
    hay = Pa_n > 1e-9
    causa_ef = np.where(hay, cau, SIN_PARADA)
    ev = (hay & (cau != SIN_PARADA)).astype(int)
    pl = pd.DataFrame(dict(
        est=np.tile(np.array([1, 2, 3]), nd), mes=np.repeat(mes, 3), causa=causa_ef.ravel(),
        Tp=Tp.ravel(), Pa=Pa_n.ravel(), To=To_n.ravel(), Teor=(cap * To_n).ravel(),
        Proc=Proc_n.ravel(), Conf=Conf_n.ravel(), NC=NC_n.ravel(), ev=ev.ravel(),
        rt=tr.ravel(), rn=atendido.ravel().astype(int)))
    fact = pl.groupby(["est", "mes", "causa"], as_index=False).sum()
    detalle = None
    if con_detalle:
        detalle = pd.DataFrame(dict(
            Fecha=np.repeat(A["fechas"][idx].to_numpy(), 3),
            Estación=[EST_CORTO[e] for e in pl["est"]], Causa_base=cau.ravel(),
            Parada_base_h=Pa.ravel(), Parada_DSS_h=Pa_n.ravel(), Operativo_base_h=To.ravel(),
            Operativo_DSS_h=To_n.ravel(), Procesados=Proc_n.ravel().round(1),
            Conformes=Conf_n.ravel().round(1), No_conformes=NC_n.ravel().round(1),
            Causa_DSS=causa_ef.ravel()))
        detalle["mes"] = np.repeat(mes, 3)
    return fact, detalle


def ejecutar_experimento(df, P, n_it, semilla, escenarios, progreso):
    """Corre n_it iteraciones por escenario (mismos dias muestreados en cada
    escenario dentro de una iteracion: comparacion pareada)."""
    A = preparar_sim(df)
    hechos, detalles = [], []
    for s in escenarios:
        for it in range(n_it):
            idx = muestrear_dias(A, np.random.default_rng([int(semilla), it]))
            rng = np.random.default_rng([int(semilla), it, s + 1])
            f, d = simular_iteracion(A, P, s, idx, rng, con_detalle=(it == 0))
            f["it"], f["scen"] = it, s
            hechos.append(f)
            if d is not None:
                d["scen"] = s
                detalles.append(d)
            progreso(s, it + 1, n_it)
    return dict(fact=pd.concat(hechos, ignore_index=True),
                detalle=pd.concat(detalles, ignore_index=True), n_it=n_it,
                semilla=semilla, escenarios=list(escenarios), P=dict(P))


def fact_base(df):
    d = df.assign(ev=_parada_ev(df))
    f = d.groupby(["Estacion_ID", "Mes", "Causa"], as_index=False).agg(
        Tp=("Tiempo_Planificado_h", "sum"), Pa=("Parada_No_Planificada_h", "sum"),
        To=("Tiempo_Operativo_h", "sum"), Teor=("Teorica", "sum"),
        Proc=("Bidones_Procesados", "sum"), Conf=("Bidones_Conformes", "sum"),
        NC=("Bidones_No_Conformes", "sum"), ev=("ev", "sum"))
    f = f.rename(columns={"Estacion_ID": "est", "Mes": "mes", "Causa": "causa"})
    f["rt"], f["rn"], f["it"], f["scen"] = 0.0, 0, 0, -1
    return f


COLS_SUMA = ["Tp", "Pa", "To", "Teor", "Proc", "Conf", "NC", "ev", "rt", "rn"]


def kpis_por_iteracion(fact):
    """Indicadores de linea por (iteracion, escenario) a partir de la tabla de hechos."""
    e = fact.groupby(["it", "scen", "est"], as_index=False)[COLS_SUMA].sum()
    e["D"] = e.To / e.Tp.replace(0, np.nan)
    e["R"] = e.Proc / e.Teor.replace(0, np.nan)
    e["Q"] = e.Conf / e.Proc.replace(0, np.nan)
    e["OEE"] = e.D * e.R * e.Q
    g = e.groupby(["it", "scen"])
    L = g[COLS_SUMA].sum()
    L["D"] = L.To / L.Tp.replace(0, np.nan)
    L["R"] = L.Proc / L.Teor.replace(0, np.nan)
    L["Q"] = L.Conf / L.Proc.replace(0, np.nan)
    L["OEE"] = g["OEE"].mean()
    L["MTBF"] = L.To / L.ev.replace(0, np.nan)
    L["MTTR"] = L.Pa / L.ev.replace(0, np.nan)
    es = e.sort_values("est")
    primera = es.groupby(["it", "scen"]).first()
    ultima = es.groupby(["it", "scen"]).last()
    L["rechazo"] = L.NC / primera.Proc.replace(0, np.nan)
    L["conformes"] = ultima.Conf
    L["t_resp"] = L.rt / L.rn.replace(0, np.nan)
    return L.reset_index(), e


def resumen_escenarios(L, claves):
    filas = []
    for s, g in L.groupby("scen"):
        for k in claves:
            x = g[k].dropna()
            if x.empty:
                continue
            filas.append(dict(scen=s, clave=k, media=x.mean(), sd=x.std(ddof=1) if len(x) > 1 else 0.0,
                              p5=x.quantile(0.05), p95=x.quantile(0.95),
                              pct_meta=(np.mean([cumple(k, v) for v in x]) * 100
                                        if k in METAS else np.nan)))
    return pd.DataFrame(filas)


def filtrar_fact(fact, est, mes, causa):
    f = fact
    if est != "Todas":
        f = f[f.est == int(est[1])]
    if mes != "Todos":
        f = f[f.mes == int(mes)]
    if causa != "Todas":
        f = f[f.causa == causa]
    return f


# =============================================================================
# DSS: VARIABLES, AMEF, C3 y C4 (logica del demostrador anterior, ya verificada)
# =============================================================================
N_LECTURAS = 120
INTERVALO_S = 5
LECTURAS_CONSEC = 3
VENTANA_OEE = 24
OEE_REF = 85.0
CICLO_OBJETIVO_S = 4.2
UMBRAL_MEDIO = 4
UMBRAL_ALTO = 7
POST_LECTURAS = 12
K_SIGMA_BANDA = 3.0
DET_OPORTUNA_LECTURAS = 10
FRACCION_PARCIAL = 0.3
EFECTOS = ("Efectiva", "Parcial", "No efectiva")
NPR_ROJO = 100
NPR_AMBAR = 70

EQUIPOS = {"Tanque": "E1", "Bomba": "E1", "Llenadora": "E2", "Tapadora": "E2",
           "Selladora": "E3", "Etiquetadora": "E3"}


def V(key, nombre, unidad, kind, lo=None, hi=None, nom=None, sd=0.0, sensor="—", entero=False):
    """kind: 'range' (lo-hi), 'max' (anomalia si supera hi), 'okng', 'runstop'."""
    return dict(key=key, nombre=nombre, unidad=unidad, kind=kind, lo=lo, hi=hi,
                nom=nom, sd=sd, sensor=sensor, entero=entero)


VARIABLES = {
    "Tanque": [V("nivel", "Nivel", "%", "range", 45, 90, 68, 4.0)],
    "Bomba": [
        V("presion", "Presión de descarga", "bar", "range", 2.0, 3.5, 2.75, 0.20),
        V("caudal", "Caudal", "L/min", "range", 80, 120, 100, 6.0, sensor="ifm SM Foodmag"),
        V("corriente", "Corriente del motor", "A", "max", 0, 7.5, 5.8, 0.40),
        V("vibracion", "Vibración", "mm/s", "range", 0, 4.5, 2.5, 0.40, sensor="ifm VVB001"),
    ],
    "Llenadora": [
        V("ciclo", "Tiempo de ciclo", "s", "max", 0, 4.2, 3.6, 0.15),
        V("presion_alim", "Presión de alimentación", "bar", "range", 1.5, 3.0, 2.25, 0.15, sensor="ifm PM1105"),
        V("conteo", "Conteo de envases", "unid/min", "range", 40, 60, 50, 3.0, sensor="SICK W4 Inox"),
        V("rech_llen", "Rechazos de llenado", "unid/50", "max", 0, 1, 0.25, 0, entero=True),
    ],
    "Tapadora": [
        V("torque", "Torque de tapado", "N·cm", "range", 10, 15, 12.5, 0.6, sensor="Mark-10 TT01"),
        V("vib_tapado", "Vibración del tapado", "mm/s", "range", 0, 4.0, 2.2, 0.40, sensor="ifm VVB001"),
    ],
    "Selladora": [
        V("rech_sello", "Rechazos de sello", "unid/50", "max", 0, 1, 0.25, 0, entero=True),
        V("cal_sello", "Calidad de sello", "OK/NG", "okng", sensor="Cognex In-Sight 2800"),
    ],
    "Etiquetadora": [
        V("estado", "Estado", "RUN/STOP", "runstop", sensor="Señal digital"),
        V("cal_etiq", "Calidad de etiqueta", "OK/NG", "okng", sensor="Cognex In-Sight 2800"),
    ],
}

FALLAS = {
    "Tanque": [("nivel", 28.0)],
    "Bomba": [("vibracion", 7.0), ("corriente", 8.6)],
    "Llenadora": [("ciclo", 4.9), ("rech_llen", 3.0)],
    "Tapadora": [("torque", 7.5)],
    "Selladora": [("rech_sello", 3.0)],
    "Etiquetadora": [("cal_etiq", "NG")],
}
PICOS = [("Bomba", "vibracion", 5.4), ("Llenadora", "presion_alim", 3.6),
         ("Etiquetadora", "cal_etiq", "NG")]

ESC_C1 = {
    "(a) Sin anomalía": "a",
    "(b) Anomalía en una sola estación": "b",
    "(c) Anomalías múltiples": "c",
    "(d) Cuello de botella oculto (OEE aceptable)": "d",
}


def C(causa, efecto, s, o, d, ejemplo=True):
    return dict(causa=causa, efecto=efecto, S=s, O=o, D=d, npr=s * o * d, ejemplo=ejemplo)


# Clave: (equipo, variable, direccion). ejemplo=False solo en las cifras de la tesis.
AMEF = {
    ("Tanque", "nivel", "baja"): [
        C("Obstrucción o falla de la válvula de alimentación", "Desabastecimiento de agua hacia la bomba", 7, 3, 3),
        C("Fuga en el tanque o conexiones", "Pérdida de agua tratada y nivel inestable", 6, 3, 4),
        C("Sensor de nivel descalibrado", "Lectura errónea; riesgo de marcha en seco", 4, 4, 5)],
    ("Tanque", "nivel", "alta"): [
        C("Control de llenado defectuoso (válvula no cierra)", "Desborde y desperdicio de agua", 6, 3, 3),
        C("Salida del tanque restringida (filtro obstruido)", "Acumulación y caudal reducido hacia la bomba", 5, 4, 4)],
    ("Bomba", "presion", "baja"): [
        C("Desgaste del impulsor", "Menor presión y caudal aguas abajo", 7, 4, 4),
        C("Cavitación por succión insuficiente", "Ruido, vibración y daño del impulsor", 7, 3, 4),
        C("Fuga en línea de descarga", "Pérdida de presión y de producto", 6, 3, 3)],
    ("Bomba", "presion", "alta"): [
        C("Obstrucción o filtro saturado en la descarga", "Sobrepresión y mayor esfuerzo del motor", 6, 4, 3),
        C("Válvula de descarga parcialmente cerrada", "Sobrepresión y menor caudal", 5, 3, 2)],
    ("Bomba", "caudal", "baja"): [
        C("Desgaste del impulsor", "Caudal insuficiente para la llenadora", 7, 4, 4),
        C("Filtro de línea obstruido", "Restricción de flujo", 5, 5, 3),
        C("Ingreso de aire en la succión", "Caudal pulsante y cavitación", 6, 3, 4)],
    ("Bomba", "caudal", "alta"): [
        C("Fuga o rotura aguas abajo", "Pérdida de producto y presión baja", 6, 2, 4),
        C("Variador o válvula fuera de consigna", "Exceso de caudal y mayor consumo", 4, 3, 3)],
    ("Bomba", "corriente", "alta"): [
        C("Desgaste de rodamiento", "Mayor fricción y consumo del motor", 7, 4, 3, ejemplo=False),
        C("Obstrucción en línea de succión", "Sobrecarga del motor y baja de caudal", 6, 3, 4, ejemplo=False),
        C("Desbalance de tensión de alimentación", "Calentamiento del motor y disparos de protección", 5, 2, 5)],
    ("Bomba", "vibracion", "alta"): [
        C("Desgaste de rodamiento", "Vibración creciente y falla del eje", 7, 5, 3),
        C("Desalineación del acople motor-bomba", "Desgaste prematuro de rodamientos y sellos", 6, 4, 4),
        C("Cavitación", "Daño del impulsor y ruido", 6, 3, 4),
        C("Pernos de anclaje flojos", "Vibración transmitida a la estructura", 4, 4, 3)],
    ("Llenadora", "ciclo", "alta"): [
        C("Desgaste de válvula dosificadora", "Ciclo prolongado: la llenadora limita el ritmo de la línea", 5, 5, 4, ejemplo=False),
        C("Baja presión de alimentación", "Llenado más lento por envase", 5, 4, 3),
        C("Aire comprimido insuficiente en actuadores", "Apertura y cierre lentos de la válvula", 5, 3, 4),
        C("Parámetros de dosificación mal ajustados", "Tiempo de llenado mayor al objetivo", 4, 4, 3)],
    ("Llenadora", "presion_alim", "baja"): [
        C("Bomba de alimentación degradada", "Llenado insuficiente o lento", 7, 3, 4),
        C("Filtro de alimentación obstruido", "Caída de presión en el cabezal", 5, 5, 3)],
    ("Llenadora", "presion_alim", "alta"): [
        C("Válvula de alivio atascada", "Sobrepresión; riesgo de sobrellenado", 6, 2, 4),
        C("Regulador de presión desajustado", "Salpicado y sobrellenado", 4, 3, 3)],
    ("Llenadora", "conteo", "baja"): [
        C("Paros menores por atasco de envases", "Menor producción por minuto", 5, 6, 3),
        C("Sensor fotoeléctrico sucio o desalineado", "Subconteo de envases", 3, 5, 4),
        C("Velocidad del transportador reducida", "Menor ritmo de la línea", 4, 4, 3)],
    ("Llenadora", "conteo", "alta"): [
        C("Sensor fotoeléctrico con doble conteo", "Sobreconteo de envases", 3, 3, 5),
        C("Velocidad de línea fuera de consigna", "Ritmo superior al diseño", 4, 3, 3)],
    ("Llenadora", "rech_llen", "alta"): [
        C("Desgaste de válvula dosificadora", "Sub/sobrellenado y envases rechazados", 6, 5, 4),
        C("Variación de presión de alimentación", "Volumen de llenado inconsistente", 5, 4, 4),
        C("Medición de dosificación descalibrada", "Dosis fuera de tolerancia", 6, 3, 5)],
    ("Tapadora", "vib_tapado", "alta"): [
        C("Desgaste de rodamiento del cabezal de tapado", "Vibración y torque inestable", 6, 4, 3),
        C("Aflojamiento mecánico del cabezal", "Tapado irregular", 5, 4, 4)],
    ("Tapadora", "torque", "baja"): [
        C("Desgaste del embrague del cabezal", "Tapas flojas; riesgo de fuga", 7, 4, 7),
        C("Ajuste de torque desregulado", "Tapado insuficiente", 5, 4, 3)],
    ("Tapadora", "torque", "alta"): [
        C("Tapas con rosca defectuosa", "Tapas dañadas por exceso de apriete", 5, 3, 5),
        C("Embrague de torque mal ajustado", "Sobreapriete y rotura de tapas", 5, 3, 3)],
    ("Selladora", "rech_sello", "alta"): [
        C("Desgaste del empaque o sello de tapa", "Fugas y envases rechazados", 6, 4, 4),
        C("Tapas o envases con defecto dimensional", "Sellado deficiente", 5, 4, 5),
        C("Torque de tapado fuera de especificación", "Tapas flojas o dañadas", 6, 4, 3)],
    ("Selladora", "cal_sello", "NG"): [
        C("Defecto en lote de tapas", "Rechazo masivo por calidad de tapa", 6, 4, 4),
        C("Tapa mal asentada por desalineación del cabezal", "Sellado incorrecto", 7, 4, 3),
        C("Lente de visión sucio o iluminación degradada", "Falsos rechazos del sistema de visión", 3, 5, 5)],
    ("Etiquetadora", "estado", "STOP"): [
        C("Falla del servomotor o del sensor de etiquetas", "Paro de la etiquetadora y de la línea", 7, 3, 7),
        C("Atasco de etiquetas o rollo agotado", "Paros frecuentes", 5, 6, 2),
        C("Paro por protecciones de seguridad", "Paro por interlock activado", 4, 4, 2)],
    ("Etiquetadora", "cal_etiq", "NG"): [
        C("Adhesivo o etiqueta defectuosa (lote)", "Etiquetas despegadas o torcidas", 5, 4, 4),
        C("Envase húmedo por condensación", "Mala adhesión de la etiqueta", 4, 5, 4),
        C("Desalineación del aplicador", "Etiquetas fuera de posición", 5, 5, 3)],
}

APTO_OPERADOR = {
    "Salida del tanque restringida (filtro obstruido)", "Obstrucción o filtro saturado en la descarga",
    "Válvula de descarga parcialmente cerrada", "Filtro de línea obstruido",
    "Variador o válvula fuera de consigna", "Parámetros de dosificación mal ajustados",
    "Filtro de alimentación obstruido", "Regulador de presión desajustado",
    "Paros menores por atasco de envases", "Sensor fotoeléctrico sucio o desalineado",
    "Torque de tapado fuera de especificación", "Ajuste de torque desregulado",
    "Lente de visión sucio o iluminación degradada", "Atasco de etiquetas o rollo agotado",
    "Desalineación del aplicador",
}
NIVELES = ("Baja", "Media", "Alta")
TABLA7_TESIS = {("Alta", "Alta", "Baja"): "Crítica", ("Alta", "Baja", "Alta"): "Alta",
                ("Media", "Media", "Media"): "Media", ("Baja", "Baja", "Alta"): "Baja"}
ACCION_PRIORIDAD = {"Crítica": "mantenimiento inmediato", "Alta": "mantenimiento planificado próximo",
                    "Media": "mantenimiento planificado", "Baja": "mantenimiento de mejora"}
PLAZO_PRIORIDAD = {
    "Crítica": "Atención inmediata, con detención de la estación",
    "Alta": "Dentro del mismo turno",
    "Media": "Mantenimiento planificado en un máximo de 48 h o en la parada preventiva del sábado",
    "Baja": "Mantenimiento autónomo en la siguiente limpieza",
}
COLOR_PRIO = {"Crítica": ("#c62828", "#fff"), "Alta": ("#ef6c00", "#fff"),
              "Media": ("#ffb300", "#3e2723"), "Baja": ("#1e88e5", "#fff")}


def nivel_s_o(x):
    return "Alta" if x >= UMBRAL_ALTO else "Media" if x >= UMBRAL_MEDIO else "Baja"


def nivel_deteccion(d):
    return "Baja" if d >= UMBRAL_ALTO else "Media" if d >= UMBRAL_MEDIO else "Alta"


def prioridad_por_regla(sev, frec, det):
    if sev == "Alta":
        if frec == "Alta" or det == "Baja":
            return "Crítica", "Severidad Alta y (Frecuencia Alta o Detectabilidad Baja) → Crítica"
        return "Alta", "Severidad Alta sin agravantes de frecuencia ni detectabilidad → Alta"
    if sev == "Media":
        if frec == "Alta" and det == "Baja":
            return "Alta", "Severidad Media con Frecuencia Alta y Detectabilidad Baja → Alta"
        return "Media", "Severidad Media sin ambos agravantes → Media"
    if frec == "Alta" or det == "Baja":
        return "Media", "Severidad Baja con Frecuencia Alta o Detectabilidad Baja → Media"
    return "Baja", "Severidad Baja sin agravantes → Baja"


MATRIZ = {}
for _s in NIVELES:
    for _f in NIVELES:
        for _d in NIVELES:
            _p, _r = prioridad_por_regla(_s, _f, _d)
            MATRIZ[(_s, _f, _d)] = dict(prioridad=_p, regla=_r, ejemplo=(_s, _f, _d) not in TABLA7_TESIS)
assert len(MATRIZ) == 27
assert all(MATRIZ[k]["prioridad"] == p for k, p in TABLA7_TESIS.items())


def evaluar_c3(causa):
    sev, frec, det = nivel_s_o(causa["S"]), nivel_s_o(causa["O"]), nivel_deteccion(causa["D"])
    fila = MATRIZ[(sev, frec, det)]
    prio = fila["prioridad"]
    apto = causa["causa"] in APTO_OPERADOR
    if prio in ("Media", "Baja") and apto:
        tipo, accion = "Autónomo", "mantenimiento autónomo (tarea del operador)"
    elif prio == "Baja":
        tipo, accion = "De mejora", ACCION_PRIORIDAD["Baja"]
    else:
        tipo, accion = "Planificado", ACCION_PRIORIDAD[prio]
    return dict(sev=sev, frec=frec, det=det, prioridad=prio, regla=fila["regla"],
                fila_ejemplo=fila["ejemplo"], apto=apto, tipo=tipo, accion=accion,
                inmediato=prio == "Crítica")


def limites(v, banda=None):
    return banda if banda else (v["lo"], v["hi"])


def direccion_desviacion(v, valor, banda=None):
    if v["kind"] == "okng":
        return "NG"
    if v["kind"] == "runstop":
        return "STOP"
    if v["kind"] == "range" and valor < limites(v, banda)[0]:
        return "baja"
    return "alta"


def causas_vigentes(equipo, key, direccion, ajustes=None):
    salida = []
    for c in AMEF.get((equipo, key, direccion), []):
        c = dict(c)
        aj = (ajustes or {}).get((equipo, key, direccion, c["causa"]))
        c["ajustada"] = aj is not None
        if aj:
            c["O"], c["D"] = aj["O"], aj["D"]
            c["npr"] = c["S"] * c["O"] * c["D"]
        salida.append(c)
    return salida


def diagnosticar(evento, ajustes=None, descartadas=None):
    """Causas probables ordenadas por NPR; las descartadas por C4 van al final."""
    v, eq = evento["var"], evento["equipo"]
    direccion = direccion_desviacion(v, evento["valor_obs"], evento.get("banda"))
    causas = causas_vigentes(eq, v["key"], direccion, ajustes)
    for c in causas:
        c["descartada"] = (eq, v["key"], direccion, c["causa"]) in (descartadas or ())
    causas.sort(key=lambda c: (c["descartada"], -c["npr"]))
    return direccion, causas


def color_npr(npr):
    return "rojo" if npr >= NPR_ROJO else "ambar" if npr >= NPR_AMBAR else "amarillo"


NPR_COL = {"rojo": ("#fde2e1", "#b42318"), "ambar": ("#fff0d6", "#b54708"),
           "amarillo": ("#fff9c4", "#7a6a00")}


# ---------------------------------------------------------------- simulador
def serie_base(v, rng, n=N_LECTURAS):
    k = v["kind"]
    if k == "okng":
        return np.array(["OK"] * n, dtype=object)
    if k == "runstop":
        return np.array(["RUN"] * n, dtype=object)
    if v["entero"]:
        return np.minimum(rng.poisson(v["nom"], n), v["hi"]).astype(float)
    margen = 0.03 * (v["hi"] - v["lo"])
    return np.clip(rng.normal(v["nom"], v["sd"], n), v["lo"] + margen, v["hi"] - margen)


def valores_post(v, rng, n, modo, objetivo=None):
    """Lecturas tras la intervencion. modo: normal | falla | parcial."""
    base = serie_base(v, rng, n)
    if modo == "normal":
        return base
    if v["kind"] in ("okng", "runstop"):
        falla = np.array(["NG" if v["kind"] == "okng" else "STOP"] * n, dtype=object)
    elif v["entero"]:
        falla = np.maximum(rng.poisson(max(objetivo, 0), n), v["hi"] + 1).astype(float)
    else:
        falla = np.maximum(rng.normal(objetivo, v["sd"], n), 0)
    if modo == "falla":
        return falla
    mezcla = base.copy()
    mask = ((n - 1 - np.arange(n)) % 5) < 3
    mezcla[mask] = falla[mask]
    return mezcla


def extender_datos(datos, fallas, eq_int, key_int, efecto, rng, n=POST_LECTURAS):
    ultimo = next(iter(datos.values())).index[-1]
    indice = pd.date_range(ultimo + pd.Timedelta(seconds=INTERVALO_S), periods=n, freq=f"{INTERVALO_S}s")
    nuevos = {}
    for eq, df in datos.items():
        cols = {}
        objetivos = dict(fallas.get(eq, {}).get("vars", []))
        for v in VARIABLES[eq]:
            obj = objetivos.get(v["key"])
            if eq == eq_int and v["key"] == key_int:
                modo = {"Efectiva": "normal", "Parcial": "parcial", "No efectiva": "falla"}[efecto]
                if obj is None and modo != "normal":
                    obj = df[v["key"]].iloc[-1]
            elif obj is not None:
                modo = "falla"
            else:
                modo = "normal"
            cols[v["key"]] = valores_post(v, rng, n, modo, obj)
        nuevos[eq] = pd.concat([df, pd.DataFrame(cols, index=indice)])
    return nuevos


def aplicar_falla(v, s, falla, rng):
    for key, objetivo in falla["vars"]:
        if key != v["key"]:
            continue
        ini, rampa = falla["inicio"], falla["rampa"]
        if v["kind"] in ("okng", "runstop"):
            s[ini:] = objetivo
            continue
        t = np.arange(N_LECTURAS)
        frac = np.clip((t - ini + 1) / rampa, 0, 1)
        media = v["nom"] + (objetivo - v["nom"]) * frac
        if v["entero"]:
            nuevo = rng.poisson(np.maximum(media, 0)).astype(float)
        else:
            nuevo = np.maximum(rng.normal(media, v["sd"]), 0)
        s[ini:] = nuevo[ini:]
    return s


def definir_fallas(escenario, rng):
    equipos = list(VARIABLES)
    if escenario == "a":
        return {}
    if escenario == "b":
        e = str(rng.choice(equipos))
        return {e: dict(inicio=int(rng.integers(45, 70)), rampa=8, vars=FALLAS[e])}
    if escenario == "c":
        k = int(rng.integers(2, 4))
        elegidos = [str(e) for e in rng.choice(equipos, size=k, replace=False)]
        return {e: dict(inicio=int(rng.integers(45, 70)), rampa=8, vars=FALLAS[e]) for e in elegidos}
    # (d) cuello de botella oculto: la llenadora se degrada lento; conteo sigue en rango
    return {"Llenadora": dict(inicio=40, rampa=25, vars=[("ciclo", 4.55), ("conteo", 47.5)])}


def generar_datos(escenario, rng):
    indice = pd.date_range(pd.Timestamp.now().floor("s"), periods=N_LECTURAS, freq=f"{INTERVALO_S}s")
    fallas = definir_fallas(escenario, rng)
    datos = {}
    for eq, variables in VARIABLES.items():
        cols = {}
        for v in variables:
            s = serie_base(v, rng)
            if eq in fallas:
                s = aplicar_falla(v, s, fallas[eq], rng)
            cols[v["key"]] = s
        datos[eq] = pd.DataFrame(cols, index=indice)
    for eq, key, valor in PICOS:
        i = int(rng.integers(8, 38))
        datos[eq].iloc[i, datos[eq].columns.get_loc(key)] = valor
    return datos, fallas


# ---------------------------------------------------------------- C1
def fuera_de_rango(v, s, banda=None):
    k = v["kind"]
    lo, hi = limites(v, banda) if k in ("range", "max") else (None, None)
    if k == "range":
        return (s < lo) | (s > hi)
    if k == "max":
        return s > hi
    if k == "okng":
        return s == "NG"
    return s == "STOP"


def analizar(datos, bandas=None):
    bandas = bandas or {}
    res = {}
    for eq, df in datos.items():
        fuera = pd.DataFrame({v["key"]: fuera_de_rango(v, df[v["key"]], bandas.get((eq, v["key"])))
                              for v in VARIABLES[eq]})
        sost = fuera.astype(int).rolling(LECTURAS_CONSEC).sum().eq(LECTURAS_CONSEC)
        res[eq] = dict(fuera=fuera, sost=sost, alarma=sost.any(axis=1))
    return res


def calcular_oee(datos):
    """OEE demostrativo de la linea (%) en ventana movil (no es el OEE de planta)."""
    w = VENTANA_OEE
    disp = (datos["Etiquetadora"]["estado"] == "RUN").astype(float).rolling(w, min_periods=w // 2).mean()
    rend = (CICLO_OBJETIVO_S / datos["Llenadora"]["ciclo"]).clip(upper=1.0).rolling(w, min_periods=w // 2).mean()
    rech = datos["Llenadora"]["rech_llen"] + datos["Selladora"]["rech_sello"]
    cal = 1.0 - (rech / 50.0).rolling(w, min_periods=w // 2).mean()
    return (disp * rend * cal * 100.0).clip(0, 100).bfill()


def oee_segmento(datos, i0, i1):
    sl = slice(i0, i1)
    disp = (datos["Etiquetadora"]["estado"].iloc[sl] == "RUN").mean()
    rend = (CICLO_OBJETIVO_S / datos["Llenadora"]["ciclo"].iloc[sl]).clip(upper=1.0).mean()
    rech = datos["Llenadora"]["rech_llen"].iloc[sl] + datos["Selladora"]["rech_sello"].iloc[sl]
    cal = 1.0 - (rech / 50.0).mean()
    return float(min(max(disp * rend * cal * 100.0, 0.0), 100.0))


def verificar_resolucion(v, post, banda=None):
    ok = ~fuera_de_rango(v, post, banda)
    frac = float(ok.mean())
    if bool(ok.iloc[-LECTURAS_CONSEC:].all()):
        return "Sí", frac
    if frac >= FRACCION_PARCIAL:
        return "Parcial", frac
    return "No", frac


def formato_valor(v, valor):
    if v["kind"] in ("okng", "runstop"):
        return str(valor)
    if v["entero"]:
        return f"{int(valor)} {v['unidad']}"
    return f"{float(valor):.2f} {v['unidad']}"


def valor_despues(v, post):
    if v["kind"] in ("okng", "runstop"):
        malo = "NG" if v["kind"] == "okng" else "STOP"
        return f"{malo} {int((post == malo).sum())}/{len(post)}"
    return formato_valor(v, float(post.astype(float).mean()))


def calcular_banda(v, post):
    """Banda de alerta temprana: promedio post ± 3σ, recortada al rango de diseño."""
    if v["kind"] not in ("range", "max") or v["entero"]:
        return None
    x = post.astype(float)
    mu, sd = float(x.mean()), float(x.std(ddof=1))
    hi = min(v["hi"], mu + K_SIGMA_BANDA * sd)
    lo = v["lo"] if v["lo"] == 0 else max(v["lo"], mu - K_SIGMA_BANDA * sd)
    if hi <= lo or (round(lo, 2), round(hi, 2)) == (v["lo"], v["hi"]):
        return None
    return (round(lo, 2), round(hi, 2))


def formato_umbral(v, banda=None):
    u = v["unidad"]
    if v["kind"] in ("range", "max") and banda:
        lo, hi = banda
        if v["kind"] == "max" or lo == 0:
            return f"≤ {hi:g} {u} ⚑ alerta"
        return f"{lo:g} – {hi:g} {u} ⚑ alerta"
    if v["kind"] == "range":
        return f"{v['lo']:g} – {v['hi']:g} {u}"
    if v["kind"] == "max":
        return f"≤ {v['hi']:g} {u}"
    return "OK" if v["kind"] == "okng" else "RUN"


def mmss(i):
    s = int(i) * INTERVALO_S
    return f"{s // 60:02d}:{s % 60:02d}"


def eventos_c1(datos, res, bandas, i):
    """Un evento por variable con anomalia confirmada en la lectura i."""
    eventos = []
    for eq in VARIABLES:
        sost = res[eq]["sost"]
        for v in VARIABLES[eq]:
            s = sost[v["key"]].values
            if i >= len(s) or not s[i]:
                continue
            j = i
            while j > 0 and s[j - 1]:
                j -= 1
            eventos.append(dict(
                equipo=eq, estacion=EQUIPOS[eq], var=v, lectura=j,
                banda=bandas.get((eq, v["key"])), valor_obs=datos[eq][v["key"]].iloc[j],
                marca=datos[eq].index[j]))
    return eventos





# =============================================================================
# VALORES POR DEFECTO (de la tesis) Y CONFIGURACION EDITABLE
# =============================================================================
AMEF_DEF = copy.deepcopy(AMEF)
VARIABLES_DEF = copy.deepcopy(VARIABLES)
MATRIZ_DEF = copy.deepcopy(MATRIZ)
METAS_DEF = dict(METAS)
META_TXT_DEF = dict(META_TXT)
LECT_DEF, MEDIO_DEF, ALTO_DEF = LECTURAS_CONSEC, UMBRAL_MEDIO, UMBRAL_ALTO

# clave: (valor mostrado, unidad, condición)
META_UI = {
    "OEE": (65.0, "%", "≥"), "D": (92.0, "%", "≥"), "R": (72.0, "%", "≥"), "Q": (99.0, "%", "≥"),
    "rechazo": (2.5, "%", "≤"), "Pa": (180.0, "h", "≤"), "MTBF": (9.0, "h", "≥"),
    "MTTR": (0.75, "h", "≤"), "t_resp": (15.0, "min", "≤"),
}
KPI_INFO = {  # (nombre, qué mide en palabras simples, fórmula, módulo)
    "OEE": ("OEE de línea", "Qué parte del tiempo planificado se produce, al ritmo ideal y bien hecho.", "D × R × Q", "C1"),
    "D": ("Disponibilidad (D)", "Qué porcentaje del tiempo planificado la línea estuvo funcionando.", "Tₒ / Tₚ", "C1, C4"),
    "R": ("Rendimiento (R)", "Qué tan rápido produce comparado con su capacidad ideal.", "(tᵢ × Nₜ) / Tₒ", "C1"),
    "Q": ("Calidad (Q)", "Qué porcentaje de lo producido sale conforme a la primera.", "Nₐ / Nₜ", "C1"),
    "rechazo": ("Índice de rechazo de línea", "De cada 100 bidones que entran, cuántos salen no conformes.",
                "No conformes / bidones que ingresan a E1", "C1, C4"),
    "Pa": ("Parada no planificada", "Horas perdidas por paradas que nadie programó.", "Σ duración de paradas", "C1, C4"),
    "MTBF": ("MTBF", "Cada cuántas horas de trabajo ocurre una parada.", "Tₒ / número de fallas", "C2, C4"),
    "MTTR": ("MTTR", "Cuánto dura, en promedio, cada parada.", "Horas de reparación / reparaciones", "C4"),
    "t_resp": ("Tiempo de respuesta a la alerta", "Minutos entre la alerta y la primera atención.",
               "Primera atención − hora de alerta", "C1, C4"),
}

CAUSAS_VALIDAS = [
    "Avería de llenadora-tapadora", "Avería de lavadora de bidones", "Avería de selladora-etiquetadora",
    "Ajuste o limpieza correctiva", "Espera de proceso (línea detenida)", "Falta de tapas o precintos",
    "Falta de bidones vacíos (retorno de reparto)", "Falta de personal", "Corte de suministro de agua",
    "Corte de energía eléctrica", SIN_PARADA,
]
COLS_PLANTILLA = ["Fecha", "Estacion_ID", "Tiempo_Turno_h", "Parada_Planificada_h", "Parada_No_Planificada_h",
                  "Causa_Parada_No_Planificada", "Capacidad_Nominal_bidones_h", "Bidones_Procesados",
                  "Bidones_Conformes"]
INSTRUCCIONES_COLS = [  # (columna, qué poner, ejemplo, obligatoria)
    ("Fecha", "Día de trabajo (dd/mm/aaaa). No incluyas domingos ni feriados.", "05/01/2026", "Sí"),
    ("Estacion_ID", "Número de la estación: 1 = Lavado y sanitizado, 2 = Llenado y tapado, 3 = Sellado y etiquetado.", "2", "Sí"),
    ("Tiempo_Turno_h", "Horas del turno de esa estación ese día.", "6", "Sí"),
    ("Parada_Planificada_h", "Horas de paradas programadas (limpieza, sanitización, mantenimiento preventivo).", "0,75", "Sí"),
    ("Parada_No_Planificada_h", "Horas perdidas por paradas que nadie programó. Pon 0 si no hubo.", "1,25", "Sí"),
    ("Causa_Parada_No_Planificada", "Causa de esa parada. Elige una de la lista de causas válidas. Si no hubo parada, «Sin parada no planificada».",
     "Avería de llenadora-tapadora", "Sí"),
    ("Capacidad_Nominal_bidones_h", "Bidones por hora que la estación puede hacer a su ritmo ideal.", "95", "Sí"),
    ("Bidones_Procesados", "Bidones que pasaron por la estación ese día (buenos y malos).", "354", "Sí"),
    ("Bidones_Conformes", "Bidones buenos a la primera. No puede ser mayor que los procesados.", "350", "Sí"),
]


def aplicar_config():
    """Aplica a la lógica (C1, C2, C3 y metas) lo que el usuario configuró en el paso 3 y en el paso 1."""
    global LECTURAS_CONSEC, UMBRAL_MEDIO, UMBRAL_ALTO
    s = st.session_state
    LECTURAS_CONSEC = int(s.get("w_p_lect", LECT_DEF))
    UMBRAL_MEDIO = int(s.get("w_p_medio", MEDIO_DEF))
    UMBRAL_ALTO = int(s.get("w_p_alto", ALTO_DEF))
    if UMBRAL_MEDIO >= UMBRAL_ALTO:
        UMBRAL_MEDIO, UMBRAL_ALTO = MEDIO_DEF, ALTO_DEF
    for (eq, key), (lo, hi) in s.get("rangos_ov", {}).items():
        for v in VARIABLES.get(eq, []):
            if v["key"] == key:
                v["lo"], v["hi"] = lo, hi
    for ident, (S, O, D) in s.get("amef_ov", {}).items():
        eq, vkey, direc, causa = ident
        for c in AMEF.get((eq, vkey, direc), []):
            if c["causa"] == causa:
                c.update(S=int(S), O=int(O), D=int(D), npr=int(S) * int(O) * int(D))
    mat = s.get("matriz")
    if mat:
        for k, prio in mat.items():
            base = MATRIZ_DEF[k]
            if prio == base["prioridad"] and not base["ejemplo"]:
                MATRIZ[k] = dict(base)
            else:
                MATRIZ[k] = dict(prioridad=prio, ejemplo=base["ejemplo"],
                                 regla=("Matriz configurada en el paso 3" if prio != base["prioridad"] else base["regla"]))
    for k, (v0, u, cond) in META_UI.items():
        v = float(s.get("metas", {}).get(k, v0))
        METAS[k] = (">=" if cond == "≥" else "<=", v / 100.0 if u == "%" else v)
        META_TXT[k] = f"{cond} {n(v, 0 if abs(v - round(v)) < 1e-9 else 2)} {u}"


# =============================================================================
# CARGA, VALIDACION Y GENERACION DE DATOS
# =============================================================================
def g_(x):
    return f"{x:g}".replace('.', ',')


def _err(errores, fila, columna, problema):
    errores.append(dict(Fila=fila, Columna=columna, Problema=problema))


def _norm_est(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        t = str(v).strip().upper()
        return int(t[1:]) if len(t) == 2 and t[0] == "E" and t[1].isdigit() else np.nan


def finalizar_df(df):
    """Agrega las columnas derivadas que usa toda la app."""
    df = df.copy()
    df["Estacion_ID"] = df["Estacion_ID"].astype(int)
    df["Fecha"] = pd.to_datetime(df["Fecha"])
    df["Mes"] = df["Fecha"].dt.year * 100 + df["Fecha"].dt.month
    df["Est"] = df["Estacion_ID"].map(EST_CORTO)
    df["Causa"] = df["Causa_Parada_No_Planificada"]
    df["Teorica"] = df["Capacidad_Nominal_bidones_h"] * df["Tiempo_Operativo_h"]
    return df.sort_values(["Fecha", "Estacion_ID"]).reset_index(drop=True)


def leer_base(origen):
    """Lee la hoja OEE_Diario con el encabezado en la fila 1 (plantilla) o en la fila 5 (base de la tesis).
    Devuelve (df, errores, avisos). Si hay errores, df es None."""
    errores, avisos = [], []
    try:
        xl = pd.ExcelFile(origen)
    except Exception as exc:
        _err(errores, "—", "—", f"No pude abrir el archivo como Excel (.xlsx): {exc}")
        return None, errores, avisos
    if "OEE_Diario" in xl.sheet_names:
        hoja = "OEE_Diario"
    elif len(xl.sheet_names) == 1:
        hoja = xl.sheet_names[0]
    else:
        _err(errores, "—", "—", "No encuentro la hoja «OEE_Diario». Usa la plantilla de la app o renombra la hoja.")
        return None, errores, avisos
    crudo = xl.parse(hoja, header=None)
    h = None
    for i in range(min(15, len(crudo))):
        vals = [str(v).strip() for v in crudo.iloc[i].tolist()]
        if "Fecha" in vals and "Estacion_ID" in vals:
            h = i
            break
    if h is None:
        _err(errores, "—", "—", "No encuentro la fila de encabezados: debe tener las columnas «Fecha» y «Estacion_ID» "
                                "(en la fila 1 si usas la plantilla).")
        return None, errores, avisos
    datos = crudo.iloc[h + 1:].copy()
    datos.columns = [str(c).strip() for c in crudo.iloc[h].tolist()]
    datos["_fila"] = datos.index + 1                      # número de fila en Excel
    obligatorias = ["Fecha", "Estacion_ID", "Parada_No_Planificada_h", "Causa_Parada_No_Planificada",
                    "Capacidad_Nominal_bidones_h", "Bidones_Procesados", "Bidones_Conformes"]
    tiene_tp = "Tiempo_Planificado_h" in datos.columns
    if not tiene_tp:
        obligatorias += ["Tiempo_Turno_h", "Parada_Planificada_h"]
    faltan = [c for c in obligatorias if c not in datos.columns]
    if faltan:
        _err(errores, h + 1, ", ".join(faltan), "Faltan columnas en el encabezado. Descarga la plantilla para ver cuáles son.")
        return None, errores, avisos
    fecha = pd.to_datetime(datos["Fecha"], errors="coerce")
    est = datos["Estacion_ID"].map(_norm_est)
    valido = fecha.notna() & est.isin([1, 2, 3])
    if not valido.any():
        _err(errores, "—", "Fecha / Estacion_ID", "No hay ninguna fila con fecha y estación (1, 2 o 3) válidas.")
        return None, errores, avisos
    ultimo = valido[valido].index.max()                    # lo que sigue son notas al pie: se ignora
    datos, fecha, est, valido = (x.loc[:ultimo] for x in (datos, fecha, est, valido))
    vacia = datos[obligatorias].isna().all(axis=1)
    for i in datos.index[~valido & ~vacia]:
        fila = int(datos.at[i, "_fila"])
        if pd.isna(fecha[i]):
            _err(errores, fila, "Fecha", "La fecha no es válida (usa dd/mm/aaaa).")
        else:
            _err(errores, fila, "Estacion_ID", "La estación debe ser 1, 2 o 3.")
    datos, fecha, est = datos[valido], fecha[valido], est[valido]
    num = {}
    cols_num = ["Parada_No_Planificada_h", "Capacidad_Nominal_bidones_h", "Bidones_Procesados", "Bidones_Conformes"]
    cols_num += ["Tiempo_Planificado_h"] if tiene_tp else ["Tiempo_Turno_h", "Parada_Planificada_h"]
    for c in cols_num:
        x = pd.to_numeric(datos[c], errors="coerce")
        for i in datos.index[x.isna()]:
            vacio = pd.isna(datos.at[i, c])
            _err(errores, int(datos.at[i, "_fila"]), c, "Está vacío: escribe un número." if vacio else
                 f"«{datos.at[i, c]}» no es un número.")
        for i in datos.index[x < 0]:
            _err(errores, int(datos.at[i, "_fila"]), c, f"No puede ser negativo ({g_(x[i])}).")
        num[c] = x
    if errores:
        return None, errores[:300], avisos
    Tp = num["Tiempo_Planificado_h"] if tiene_tp else num["Tiempo_Turno_h"] - num["Parada_Planificada_h"]
    Pa, cap, proc, conf = (num[c] for c in ("Parada_No_Planificada_h", "Capacidad_Nominal_bidones_h",
                                            "Bidones_Procesados", "Bidones_Conformes"))
    for i in datos.index:
        fila = int(datos.at[i, "_fila"])
        if Tp[i] <= 0:
            _err(errores, fila, "Tiempo planificado", f"El tiempo planificado debe ser mayor que 0 (es {g_(Tp[i])} h).")
        elif Pa[i] > Tp[i] + 1e-9:
            _err(errores, fila, "Parada_No_Planificada_h",
                 f"La parada no planificada ({g_(Pa[i])} h) no puede ser mayor que el tiempo planificado ({g_(Tp[i])} h).")
        if cap[i] <= 0:
            _err(errores, fila, "Capacidad_Nominal_bidones_h", "La capacidad debe ser mayor que 0.")
        if conf[i] > proc[i] + 1e-9:
            _err(errores, fila, "Bidones_Conformes",
                 f"Los bidones conformes ({g_(conf[i])}) no pueden ser más que los procesados ({g_(proc[i])}).")
        elif proc[i] > cap[i] * max(Tp[i] - Pa[i], 0) * 1.02 + 1:
            avisos.append(f"Fila {fila}: los bidones procesados ({g_(proc[i])}) superan la capacidad posible en el tiempo "
                          f"operativo; revisa la capacidad o el tiempo.")
    validas = {c.lower(): c for c in CAUSAS_VALIDAS}
    causas = []
    for i in datos.index:
        fila = int(datos.at[i, "_fila"])
        txt = "" if pd.isna(datos.at[i, "Causa_Parada_No_Planificada"]) else str(datos.at[i, "Causa_Parada_No_Planificada"]).strip()
        if txt == "":
            if Pa[i] > 0:
                _err(errores, fila, "Causa_Parada_No_Planificada", f"Hay {g_(Pa[i])} h de parada pero falta la causa.")
            causas.append(SIN_PARADA)
            continue
        c = validas.get(txt.lower())
        if c is None:
            _err(errores, fila, "Causa_Parada_No_Planificada", f"«{txt}» no es una causa válida. Elige una de la lista de la hoja Instrucciones.")
            causas.append(txt)
            continue
        if Pa[i] > 0 and c == SIN_PARADA:
            _err(errores, fila, "Causa_Parada_No_Planificada", f"Hay {g_(Pa[i])} h de parada pero la causa dice «{SIN_PARADA}».")
        if Pa[i] == 0 and c != SIN_PARADA:
            avisos.append(f"Fila {fila}: la causa es «{c}» pero la parada es 0 h; se tomará como «{SIN_PARADA}».")
            c = SIN_PARADA
        causas.append(c)
    t = pd.DataFrame(dict(Fecha=fecha, Estacion_ID=est.astype(int), _fila=datos["_fila"]))
    dup = t.duplicated(["Fecha", "Estacion_ID"], keep="first")
    for i in t.index[dup]:
        _err(errores, int(t.at[i, "_fila"]), "Fecha / Estacion_ID",
             f"La estación {t.at[i, 'Estacion_ID']} aparece dos veces el {t.at[i, 'Fecha']:%d/%m/%Y}.")
    cuenta = t.groupby("Fecha")["Estacion_ID"].nunique()
    for f_, k in cuenta[cuenta != 3].items():
        fila = int(t.loc[t["Fecha"] == f_, "_fila"].min())
        _err(errores, fila, "Estacion_ID", f"El {f_:%d/%m/%Y} tiene {k} estación(es); cada día debe tener las 3 (1, 2 y 3).")
    if errores:
        return None, errores[:300], avisos
    out = pd.DataFrame(dict(
        Fecha=fecha.values, Estacion_ID=est.astype(int).values, Tiempo_Planificado_h=Tp.values,
        Parada_No_Planificada_h=Pa.values, Tiempo_Operativo_h=(Tp - Pa).values,
        Causa_Parada_No_Planificada=causas, Capacidad_Nominal_bidones_h=cap.values,
        Bidones_Procesados=proc.values, Bidones_Conformes=conf.values,
        Bidones_No_Conformes=(proc - conf).values))
    out["Produccion_Teorica_bidones"] = out["Capacidad_Nominal_bidones_h"] * out["Tiempo_Operativo_h"]
    return finalizar_df(out), [], avisos


@st.cache_data(show_spinner=False)
def leer_ruta_cache(ruta: str, mtime: float):
    return leer_base(ruta)


def generar_base(df_ref, ini, fin, semilla):
    """Base nueva: remuestrea días completos de la base incluida (mismo tipo de día) y
    vuelve a encadenar el flujo E1 → E2 → E3 (cada estación procesa lo que la anterior dejó conforme)."""
    rng = np.random.default_rng(int(semilla))
    dias = {f: g.sort_values("Estacion_ID").reset_index(drop=True) for f, g in df_ref.groupby("Fecha")}
    sab = [f for f in dias if pd.Timestamp(f).weekday() == 5]
    sem = [f for f in dias if pd.Timestamp(f).weekday() != 5]
    partes = []
    for d in pd.date_range(ini, fin):
        if d.weekday() == 6:
            continue
        pool = sab if d.weekday() == 5 and sab else sem
        g = dias[pool[int(rng.integers(len(pool)))]].copy()
        g["Fecha"] = d
        proc = g["Bidones_Procesados"].to_numpy(float).copy()
        nc = g["Bidones_No_Conformes"].to_numpy(float).copy()
        q = np.where(proc > 0, nc / np.where(proc > 0, proc, 1), 0.0)
        tope = (g["Capacidad_Nominal_bidones_h"] * g["Tiempo_Operativo_h"]).to_numpy(float)
        for i in (1, 2):
            proc[i] = float(max(0, min(round(proc[i - 1] - nc[i - 1]), np.floor(tope[i]))))
            nc[i] = float(round(proc[i] * q[i]))
        g["Bidones_Procesados"], g["Bidones_No_Conformes"] = proc, nc
        g["Bidones_Conformes"] = proc - nc
        partes.append(g)
    if not partes:
        return None
    out = pd.concat(partes, ignore_index=True)
    return finalizar_df(out.drop(columns=["Mes", "Est", "Causa", "Teorica"], errors="ignore"))


@st.cache_data(show_spinner=False)
def _leeme(ruta: str, mtime: float):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        hoja = next((n_ for n_ in wb.sheetnames if n_.lower().startswith("lee")), None)
        if hoja is None:
            return None, None, None
        filas = [str(r[0]) for r in wb[hoja].iter_rows(values_only=True) if r and r[0]]
    except Exception:
        return None, None, None
    titulo = filas[0] if filas else None
    caso = next((f.split(":", 1)[1].strip() for f in filas if f.startswith("Caso:")), None)
    ver = next((f.split(":", 1)[1].strip() for f in filas if f.lower().startswith("qué deberías ver")), None)
    return titulo, caso, ver


def listar_casos_prueba():
    salida = []
    for f in sorted(RUTA_BASE.parent.glob("Base_0*.xlsx")):
        titulo, caso, ver = _leeme(str(f), f.stat().st_mtime)
        salida.append(dict(ruta=f, nombre=f.name, titulo=titulo or f.stem.replace("_", " "), caso=caso, ver=ver))
    return salida


@st.cache_data(show_spinner=False)
def construir_plantilla() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "OEE_Diario"
    ws.append(COLS_PLANTILLA)
    ejemplo = [
        ("2026-01-05", 1, 6, 0.5, 0.75, "Falta de bidones vacíos (retorno de reparto)", 100, 360, 354),
        ("2026-01-05", 2, 6, 0.75, 0, SIN_PARADA, 95, 354, 350),
        ("2026-01-05", 3, 6, 0.25, 0.25, "Avería de selladora-etiquetadora", 95, 350, 347),
    ]
    for fila in ejemplo:
        ws.append([pd.Timestamp(fila[0]).to_pydatetime(), *fila[1:]])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0D1117")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for col, ancho in zip("ABCDEFGHI", (13, 12, 14, 18, 22, 44, 24, 18, 17)):
        ws.column_dimensions[col].width = ancho
    for r in range(2, 400):
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
    ins = wb.create_sheet("Instrucciones")
    ins.append(["Cómo llenar la hoja OEE_Diario"])
    ins["A1"].font = Font(bold=True, size=14)
    ins.append(["Una fila por estación y por día: cada día lleva 3 filas (estaciones 1, 2 y 3). "
                "Las 3 filas de ejemplo (05/01/2026) son solo un modelo: bórralas o reemplázalas con tus datos."])
    ins.append([])
    ins.append(["Columna", "Qué poner", "Ejemplo", "¿Obligatoria?"])
    for c in ins[4]:
        c.font = Font(bold=True)
    for fila in INSTRUCCIONES_COLS:
        ins.append(list(fila))
    ins.append([])
    ins.append(["Causas válidas (copia el texto exacto)"])
    ins.cell(row=ins.max_row, column=1).font = Font(bold=True)
    ini_causas = ins.max_row + 1
    for c in CAUSAS_VALIDAS:
        ins.append([c])
    fin_causas = ins.max_row
    ins.append([])
    ins.append(["Qué calcula la app por ti: tiempo planificado (turno − parada planificada), tiempo operativo "
                "(planificado − parada no planificada), bidones no conformes (procesados − conformes), "
                "disponibilidad, rendimiento, calidad y OEE."])
    for col, ancho in zip("ABCD", (34, 90, 30, 14)):
        ins.column_dimensions[col].width = ancho
    for row in ins.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1=f"=Instrucciones!$A${ini_causas}:$A${fin_causas}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("F2:F2000")
    dv2 = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False)
    ws.add_data_validation(dv2)
    dv2.add("B2:B2000")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def activar_datos(df, origen, archivo, detalle=""):
    """Deja una base como los datos activos y reinicia lo que depende de ella."""
    s = st.session_state
    s.datos = dict(df=df, origen=origen, archivo=archivo, detalle=detalle, ini=df["Fecha"].min(),
                   fin=df["Fecha"].max(), dias=int(df["Fecha"].nunique()))
    s.casos, s.caso_seq, s.wz = [], 0, None
    s.umb = referencias_p10_p90(df)
    s.umb_ver = s.get("umb_ver", 0) + 1
    for k in ("v_res", "v_xlsx", "dia_tipico", "errores_carga", "avisos_carga"):
        s.pop(k, None)


def periodo_txt(d):
    return f"{d['ini']:%d/%m/%Y} – {d['fin']:%d/%m/%Y}"



# =============================================================================
# ALERTAS (PARTE POST), CASOS Y VERIFICACION DE LA RESULTANTE
# =============================================================================
T7 = {  # Tabla 7 de la tesis: decisión inmediata, responsable y plazo por tipo de señal
    "oee": dict(nombre="OEE diario por estación", decision="Identificar el factor dominante: D → paradas; R → esperas o velocidad; Q → rechazos",
                resp="Jefe de Producción", plazo="≤ 24 h", roles=["Jefe de Producción"]),
    "rechazo": dict(nombre="Índice de rechazo diario", decision="Retener el lote e inspeccionar llenado, tapado o sellado según el NPR",
                    resp="Producción y Calidad", plazo="Mismo turno", roles=["Jefe de Producción"]),
    "parada": dict(nombre="Parada no planificada", decision="Priorizar con la matriz de criticidad (C3); si es crítica, correctivo inmediato",
                   resp="Mantenimiento", plazo="Inmediato o en el turno", roles=["Mantenimiento", "Operador"]),
    "recurrencia": dict(nombre="Parada repetida", decision="Priorizar con la matriz de criticidad (C3); si es crítica, correctivo inmediato",
                        resp="Mantenimiento", plazo="Inmediato o en el turno", roles=["Mantenimiento", "Operador"]),
    "insumos": dict(nombre="Parada por falta de insumos", decision="Registrar la causa, avisar a Logística y resecuenciar la producción",
                    resp="Logística (fuera del alcance)", plazo="Mismo día", roles=["Jefe de Producción"]),
    "sensor": dict(nombre="Variable de sensor fuera de rango (C1)", decision="Abrir el evento y diagnosticar por NPR (C2)",
                   resp="Operador avisa; Mantenimiento según prioridad C3", plazo="Según la prioridad (C3)",
                   roles=["Operador", "Mantenimiento"]),
}
TIPO_EMOJI = {"rojo": "🔴", "naranja": "🟠", "amarillo": "🟡"}

# Relación alerta → equipos y variables del AMEF (supuesto de demostración, "(ejemplo)")
MAPA_AMEF = {
    ("rechazo", 1): [("Tanque", "nivel", "baja"), ("Bomba", "presion", "baja"), ("Bomba", "caudal", "baja")],
    ("rechazo", 2): [("Llenadora", "rech_llen", "alta"), ("Tapadora", "torque", "baja")],
    ("rechazo", 3): [("Selladora", "rech_sello", "alta"), ("Selladora", "cal_sello", "NG"), ("Etiquetadora", "cal_etiq", "NG")],
    ("equipo", 1): [("Bomba", "vibracion", "alta"), ("Bomba", "corriente", "alta")],
    ("equipo", 2): [("Llenadora", "ciclo", "alta"), ("Tapadora", "vib_tapado", "alta")],
    ("equipo", 3): [("Etiquetadora", "estado", "STOP"), ("Selladora", "rech_sello", "alta")],
    ("velocidad", 1): [("Bomba", "caudal", "baja")],
    ("velocidad", 2): [("Llenadora", "conteo", "baja"), ("Llenadora", "ciclo", "alta")],
    ("velocidad", 3): [("Etiquetadora", "estado", "STOP")],
}


def leer_par():
    s = st.session_state
    return dict(parada_min=float(s.get("w_parada_min", 30)), rec_n=int(s.get("w_rec_n", 2)),
                rec_ventana=int(s.get("w_rec_vent", 1)))


def factores_base(df):
    g = tabla_estaciones(df)
    return {e: (float(g.loc[e, "D"]), float(g.loc[e, "R"]), float(g.loc[e, "Q"])) for e in g.index}


def alertas_en(df_rows, df_all, umb, par):
    """Alertas de la parte post (Tabla 7) para las filas estación-día de df_rows, ordenadas por gravedad."""
    fb = factores_base(df_all)
    alertas = []
    for r in df_rows.itertuples(index=False):
        e, fecha = int(r.Estacion_ID), r.Fecha
        u = umb[e]
        tp, top, cap = r.Tiempo_Planificado_h, r.Tiempo_Operativo_h, r.Capacidad_Nominal_bidones_h
        proc, conf, nc = r.Bidones_Procesados, r.Bidones_Conformes, r.Bidones_No_Conformes
        pa, causa = float(r.Parada_No_Planificada_h), r.Causa
        nom = EST_NOM[e]
        base = dict(fecha=fecha, est=e, causa_parada=causa)

        def nueva(tipo, senal, valor, umbral, vnum, unum, gravedad, color, cierre, **extra):
            t7 = T7[tipo]
            return dict(base, id=f"{fecha:%Y-%m-%d}|E{e}|{tipo}", tipo=tipo, senal=f"{senal} · {nom}", valor=valor,
                        umbral=umbral, valor_num=vnum, umbral_num=unum, gravedad=gravedad, color=color,
                        decision=t7["decision"], resp=t7["resp"], plazo=t7["plazo"], roles=t7["roles"],
                        cierre=cierre, **extra)
        oee = conf / (cap * tp) if cap * tp > 0 and top > 0 else np.nan
        if np.isfinite(oee) and oee < u["oee_p10"]:
            d_, r_, q_ = (top / tp if tp else np.nan), (proc / (cap * top) if cap * top > 0 else np.nan), (conf / proc if proc > 0 else np.nan)
            rel = {"D": d_ / fb[e][0], "R": r_ / fb[e][1], "Q": q_ / fb[e][2]}
            rel = {k: v for k, v in rel.items() if np.isfinite(v)}
            dom = min(rel, key=rel.get) if rel else "D"
            grav = (u["oee_p10"] - oee) / max(u["oee_p10"], 1e-9)
            alertas.append(nueva("oee", T7["oee"]["nombre"], f"{n(oee * 100)} %", f"< P10: {n(u['oee_p10'] * 100)} %", oee,
                                 u["oee_p10"], grav, "rojo" if grav >= 0.15 else "naranja",
                                 f"OEE ≥ mediana base en la semana siguiente ({n(u['oee_med'] * 100)} %)", factor=dom))
        rech = nc / proc if proc > 0 else np.nan
        if np.isfinite(rech) and rech > u["rech_p90"]:
            grav = (rech - u["rech_p90"]) / max(u["rech_p90"], 1e-9)
            alertas.append(nueva("rechazo", T7["rechazo"]["nombre"], f"{n(rech * 100)} %", f"> P90: {n(u['rech_p90'] * 100)} %", rech,
                                 u["rech_p90"], grav, "rojo" if grav >= 0.3 else "naranja",
                                 f"Rechazo ≤ promedio base al día siguiente ({n(u['rech_med'] * 100)} %)"))
        if pa > 0 and causa in CAUSAS_EQUIPO:
            if pa * 60 >= par["parada_min"]:
                alertas.append(nueva("parada", T7["parada"]["nombre"], f"{n(pa)} h · {causa}", f"≥ {n(par['parada_min'], 0)} min", pa,
                                     par["parada_min"] / 60, pa, "rojo" if pa >= 1.0 else "naranja",
                                     "MTBF de la causa mayor que la línea base en el mes siguiente"))
            ventana = pd.Timedelta(days=max(par["rec_ventana"], 1) - 1)
            previos = df_all[(df_all["Estacion_ID"] == e) & (df_all["Causa"] == causa) & (df_all["Fecha"] <= fecha)
                             & (df_all["Fecha"] >= fecha - ventana) & (df_all["Parada_No_Planificada_h"] > 0)]
            if len(previos) >= par["rec_n"]:
                alertas.append(nueva("recurrencia", T7["recurrencia"]["nombre"], f"{len(previos)} eventos · {causa}",
                                     f"≥ {par['rec_n']} eventos en {par['rec_ventana']} día(s)", pa, par["rec_n"], len(previos) / par["rec_n"],
                                     "rojo", "MTBF de la causa mayor que la línea base en el mes siguiente"))
        elif pa > 0 and causa in CAUSAS_INSUMO:
            alertas.append(nueva("insumos", T7["insumos"]["nombre"], f"{n(pa)} h · {causa}", "Cualquier parada por insumos", pa, 0,
                                 0.0, "amarillo", "Sin paradas por falta de insumos en la semana siguiente"))
    orden = {"rojo": 0, "naranja": 1, "amarillo": 2}
    alertas.sort(key=lambda a: (orden[a["color"]], -a["gravedad"], a["fecha"], a["est"]))
    return alertas


def dia_tipico(df, umb, par):
    """Día típico: sin cortes de servicio, con 1-2 alertas y OEE de línea cercano a la mediana;
    si no hay, el día del medio del periodo."""
    dias = sorted(df["Fecha"].unique())
    oee_d = df["Bidones_Conformes"] / (df["Capacidad_Nominal_bidones_h"] * df["Tiempo_Planificado_h"])
    oee_linea = oee_d.groupby(df["Fecha"]).mean()
    mediana = float(oee_linea.median())
    con_corte = set(df.loc[df["Causa"].isin(["Corte de energía eléctrica", "Corte de suministro de agua"]), "Fecha"])
    cand = [d for d in dias if d not in con_corte and 1 <= len(alertas_en(df[df["Fecha"] == d], df, umb, par)) <= 2]
    if cand:
        return min(cand, key=lambda d: abs(float(oee_linea[d]) - mediana))
    return dias[len(dias) // 2]


def causas_para_alerta(al, ajustes, descartadas):
    """Causas candidatas del AMEF para una alerta, ordenadas por NPR (las descartadas al final)."""
    e, tipo = al["est"], al["tipo"]
    if tipo == "insumos":
        return []
    if tipo == "sensor":
        triples = [tuple(al["sensor_ref"])]
    else:
        if tipo in ("parada", "recurrencia"):
            g = ("equipo", e)
        elif tipo == "rechazo":
            g = ("rechazo", e)
        else:
            g = ({"D": "equipo", "R": "velocidad", "Q": "rechazo"}[al.get("factor", "D")], e)
        triples = MAPA_AMEF.get(g, [])
    salida = []
    if True:
        for eq, vkey, direc in triples:
            v = next(x for x in VARIABLES[eq] if x["key"] == vkey)
            for c in causas_vigentes(eq, vkey, direc, ajustes):
                c.update(equipo=eq, vkey=vkey, vnombre=v["nombre"], direccion=direc,
                         ident=(eq, vkey, direc, c["causa"]))
                c["descartada"] = c["ident"] in descartadas
                salida.append(c)
    salida.sort(key=lambda c: (c["descartada"], -c["npr"]))
    return salida


def evaluar_causa(c):
    """Prioridad (matriz configurada) y tipo TPM de una causa del AMEF."""
    return evaluar_c3(c)


ART_EQ = {"Tanque": "del tanque", "Bomba": "de la bomba", "Llenadora": "de la llenadora", "Tapadora": "de la tapadora",
          "Selladora": "de la selladora", "Etiquetadora": "de la etiquetadora"}


def plazo_prio(prio, tpm):
    """Plazo de atención según la prioridad de C3 (y el tipo de mantenimiento TPM)."""
    if prio == "Media":
        quien = "Mantenimiento autónomo del operador" if tpm == "Autónomo" else "Mantenimiento planificado"
        return f"{quien} en un máximo de 48 h o en la parada preventiva del sábado"
    if prio == "Baja" and tpm != "Autónomo":
        return "Mantenimiento de mejora en la siguiente limpieza"
    return PLAZO_PRIORIDAD[prio]


def decision_info(al, c, c3):
    """Acción concreta y plazo coherentes con la prioridad C3 (la Tabla 7 solo da «inmediato» a las Críticas)."""
    if c3 is None or c is None:                              # falta de insumos: sin prioridad de mantenimiento
        return dict(accion=al["decision"], plazo=al["plazo"], plazo_ref="", prio_txt="No aplica (fuera del alcance)")
    prio, tpm = c3["prioridad"], c3["tipo"]
    est = EST_CORTO[al["est"]]
    eq = ART_EQ.get(c.get("equipo"), "del equipo")
    causa = c["causa"][:1].lower() + c["causa"][1:]
    plazo_c3 = plazo_prio(prio, tpm)
    prio_txt = f"{prio} · mantenimiento {tpm.lower()}"
    if al["tipo"] in ("parada", "recurrencia", "sensor"):
        if prio == "Crítica":
            accion = f"Detener {est} y hacer el mantenimiento correctivo inmediato {eq}: revisar {causa}"
        elif tpm == "Autónomo":
            accion = f"Pedir al operador de {est} el mantenimiento autónomo {eq}: revisar {causa}"
        elif tpm == "De mejora":
            accion = f"Incluir en el plan de mejora {eq} ({est}) la revisión de {causa}"
        elif prio == "Alta":
            accion = f"Programar el mantenimiento planificado {eq} ({est}) dentro del mismo turno para revisar {causa}"
        else:
            accion = f"Programar mantenimiento planificado {eq} ({est}) para revisar {causa}"
        return dict(accion=accion, plazo=plazo_c3, plazo_ref="(Tabla 7: correctivo inmediato solo si la prioridad es Crítica)",
                    prio_txt=prio_txt)
    # rechazo / OEE: se conserva la decisión de la Tabla 7; el plazo del mantenimiento sale de la prioridad
    return dict(accion=al["decision"], plazo=al["plazo"], plazo_ref="", prio_txt=f"{prio_txt} · plazo del mantenimiento: {plazo_c3}")


def nuevo_caso(al, causa, c3, origen="datos"):
    s = st.session_state
    s.caso_seq += 1
    info = decision_info(al, causa, c3)
    return dict(
        id=f"C-{s.caso_seq:03d}", origen=origen, alerta={k: v for k, v in al.items()}, causa=dict(causa),
        prioridad=c3["prioridad"] if c3 else "No aplica", tipo_tpm=c3["tipo"] if c3 else "Fuera del alcance (Logística)",
        regla=c3["regla"] if c3 else "", responsable=al["resp"], plazo=info["plazo"], plazo_ref=info["plazo_ref"],
        accion=info["accion"], estado="Abierto",
        descartadas=[], intervenciones=[], retro=[], creado=pd.Timestamp.now(), resultado=None)


def _etq_despues(df, w, fecha, dias, nombre):
    """Etiqueta del «después»: usa «semana/mes siguiente» solo si los datos cubren ese periodo completo."""
    completo = df["Fecha"].max() >= pd.Timestamp(fecha) + pd.Timedelta(days=dias)
    return f"Después ({nombre})" if completo else f"Después ({len(w)} días laborables siguientes)"


def _ventana(df, est, fecha, dias, primero=False):
    d = df[(df["Estacion_ID"] == est) & (df["Fecha"] > fecha)].sort_values("Fecha")
    if primero:
        return d.head(1)
    return d[d["Fecha"] <= fecha + pd.Timedelta(days=dias)]


def verificar_caso(caso, df, umb):
    """Compara con los datos del periodo SIGUIENTE según la Tabla 7. Devuelve un dict con el resultado."""
    al = caso["alerta"]
    e, fecha, tipo = al["est"], pd.Timestamp(al["fecha"]), al["tipo"]
    nom = EST_NOM[e]
    sin = dict(sin_datos=True, ok=None, mensaje="No hay datos posteriores para verificar; carga un periodo más largo.")
    if tipo == "rechazo":
        w = _ventana(df, e, fecha, 1, primero=True)
        if w.empty:
            return sin
        r = w.iloc[0]
        obs = r.Bidones_No_Conformes / r.Bidones_Procesados if r.Bidones_Procesados > 0 else np.nan
        esp = umb[e]["rech_med"]
        return dict(sin_datos=False, ok=bool(obs <= esp), tipo_u="pct", antes=al["valor_num"], despues=obs, esperado=esp, cond="≤",
                    etiqueta_antes="Antes (día de la alerta)", etiqueta_despues=f"Después ({r.Fecha:%d/%m/%Y})",
                    esperado_txt=f"Rechazo de {nom} ≤ promedio base ({n(esp * 100)} %) al día siguiente",
                    observado_txt=f"{n(obs * 100)} % el {r.Fecha:%d/%m/%Y}", ventana=f"1 día ({r.Fecha:%d/%m/%Y})", n_obs=1)
    if tipo == "oee":
        w = _ventana(df, e, fecha, 7)
        if w.empty:
            return sin
        o = w["Bidones_Conformes"] / (w["Capacidad_Nominal_bidones_h"] * w["Tiempo_Planificado_h"])
        obs, esp = float(o.mean()), umb[e]["oee_med"]
        return dict(sin_datos=False, ok=bool(obs >= esp), tipo_u="pct", antes=al["valor_num"], despues=obs, esperado=esp, cond="≥",
                    etiqueta_antes="Antes (día de la alerta)", etiqueta_despues=_etq_despues(df, w, fecha, 7, "semana siguiente"),
                    esperado_txt=f"OEE de {nom} ≥ mediana base ({n(esp * 100)} %) en la semana siguiente",
                    observado_txt=f"{n(obs * 100)} % (promedio de {len(w)} días)", ventana=f"{len(w)} día(s) tras la alerta", n_obs=len(w),
                    serie=o.tolist())
    if tipo in ("parada", "recurrencia"):
        w = _ventana(df, e, fecha, 30)
        if w.empty:
            return sin
        causa = al["causa_parada"]
        todo = df[df["Estacion_ID"] == e]
        ev_base = int(((todo["Causa"] == causa) & (todo["Parada_No_Planificada_h"] > 0)).sum())
        mtbf_base = todo["Tiempo_Operativo_h"].sum() / ev_base if ev_base else float("inf")
        ev_post = int(((w["Causa"] == causa) & (w["Parada_No_Planificada_h"] > 0)).sum())
        mtbf_post = w["Tiempo_Operativo_h"].sum() / ev_post if ev_post else float(w["Tiempo_Operativo_h"].sum())
        return dict(sin_datos=False, ok=bool(mtbf_post > mtbf_base), tipo_u="h", antes=mtbf_base, despues=mtbf_post, esperado=mtbf_base,
                    cond=">", etiqueta_antes="Línea base (todo el periodo)", etiqueta_despues=_etq_despues(df, w, fecha, 30, "mes siguiente"),
                    esperado_txt=f"MTBF de «{causa}» en {nom} mayor que la línea base ({n(mtbf_base)} h)",
                    observado_txt=f"{n(mtbf_post)} h ({ev_post} evento(s) en {len(w)} días)", ventana=f"{len(w)} día(s) tras la alerta",
                    n_obs=len(w))
    if tipo == "insumos":
        w = _ventana(df, e, fecha, 7)
        if w.empty:
            return sin
        k = int(((w["Causa"].isin(CAUSAS_INSUMO)) & (w["Parada_No_Planificada_h"] > 0)).sum())
        return dict(sin_datos=False, ok=bool(k == 0), tipo_u="int", antes=1, despues=k, esperado=0, cond="=",
                    etiqueta_antes="Antes (alerta)", etiqueta_despues=_etq_despues(df, w, fecha, 7, "semana siguiente"),
                    esperado_txt=f"Sin paradas por falta de insumos en {nom} en la semana siguiente",
                    observado_txt=f"{k} parada(s) por insumos en {len(w)} días", ventana=f"{len(w)} día(s) tras la alerta", n_obs=len(w))
    return sin


def fmt_u(tipo_u, x):
    if x is None or (isinstance(x, float) and not np.isfinite(x)):
        return "—"
    return {"pct": f"{n(x * 100)} %", "h": f"{n(x)} h", "int": n(x, 0)}.get(tipo_u, n(x))


def retro_kpi(caso, res, umb_store):
    """Retroalimentación al cerrar un caso de KPI: banda de C1, O/D de C2 y prioridad de C3."""
    s = st.session_state
    lineas = []
    al, c = caso["alerta"], caso["causa"]
    e = al["est"]
    if al["tipo"] in ("rechazo", "oee") and res.get("n_obs", 0) >= 3 and len(res.get("serie", [])) >= 3:
        x = np.array(res["serie"], float)
        if al["tipo"] == "oee":
            nuevo = max(umb_store[e]["oee_p10"], float(x.mean() - 3 * x.std(ddof=1)))
            if nuevo > umb_store[e]["oee_p10"] + 1e-6:
                lineas.append(f"C1: umbral de OEE bajo (P10) de {EST_CORTO[e]}: {n(umb_store[e]['oee_p10'] * 100)} % → {n(nuevo * 100)} %.")
                umb_store[e]["oee_p10"] = nuevo
                s.umb_ver += 1
            else:
                lineas.append("C1: umbral sin cambio (la banda calculada no es más estricta).")
    else:
        lineas.append("C1: umbral sin cambio (hay muy pocos datos posteriores para calcular una banda).")
    ident = c.get("ident")
    if ident:
        aj = s.d_ajustes.get(ident, dict(O=c["O"], D=c["D"]))
        o0, d0 = aj["O"], aj["D"]
        o1, d1 = min(10, o0 + 1), max(1, d0 - 1)
        s.d_ajustes[ident] = dict(O=o1, D=d1)
        c2 = dict(c, O=o1, D=d1, npr=c["S"] * o1 * d1)
        p_a, p_d = evaluar_c3(dict(c, O=o0, D=d0))["prioridad"], evaluar_c3(c2)["prioridad"]
        lineas.append(f"C2: causa confirmada · O {o0}→{o1}; D {d0}→{d1} · NPR {c['S'] * o0 * d0}→{c2['npr']}.")
        lineas.append(f"C3: prioridad {p_a} → {p_d}.")
    return lineas


def verificar_sensor(caso, efecto):
    """Verifica un caso creado desde sensores con la lógica C4 (12 lecturas posteriores; selector de prueba)."""
    s = st.session_state
    se = caso["sensor"]
    eq, key, direccion = se["equipo"], se["key"], se["direccion"]
    v = next(x for x in VARIABLES[eq] if x["key"] == key)
    datos = s.d_datos
    n0 = len(datos[eq])
    res_prev = analizar(datos, s.d_bandas)
    sost_eq = res_prev[eq]["sost"]
    primera = int(np.argmax(sost_eq[key].values))
    inicio = s.d_fallas.get(eq, {}).get("inicio")
    retraso = None if inicio is None else primera - inicio
    a_tiempo = retraso is not None and retraso <= DET_OPORTUNA_LECTURAS
    anomalas = {k for k in sost_eq.columns if sost_eq[k].iloc[n0 - 1]}
    banda_prev = s.d_bandas.get((eq, key))
    s.d_seed += 1
    nuevos = extender_datos(datos, s.d_fallas_act, eq, key, efecto, np.random.default_rng(s.d_seed))
    post = nuevos[eq][key].iloc[n0:]
    resultado, frac = verificar_resolucion(v, post, banda_prev)
    s.d_datos = nuevos
    s.w_d_lect = len(nuevos[eq])
    c = caso["causa"]
    lineas = []
    ok = resultado == "Sí"
    if ok:
        aj = s.d_ajustes.get(c["ident"], dict(O=c["O"], D=c["D"]))
        o0, d0 = aj["O"], aj["D"]
        o1, d1 = min(10, o0 + 1), (max(1, d0 - 1) if a_tiempo else d0)
        s.d_ajustes[c["ident"]] = dict(O=o1, D=d1)
        banda = calcular_banda(v, post)
        if banda:
            s.d_bandas[(eq, key)] = banda
        f = s.d_fallas_act.get(eq)
        if f:
            resto = [(k, t) for k, t in f["vars"] if k in anomalas and k != key]
            if resto:
                s.d_fallas_act[eq] = dict(f, vars=resto)
            else:
                del s.d_fallas_act[eq]
        lineas.append(f"C1: banda de alerta temprana ⚑ {formato_umbral(v, banda)} (promedio posterior ± {K_SIGMA_BANDA:g}σ; ejemplo)."
                      if banda else "C1: umbral sin cambio (la variable no admite banda).")
        dtxt = ("C1 detectó a tiempo" if a_tiempo else "detección tardía: D no se reduce")
        lineas.append(f"C2: causa confirmada · O {o0}→{o1}; D {d0}→{d1} ({dtxt}) · NPR {c['S'] * o0 * d0}→{c['S'] * o1 * d1}.")
        p_a = evaluar_c3(dict(c, O=o0, D=d0))["prioridad"]
        p_d = evaluar_c3(dict(c, O=o1, D=d1))["prioridad"]
        lineas.append(f"C3: prioridad {p_a} → {p_d}.")
    return dict(sin_datos=False, ok=ok, tipo_u="num", antes=float(se["valor_obs"]) if not isinstance(se["valor_obs"], str) else np.nan,
                despues=float(post.astype(float).mean()) if v["kind"] not in ("okng", "runstop") else np.nan,
                esperado=None, cond="en rango", etiqueta_antes="Antes (al alertar)", etiqueta_despues="Después (12 lecturas)",
                esperado_txt="Variable en rango en las 3 últimas de 12 lecturas",
                observado_txt=f"{valor_despues(v, post)} · {frac * 100:.0f} % de las 12 lecturas en rango · resultado «{resultado}»",
                ventana="12 lecturas simuladas", n_obs=12, lineas=lineas, parcial=(resultado == "Parcial"),
                rango=(limites(v, banda_prev) if v["kind"] in ("range", "max") else None), unidad=v["unidad"])



# =============================================================================
# ESTILOS Y COMPONENTES DE INTERFAZ
# =============================================================================
CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
html, body, .stApp { font-family: 'Inter', 'Segoe UI', system-ui, sans-serif; }
[data-testid="stIconMaterial"], span[class*="material-symbols"] { font-family: 'Material Symbols Rounded' !important; }
.stApp { background: #ffffff; color: #0d1117; }
header[data-testid="stHeader"] { background: transparent; }
.block-container { max-width: 1180px; padding-top: 2.2rem; padding-bottom: 3rem; }
section[data-testid="stSidebar"] { background: #fafafa; border-right: 1px solid #e5e7eb; }
section[data-testid="stSidebar"] .block-container { padding-top: 1.2rem; }
section[data-testid="stSidebar"] button { min-height: 2.5rem !important; justify-content: flex-start; text-align: left; font-size: .88rem; }
.hdr { display:flex; align-items:center; justify-content:space-between; padding-bottom:14px;
       border-bottom:1px solid #e5e7eb; margin-bottom:20px; gap:12px; flex-wrap:wrap; }
.brand { display:flex; align-items:center; gap:12px; }
.logo { width:40px; height:40px; border-radius:10px; background:#0d1117; display:flex; align-items:center; justify-content:center; }
.bn { font-weight:700; font-size:1.02rem; line-height:1.1; }
.bs { font-size:.66rem; letter-spacing:.14em; color:#6b7280; font-weight:500; }
.pill { font-size:.78rem; color:#374151; border:1px solid #e5e7eb; border-radius:999px; padding:5px 12px; background:#fff;
        display:flex; align-items:center; gap:7px; }
.pill .dot { width:7px; height:7px; border-radius:50%; background:#12a150; display:inline-block; }
.pill .dot.off { background:#d1d5db; }
.eyebrow { font-size:.7rem; letter-spacing:.16em; text-transform:uppercase; color:#6b7280; font-weight:600; }
.h1 { font-size:2.05rem; font-weight:700; margin:.25rem 0 .45rem; letter-spacing:-.01em; }
.h2 { font-size:1.6rem; font-weight:700; margin:.1rem 0 .5rem; letter-spacing:-.01em; }
.lead { color:#4b5563; line-height:1.6; max-width:860px; font-size:.97rem; }
.ayuda { border:1px solid #bfdbfe; background:#eff6ff; color:#1e3a8a; border-radius:12px; padding:12px 16px; margin:6px 0 18px;
         font-size:.92rem; line-height:1.55; }
.ayuda .ayuda-t { display:block; font-size:.74rem; font-weight:700; letter-spacing:.12em; text-transform:uppercase; margin-bottom:2px; }
.kpi { border:1px solid #e5e7eb; border-radius:12px; padding:14px 16px; background:#fff; box-shadow:0 1px 2px rgba(16,24,40,.04); min-height:96px; }
.kpi .lab { font-size:.66rem; letter-spacing:.13em; text-transform:uppercase; color:#6b7280; font-weight:600; }
.kpi .val { font-size:1.55rem; font-weight:700; margin-top:5px; line-height:1.15; }
.kpi .sub { font-size:.74rem; color:#6b7280; margin-top:3px; }
.kpi.ok { border-color:#abefc6; background:#f6fef9; } .kpi.ok .val { color:#067647; }
.kpi.bad { border-color:#fecdca; background:#fffbfa; } .kpi.bad .val { color:#b42318; }
.card-t { font-size:1.2rem; font-weight:700; margin:.2rem 0 .5rem; }
.card-p { color:#4b5563; font-size:.92rem; line-height:1.55; }
.mini-t { font-weight:700; font-size:1rem; margin:.15rem 0 .35rem; }
.mini-p { color:#4b5563; font-size:.86rem; line-height:1.5; }
.paso-card { border:1px solid #e5e7eb; border-radius:12px; padding:12px 14px; background:#fff; min-height:112px; }
.paso-card .num { display:inline-flex; width:26px; height:26px; border-radius:50%; background:#0d1117; color:#fff; align-items:center;
                  justify-content:center; font-size:.8rem; font-weight:700; margin-bottom:6px; }
.paso-card .t { font-weight:700; font-size:.92rem; } .paso-card .p { color:#6b7280; font-size:.8rem; line-height:1.4; margin-top:2px; }
.box { border-radius:10px; padding:11px 14px; font-size:.9rem; margin:8px 0; line-height:1.5; }
.box.ok { background:#ecfdf3; color:#067647; border:1px solid #abefc6; }
.box.err { background:#fef3f2; color:#b42318; border:1px solid #fecdca; }
.box.info { background:#eff6ff; color:#1d4ed8; border:1px solid #bfdbfe; }
.box.warn { background:#fffaeb; color:#b54708; border:1px solid #fedf89; }
.prog-row { display:grid; grid-template-columns: 190px 120px 70px 1fr; gap:14px; align-items:center; padding:9px 0; font-size:.92rem; }
.badge { font-size:.74rem; border-radius:999px; padding:3px 10px; text-align:center; font-weight:600; }
.badge.c { background:#ecfdf3; color:#067647; border:1px solid #abefc6; }
.badge.e { background:#eff6ff; color:#1d4ed8; border:1px solid #bfdbfe; }
.badge.p { background:#f3f4f6; color:#6b7280; border:1px solid #e5e7eb; }
.bar { height:8px; border-radius:6px; background:#eef0f3; overflow:hidden; }
.bar > div { height:100%; background:#2f80ed; border-radius:6px; }
.sem { border:1px solid #e5e7eb; border-radius:12px; padding:12px 14px; background:#fff; }
.sem .st { display:flex; align-items:center; gap:8px; font-weight:700; margin-bottom:6px; }
.sem .eq { display:flex; align-items:center; gap:8px; font-size:.86rem; padding:3px 0; color:#374151; }
.dotc { width:10px; height:10px; border-radius:50%; display:inline-block; flex:none; }
.dotc.g { background:#12a150; } .dotc.r { background:#d92d20; } .dotc.a { background:#e0a100; } .dotc.n { background:#f08c00; }
.tbl { width:100%; border-collapse:collapse; font-size:.86rem; }
.tbl th { text-align:left; font-size:.68rem; letter-spacing:.1em; text-transform:uppercase; color:#6b7280; padding:8px 10px; border-bottom:1px solid #e5e7eb; }
.tbl td { padding:9px 10px; border-bottom:1px solid #f1f2f4; vertical-align:top; }
.npr { display:inline-block; min-width:44px; text-align:center; border-radius:8px; padding:3px 8px; font-weight:700; }
.prio { display:inline-block; font-size:1.5rem; font-weight:700; padding:14px 22px; border-radius:12px; }
.alerta { border:1px solid #e5e7eb; border-left:5px solid #d92d20; border-radius:10px; padding:12px 16px; margin:8px 0; background:#fff; }
.alerta .t { font-weight:700; margin-bottom:4px; } .alerta .m { font-size:.88rem; color:#374151; line-height:1.55; }
.decision { border:2px solid #0d1117; border-radius:14px; padding:18px 20px; background:#fff; margin:10px 0; }
.decision .t { font-size:1.15rem; font-weight:700; margin-bottom:8px; }
.decision .m { font-size:.95rem; line-height:1.7; }
.ev { border:1px solid #e5e7eb; border-radius:14px; padding:18px 20px; background:#fff; min-height:120px; }
.ev .lab { font-size:.7rem; letter-spacing:.14em; text-transform:uppercase; color:#6b7280; font-weight:600; }
.ev .big { font-size:1.35rem; font-weight:700; margin-top:6px; line-height:1.3; }
.resultado { font-size:2rem; font-weight:800; text-align:center; border-radius:14px; padding:18px; margin:12px 0; }
.resultado.ok { background:#ecfdf3; color:#067647; border:2px solid #abefc6; }
.resultado.bad { background:#fef3f2; color:#b42318; border:2px solid #fecdca; }
.chip { display:inline-block; font-size:.72rem; font-weight:700; border-radius:999px; padding:3px 10px; letter-spacing:.04em; }
.chip.emp { background:#eff6ff; color:#1d4ed8; border:1px solid #bfdbfe; }
.chip.sin { background:#f5f3ff; color:#6d28d9; border:1px solid #ddd6fe; }
.activos { font-size:.82rem; line-height:1.55; border:1px solid #e5e7eb; border-radius:12px; padding:10px 12px; background:#fff; }
.foot { text-align:center; color:#9ca3af; font-size:.76rem; padding-top:26px; }
button[data-testid="stBaseButton-primary"], button[kind="primary"] {
    background:#0d1117 !important; color:#fff !important; border:1px solid #0d1117 !important; border-radius:10px !important; min-height:3rem; font-weight:600; }
button[data-testid="stBaseButton-primary"]:hover, button[kind="primary"]:hover { background:#1f2937 !important; }
button[data-testid="stBaseButton-primary"]:disabled, button[kind="primary"]:disabled {
    background:#e5e7eb !important; color:#9ca3af !important; border-color:#e5e7eb !important; }
button[data-testid="stBaseButton-secondary"], button[kind="secondary"] {
    background:#fff !important; color:#0d1117 !important; border:1px solid #d1d5db !important; border-radius:10px !important; min-height:3rem; font-weight:500; }
div[data-testid="stVerticalBlockBorderWrapper"] { border-radius:14px; border-color:#e5e7eb; box-shadow:0 1px 3px rgba(16,24,40,.05); }
div[data-testid="stTabs"] button[role="tab"] { font-weight:500; }
/* cargadores de archivos: textos en español (el ícono se conserva) y alto parejo con los botones */
[data-testid="stFileUploaderDropzone"] { min-height:3rem; padding:.35rem .75rem; align-items:center; box-sizing:border-box; gap:.6rem; }
[data-testid="stFileUploaderDropzone"] button[data-testid="stBaseButton-secondary"] { min-height:2.2rem !important; }
[data-testid="stFileUploaderDropzone"] button [data-testid="stMarkdownContainer"] p { font-size:0 !important; line-height:0; margin:0; }
[data-testid="stFileUploaderDropzone"] button [data-testid="stMarkdownContainer"] p::after { content:"Cargar archivo"; font-size:.875rem; line-height:1.4; }
[data-testid="stFileUploaderDropzone"]:has(input[accept*=".csv"]) button [data-testid="stMarkdownContainer"] p::after { content:"Cargar CSV"; }
[data-testid="stFileUploaderDropzone"]:has(input[accept*=".xlsx"]) button [data-testid="stMarkdownContainer"] p::after { content:"Cargar Excel"; }
[data-testid="stFileUploaderDropzone"] { flex-direction:row !important; flex-wrap:nowrap !important; overflow:hidden; }
[data-testid="stFileUploaderDropzone"] > span { flex:none; }
[data-testid="stFileUploaderDropzone"] button { flex:none; }
[data-testid="stFileUploaderDropzoneInstructions"], [data-testid="stFileUploaderDropzoneInstructions"] > div { min-width:0; overflow:hidden; }
[data-testid="stFileUploaderDropzoneInstructions"] span { font-size:0 !important; display:block; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }
[data-testid="stFileUploaderDropzoneInstructions"] span::after { content:"Archivo (máx. 200 MB)"; font-size:.75rem; }
[data-testid="stFileUploaderDropzone"]:has(input[accept*=".csv"]) [data-testid="stFileUploaderDropzoneInstructions"] span::after { content:"Archivo .csv (máx. 200 MB)"; }
[data-testid="stFileUploaderDropzone"]:has(input[accept*=".xlsx"]) [data-testid="stFileUploaderDropzoneInstructions"] span::after { content:"Archivo .xlsx (máx. 200 MB)"; }
</style>
"""

LOGO = ('<svg width="22" height="22" viewBox="0 0 24 24" fill="none" stroke="#fff" stroke-width="1.8" '
        'stroke-linecap="round" stroke-linejoin="round"><path d="M12 3.2c3.2 4 6 6.9 6 10.4a6 6 0 0 1-12 0c0-3.5 2.8-6.4 6-10.4z"/>'
        '<path d="M9.2 14.2a2.9 2.9 0 0 0 2.4 2.4"/></svg>')


def envolver(texto, ancho=14):
    """Parte un texto largo en varias líneas (<br>) para las etiquetas de los ejes."""
    return "<br>".join(textwrap.wrap(str(texto), ancho)) or str(texto)


def estilo_fig(fig, h=380, titulo=None):
    fig.update_layout(
        template="plotly_white", height=h, paper_bgcolor="#fff", plot_bgcolor="#fff",
        margin=dict(l=10, r=10, t=88 if titulo else 44, b=10),
        font=dict(family="Segoe UI, Arial, sans-serif", size=12, color=INK),
        legend=dict(orientation="h", y=1.02, yanchor="bottom", x=0, xanchor="left"),
        title=dict(text=titulo, x=0, y=0.97, yanchor="top", font=dict(size=14)) if titulo else None)
    fig.update_xaxes(automargin=True, tickangle=0, tickfont=dict(size=10.5))
    fig.update_yaxes(automargin=True)
    fig.update_traces(cliponaxis=False, selector=dict(type="bar"))
    fig.update_traces(cliponaxis=False, selector=dict(type="scatter"))
    return fig


def caja(texto, tipo="info"):
    st.markdown(f'<div class="box {tipo}">{texto}</div>', unsafe_allow_html=True)


def caja_ayuda(texto):
    st.markdown(f'<div class="ayuda"><span class="ayuda-t">¿Qué hago aquí?</span>{texto}</div>', unsafe_allow_html=True)


def kpi(col, etiqueta, valor, sub="", estado=None):
    sub_html = f'<div class="sub">{sub}</div>' if sub else ""
    clase = f" {estado}" if estado in ("ok", "bad") else ""
    col.markdown(f'<div class="kpi{clase}"><div class="lab">{etiqueta}</div><div class="val">{valor}</div>'
                 f'{sub_html}</div>', unsafe_allow_html=True)


def tabla_html(cabeceras, filas):
    th = "".join(f"<th>{c}</th>" for c in cabeceras)
    tr = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in f) + "</tr>" for f in filas)
    st.markdown(f'<table class="tbl"><thead><tr>{th}</tr></thead><tbody>{tr}</tbody></table>', unsafe_allow_html=True)


# =============================================================================
# NAVEGACION: PASOS, BLOQUEOS, MENU LATERAL
# =============================================================================
PASOS = [
    ("inicio", "Inicio"),
    ("kpis", "1. KPIs de entrada"),
    ("datos", "2. Datos"),
    ("params", "3. Parámetros"),
    ("tablero", "4. Tablero (tendencias y OEE)"),
    ("decision", "5. Decisión (parte post)"),
    ("ejecucion", "6. Ejecución y resultante"),
    ("valid", "7. Validación (Monte Carlo)"),
    ("export", "8. Exportación"),
]
ORDEN = [c for c, _ in PASOS]
NOMBRE = dict(PASOS)


def nombre_corto(clave):
    return NOMBRE[clave].split(". ", 1)[-1]


def hay_datos():
    return st.session_state.get("datos") is not None


def df_activo():
    return st.session_state.datos["df"]


def completo(clave):
    s = st.session_state
    if clave == "inicio":
        return True
    if clave in ("kpis", "params", "tablero", "export"):
        return clave in s.visto
    if clave == "datos":
        return hay_datos()
    if clave == "decision":
        return len(s.casos) > 0
    if clave == "ejecucion":
        return any(c["estado"] in ("Cerrado", "Reabierto") for c in s.casos)
    if clave == "valid":
        return s.get("v_res") is not None
    return False


def bloqueo(clave):
    """Texto que explica por qué un paso está bloqueado (None si está habilitado)."""
    s = st.session_state
    if clave in ("inicio", "kpis"):
        return None
    if clave == "datos":
        return None if "kpis" in s.visto else "Primero revisa los KPIs de entrada (paso 1)."
    if clave in ("valid", "export"):
        return None if hay_datos() else "Primero carga tus datos en el paso 2."
    if clave == "params":
        return None if hay_datos() else "Primero carga tus datos en el paso 2."
    if clave == "tablero":
        if not hay_datos():
            return "Primero carga tus datos en el paso 2."
        return None if "params" in s.visto else "Primero revisa los parámetros (paso 3)."
    if clave == "decision":
        if not hay_datos():
            return "Primero carga tus datos en el paso 2."
        return None if "tablero" in s.visto else "Primero mira el tablero (paso 4)."
    if clave == "ejecucion":
        return None if s.casos else "Primero confirma una decisión en el paso 5."
    return None


def ir(clave):
    st.session_state.pagina = clave


def pie_siguiente(clave, habilitado=True, tip=None, etiqueta=None):
    """Botón «Siguiente: <paso> ›» al final de cada página."""
    i = ORDEN.index(clave)
    st.divider()
    if i + 1 >= len(ORDEN):
        st.button("‹ Volver al inicio", key=f"b_fin_{clave}", width="stretch", on_click=ir, args=("inicio",))
        return
    sig = ORDEN[i + 1]
    msg = bloqueo(sig)
    st.button(etiqueta or f"Siguiente: {nombre_corto(sig)} ›", key=f"b_sig_{clave}", type="primary", width="stretch",
              disabled=bool(msg) or not habilitado, on_click=ir, args=(sig,))
    if msg:
        st.caption("🔒 " + msg)
    elif tip and not habilitado:
        st.caption("🔒 " + tip)


def abrir_pagina(clave, titulo, que_hago, eyebrow=None):
    """Encabezado común de cada página: logo, estado de datos, título y recuadro «¿Qué hago aquí?»."""
    s = st.session_state
    s.visto.add(clave)
    d = s.get("datos")
    estado = (f'<span class="dot"></span>{d["origen"]} · {d["archivo"]}' if d
              else '<span class="dot off"></span>Sin datos cargados')
    st.markdown(f'<div class="hdr"><div class="brand"><div class="logo">{LOGO}</div><div><div class="bn">Agua Bora</div>'
                f'<div class="bs">DECISION SUPPORT SYSTEM</div></div></div><div class="pill">{estado}</div></div>',
                unsafe_allow_html=True)
    st.markdown(f'<div class="eyebrow">{eyebrow or NOMBRE[clave]}</div><div class="h2">{titulo}</div>', unsafe_allow_html=True)
    caja_ayuda(que_hago)


GLOSARIO = [
    ("OEE", "Efectividad global del equipo: qué parte del tiempo planificado produce bidones buenos al ritmo ideal (D × R × Q)."),
    ("D · Disponibilidad", "Porcentaje del tiempo planificado en que la línea funcionó (sin paradas no planificadas)."),
    ("R · Rendimiento", "Qué tan rápido produce comparado con su capacidad ideal."),
    ("Q · Calidad", "Porcentaje de bidones buenos a la primera."),
    ("MTBF", "Tiempo medio entre fallas: horas de trabajo por cada parada."),
    ("MTTR", "Tiempo medio de reparación: duración promedio de cada parada."),
    ("NPR", "Número de prioridad de riesgo = Severidad × Ocurrencia × Detección. Ordena qué causa revisar primero; no es una probabilidad."),
    ("P10 / P90", "Valor que deja por debajo al 10 % (P10) o al 90 % (P90) de los días. El OEE bajo el P10 o el rechazo sobre el P90 son «días malos»."),
    ("TPM", "Mantenimiento productivo total: operadores y mantenimiento cuidan juntos los equipos (autónomo, planificado o de mejora)."),
    ("C1 · C2 · C3 · C4", "C1 detecta, C2 diagnostica, C3 prioriza, C4 ejecuta y verifica."),
]


def glosario():
    for t, d in GLOSARIO:
        st.markdown(f"**{t}** — {d}")


@st.dialog("Guía rápida", width="large")
def dialogo_guia():
    st.markdown("#### Los 6 pasos")
    for i, (t, d) in enumerate(PASOS_RAPIDOS, 1):
        st.markdown(f"**{i}. {t}** — {d}")
    st.markdown("#### ¿Quién hace qué?")
    for rol, filas in ROL_DECIDE.items():
        st.markdown(f"**{rol}**")
        for senal, decide, plazo in filas:
            st.markdown(f"- *{senal}*: {decide} ({plazo}).")
    st.caption("Los pasos 7 (Validación) y 8 (Exportación) son opcionales y solo piden tener datos cargados.")


PASOS_RAPIDOS = [
    ("KPIs de entrada", "mira qué indicadores se miden y cuáles son las metas."),
    ("Datos", "carga los registros de tu empresa o usa datos de ejemplo."),
    ("Parámetros", "revisa los umbrales, rangos y reglas (puedes dejar los de la tesis)."),
    ("Tablero", "mira cómo va la línea: OEE, paradas y rechazo."),
    ("Decisión", "elige una alerta, entiende la causa y decide qué hacer."),
    ("Resultante", "comprueba con datos posteriores si la acción funcionó."),
]

ROL_DECIDE = {  # (señal / situación, qué decide, plazo) según la Tabla 7 de la tesis
    "Operador": [
        ("Variable fuera de rango 3 lecturas (C1)", "Avisa a Mantenimiento y registra la observación", "Al confirmarse la alerta"),
        ("Sin alerta", "Continúa el monitoreo y el mantenimiento autónomo programado", "Turno en curso"),
        ("Prioridad baja o causa apta para operador", "Ejecuta el mantenimiento autónomo", "En la siguiente limpieza"),
    ],
    "Mantenimiento": [
        ("Variable fuera de rango (C1)", "Inspecciona la causa de mayor NPR (C2) y registra la evidencia", "Según la prioridad (C3)"),
        ("Parada no planificada ≥ 30 min o 2 eventos de la misma causa", "Prioriza con la matriz de criticidad (C3); si es crítica, correctivo inmediato", "Inmediato o en el turno"),
        ("Intervención realizada", "Registra la acción y verifica la resultante antes–después", "Tras la intervención"),
    ],
    "Jefe de Producción": [
        ("OEE diario por estación < P10", "Identifica el factor dominante: D → paradas; R → esperas o velocidad; Q → rechazos", "≤ 24 h"),
        ("Índice de rechazo diario > P90", "Con Calidad: retiene el lote e inspecciona llenado, tapado o sellado según NPR", "Mismo turno"),
        ("Parada por falta de insumos", "Registra la causa y avisa a Logística (fuera del alcance); resecuencia la producción", "Mismo día"),
    ],
}


def barra_lateral():
    s = st.session_state
    with st.sidebar:
        st.markdown(f'<div class="brand" style="margin-bottom:10px"><div class="logo">{LOGO}</div><div><div class="bn">Agua Bora</div>'
                    '<div class="bs">DECISION SUPPORT SYSTEM</div></div></div>', unsafe_allow_html=True)
        st.markdown('<div class="eyebrow" style="margin:8px 0 6px">Flujo de trabajo</div>', unsafe_allow_html=True)
        aviso_puesto = False
        for clave, nombre in PASOS:
            actual = clave == s.pagina
            marca = "●" if actual else ("✓" if completo(clave) else "○")
            msg = bloqueo(clave)
            st.button(f"{marca}  {nombre}", key=f"b_nav_{clave}", type="primary" if actual else "secondary", width="stretch",
                      disabled=bool(msg) and not actual, help=msg, on_click=ir, args=(clave,))
            if msg and not actual and not aviso_puesto:   # el motivo de cada bloqueo sale al pasar el cursor; aquí solo el primero
                st.caption("🔒 " + msg)
                aviso_puesto = True
        st.divider()
        d = s.get("datos")
        abiertos = sum(1 for c in s.casos if c["estado"] in ("Abierto", "Reabierto"))
        if d:
            chip = ('<span class="chip emp">DATOS DE LA EMPRESA</span>' if d["origen"] == "Datos de la empresa"
                    else '<span class="chip sin">DATOS SINTÉTICOS</span>')
            st.markdown(f'<div class="activos"><div class="eyebrow">Datos activos</div>{chip}<br><b>{d["archivo"]}</b><br>'
                        f'Periodo: {periodo_txt(d)}<br>Días: {d["dias"]}<br>Casos abiertos: <b>{abiertos}</b></div>',
                        unsafe_allow_html=True)
        else:
            st.markdown('<div class="activos"><div class="eyebrow">Datos activos</div>Ninguno todavía.<br>'
                        f'Casos abiertos: <b>{abiertos}</b></div>', unsafe_allow_html=True)
        st.write("")
        with st.popover("Reiniciar casos", width="stretch", help="Borra todos los casos sin perder los datos ni los parámetros."):
            st.write("Se borrarán todos los casos. Los datos cargados y los parámetros se conservan.")
            st.button("Sí, borrar todos los casos", key="b_reiniciar_ok", type="primary", width="stretch",
                      disabled=not s.casos, on_click=cb_reiniciar_casos)
        if st.button("Guía rápida", key="b_guia", width="stretch"):
            dialogo_guia()
        with st.expander("Glosario"):
            glosario()



# =============================================================================
# INICIO
# =============================================================================
def usar_ejemplo():
    """Atajo: carga la base de la tesis y salta directo al tablero."""
    s = st.session_state
    if not RUTA_BASE.exists():
        s.aviso = ("err", f"No encontré {NOMBRE_BASE} en la carpeta de la app.")
        s.pagina = "datos"
        return
    df, errs, _ = leer_ruta_cache(str(RUTA_BASE), RUTA_BASE.stat().st_mtime)
    activar_datos(df, "Datos sintéticos", NOMBRE_BASE, "Base de la tesis (ene–jun 2026)")
    s.visto.update({"kpis", "params"})
    s.pagina = "tablero"


def atajo_validacion():
    s = st.session_state
    s.visto.add("kpis")
    if hay_datos():
        s.pagina = "valid"
    else:
        s.aviso = ("info", "Para correr la validación primero carga datos (puedes usar los de la tesis).")
        s.pagina = "datos"


def pagina_inicio():
    abrir_pagina("inicio", "Plataforma de apoyo a la decisión",
                 "Esta app te ayuda a decidir qué hacer cuando la línea rinde mal. Sigue los pasos del menú de la izquierda, "
                 "de arriba hacia abajo: primero mira los indicadores, luego carga datos y al final decide y comprueba si funcionó.",
                 eyebrow="Agua Bora / Lima, Perú")
    st.markdown('<div class="lead">Línea de envasado de bidones retornables de 20 L con <b>3 estaciones en serie</b>: '
                'E1 lavado y sanitizado, E2 llenado y tapado, E3 sellado y etiquetado. Un turno diurno de lunes a sábado. '
                'Los datos de ejemplo son simulados, con fines académicos.</div>', unsafe_allow_html=True)
    st.write("")
    st.markdown("##### Cómo usar esta app en 5 minutos")
    cols = st.columns(6)
    textos = [("KPIs", "Mira qué se mide y cuáles son las metas."), ("Datos", "Carga tus registros o usa datos de ejemplo."),
              ("Parámetros", "Revisa umbrales y reglas (o deja los de la tesis)."), ("Tablero", "Mira cómo va la línea."),
              ("Decisión", "Elige una alerta y decide qué hacer."), ("Resultante", "Comprueba si la acción funcionó.")]
    for i, (col, (t, p)) in enumerate(zip(cols, textos), 1):
        col.markdown(f'<div class="paso-card"><div class="num">{i}</div><div class="t">{t}</div><div class="p">{p}</div></div>',
                     unsafe_allow_html=True)
    st.write("")
    st.button("Empezar", key="b_empezar", type="primary", width="stretch", on_click=ir, args=("kpis",))
    a, b = st.columns(2)
    a.button("Validación (Monte Carlo)", key="b_atajo_val", width="stretch", on_click=atajo_validacion,
             help="Responde: ¿cuánto mejoraría la línea con el DSS? Necesita datos cargados.")
    b.button("Probar con datos de ejemplo", key="b_ejemplo", width="stretch", on_click=usar_ejemplo,
             help="Carga la base de la tesis y te lleva directo al tablero.")
    st.write("")
    st.markdown("##### ¿Quién hace qué?")
    cols = st.columns(3, gap="medium")
    for col, (rol, filas) in zip(cols, ROL_DECIDE.items()):
        html = f'<div class="mini-t">{rol}</div><div style="min-height:15.5rem">'
        for senal, decide, plazo in filas:
            html += f'<div class="mini-p" style="margin-bottom:6px"><b>{senal}.</b> {decide}. <i>({plazo})</i></div>'
        with col:
            with st.container(border=True):
                st.markdown(html + "</div>", unsafe_allow_html=True)
    st.caption("Fuente: reglas de decisión de la parte post (Tabla 7 de la tesis).")
    pie_siguiente("inicio")


# =============================================================================
# PASO 1 - KPIs DE ENTRADA
# =============================================================================
def pagina_kpis():
    s = st.session_state
    abrir_pagina("kpis", "KPIs de entrada",
                 "Aquí ves qué indicadores se miden, qué significan y cuál es la meta de cada uno. Puedes cambiar una meta "
                 "haciendo doble clic sobre ella. La línea base se llena sola cuando cargas tus datos (paso 2).")
    base = kpis_base(df_activo()) if hay_datos() else None
    filas = []
    for k, (v0, u, cond) in META_UI.items():
        nombre, que, formula, modulo = KPI_INFO[k]
        if k == "t_resp":
            lb = "No registrado (planilla manual)"
        elif base is None:
            lb = "—"
        else:
            lb = fmt(TIPO[k], base[k])
        filas.append({"KPI": nombre, "Qué mide": que, "Fórmula": formula, "Línea base": lb,
                      "Condición": cond, "Meta": float(s.metas.get(k, v0)), "Unidad": u, "Módulo": modulo})
    src = pd.DataFrame(filas)
    if base is None:
        caja("Se calcula al cargar los datos (paso 2). Por ahora la línea base aparece como «—».", "warn")
    ed = st.data_editor(
        src, hide_index=True, width="stretch", key="w_ed_metas",
        disabled=["KPI", "Qué mide", "Fórmula", "Línea base", "Condición", "Unidad", "Módulo"],
        column_config={
            "Meta": st.column_config.NumberColumn("Meta (editable)", min_value=0.0, format="%.2f",
                                                  help="Haz doble clic para cambiar la meta. Se usa en el tablero y en la validación."),
            "Línea base": st.column_config.TextColumn(help="Valor actual calculado con los datos cargados."),
            "Qué mide": st.column_config.TextColumn(width="large"),
            "Condición": st.column_config.TextColumn("Cumple si…", help="≥ significa «al menos esta meta»; ≤ significa «como máximo»."),
        })
    nuevas = {}
    for (k, (v0, u, cond)), v in zip(META_UI.items(), ed["Meta"]):
        if pd.notna(v) and abs(float(v) - v0) > 1e-9:
            nuevas[k] = float(v)
    s.metas = nuevas
    st.button("Restablecer metas de la tesis", key="b_reset_metas", width="stretch",
              on_click=lambda: (s.__setitem__("metas", {}), s.pop("w_ed_metas", None)))
    st.caption("La meta de horas de parada (180 h) corresponde a 6 meses de datos; cámbiala si tu periodo es distinto. "
               "Las metas vienen de la Tabla 2 de la tesis y deben validarse con Producción y Mantenimiento.")
    with st.expander("¿Qué significan estas siglas?"):
        glosario()
    pie_siguiente("kpis")


# =============================================================================
# PASO 2 - DATOS
# =============================================================================
def pagina_datos():
    s = st.session_state
    abrir_pagina("datos", "Datos",
                 "Elige de dónde salen los datos: de tu empresa (con la plantilla de Excel) o de datos sintéticos de ejemplo. "
                 "Cuando termines, verás un resumen corto. El detalle está en el tablero (paso 4).")
    aviso = s.pop("aviso", None)
    if aviso:
        caja(aviso[1], aviso[0])
    t1, t2 = st.tabs(["Datos de mi empresa", "Datos sintéticos (de ejemplo)"])
    with t1:
        st.markdown("**1) Descarga la plantilla**")
        st.download_button("Descargar plantilla Excel", construir_plantilla(), file_name="plantilla_agua_bora.xlsx",
                           key="dl_plantilla", width="stretch", type="primary",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           help="Incluye la hoja OEE_Diario con una fila de ejemplo y la hoja Instrucciones.")
        st.markdown("**2) Llénala con tus registros diarios**")
        st.caption("Una fila por estación y por día (3 filas por día). La hoja «Instrucciones» explica cada columna y trae la lista de causas válidas.")
        st.markdown("**3) Súbela**")
        up = st.file_uploader("Archivo .xlsx con la hoja OEE_Diario", type=["xlsx"], key="w_up_emp")
        if up is None:
            s.up_emp_id = None
            s.errores_carga = None
        elif s.get("up_emp_id") != up.file_id:
            s.up_emp_id = up.file_id
            df, errs, avisos = leer_base(io.BytesIO(up.getvalue()))
            if errs:
                s.errores_carga, s.avisos_carga = errs, avisos
            else:
                activar_datos(df, "Datos de la empresa", up.name, "Archivo subido")
                s.avisos_carga = avisos
        errs = s.get("errores_carga")
        if errs:
            st.error(f"Encontré {len(errs)} problema(s) en tu archivo. Corrígelos y vuelve a subirlo.")
            dfe = pd.DataFrame(errs)
            dfe["Resumen"] = [f"Fila {r.Fila}: {r.Problema}" if str(r.Fila) != "—" else r.Problema for r in dfe.itertuples()]
            st.dataframe(dfe[["Fila", "Columna", "Problema"]], hide_index=True, width="stretch")
            st.download_button("Descargar lista de errores (CSV)", dfe[["Fila", "Columna", "Problema"]].to_csv(index=False).encode("utf-8-sig"),
                               file_name="errores_de_carga.csv", mime="text/csv", key="dl_errores")
        elif up is not None and hay_datos() and s.datos["origen"] == "Datos de la empresa":
            caja("Archivo cargado y validado correctamente.", "ok")
        if s.get("avisos_carga"):
            with st.expander(f"Avisos que no impiden cargar ({len(s.avisos_carga)})"):
                for a in s.avisos_carga[:40]:
                    st.write("• " + a)
    with t2:
        op = st.radio("Elige una opción", ["Base de la tesis (ene–jun 2026)", "Casos de prueba", "Generar una base nueva"],
                      key="w_sint_op", horizontal=True)
        if op.startswith("Base de la tesis"):
            st.write("Es la base sintética de 6 meses (150 días laborables) con la que se escribió el Capítulo IV: línea base de OEE 56,72 %.")
            if st.button("Usar la base de la tesis", key="b_usar_tesis", type="primary", width="stretch"):
                if not RUTA_BASE.exists():
                    caja(f"No encontré {NOMBRE_BASE} en la carpeta de la app.", "err")
                else:
                    df, errs, _ = leer_ruta_cache(str(RUTA_BASE), RUTA_BASE.stat().st_mtime)
                    activar_datos(df, "Datos sintéticos", NOMBRE_BASE, "Base de la tesis (ene–jun 2026)")
        elif op == "Casos de prueba":
            casos = listar_casos_prueba()
            if not casos:
                caja("No encontré archivos Base_01 … Base_06 en la carpeta de la app.", "warn")
            else:
                nombres = {c["titulo"]: c for c in casos}
                if s.get("w_caso_sel") not in nombres:
                    s.w_caso_sel = list(nombres)[0]
                elegido = nombres[st.selectbox("Caso de prueba", list(nombres), key="w_caso_sel")]
                if elegido["caso"]:
                    st.info(f"**Caso:** {elegido['caso']}")
                if elegido["ver"]:
                    st.caption(f"Qué deberías ver: {elegido['ver']}")
                if st.button("Usar este caso de prueba", key="b_usar_caso", type="primary", width="stretch"):
                    df, errs, avisos = leer_ruta_cache(str(elegido["ruta"]), elegido["ruta"].stat().st_mtime)
                    if errs:
                        st.error("Este archivo no pasó la validación: " + errs[0]["Problema"])
                    else:
                        activar_datos(df, "Datos sintéticos", elegido["nombre"], elegido["titulo"])
                        s.avisos_carga = avisos
        else:
            st.write("Crea una base nueva remuestreando días de la base de la tesis, con el flujo E1 → E2 → E3 "
                     "(cada estación procesa lo que la anterior dejó conforme). Los domingos no se laboran; los feriados no se excluyen.")
            c1, c2, c3 = st.columns(3)
            ini = c1.date_input("Desde", key="w_gen_ini")
            fin = c2.date_input("Hasta", key="w_gen_fin")
            sem = c3.number_input("Semilla", min_value=0, max_value=10_000_000, step=1, key="w_gen_sem",
                                  help="Con la misma semilla y el mismo periodo siempre sale la misma base.")
            if st.button("Generar y usar esta base", key="b_generar", type="primary", width="stretch"):
                if fin <= ini or (fin - ini).days > 366:
                    caja("Elige un periodo válido: «Hasta» debe ser posterior a «Desde» y no pasar de 12 meses.", "err")
                elif not RUTA_BASE.exists():
                    caja(f"No encontré {NOMBRE_BASE} para remuestrear.", "err")
                else:
                    ref, _, _ = leer_ruta_cache(str(RUTA_BASE), RUTA_BASE.stat().st_mtime)
                    nueva = generar_base(ref, pd.Timestamp(ini), pd.Timestamp(fin), sem)
                    if nueva is None:
                        caja("No se pudo generar la base con ese periodo.", "err")
                    else:
                        activar_datos(nueva, "Datos sintéticos", f"Base generada {ini:%d/%m/%Y}–{fin:%d/%m/%Y} (semilla {sem})",
                                      "Base generada a partir de la base de la tesis")
    st.divider()
    if hay_datos():
        d = s.datos
        df = d["df"]
        chip = ('<span class="chip emp">DATOS DE LA EMPRESA</span>' if d["origen"] == "Datos de la empresa"
                else '<span class="chip sin">DATOS SINTÉTICOS</span>')
        st.markdown(f'##### Datos activos {chip}', unsafe_allow_html=True)
        st.caption(f"{d['archivo']} · {d['detalle']} · periodo {periodo_txt(d)}")
        k = kpis_base(df)
        cols = st.columns(4)
        kpi(cols[0], "Días laborables", n(d["dias"], 0), periodo_txt(d))
        kpi(cols[1], "Eventos de parada", n(k["ev"], 0), "no planificados")
        kpi(cols[2], "Horas de parada", f"{n(k['Pa'])} h", "no planificadas")
        kpi(cols[3], "OEE de línea", f"{n(k['OEE'] * 100)} %", "promedio de las 3 estaciones")
        st.markdown("**Vista previa (10 filas)**")
        pv = df.head(10)[["Fecha", "Est", "Causa", "Parada_No_Planificada_h", "Tiempo_Operativo_h", "Bidones_Procesados",
                          "Bidones_Conformes", "Bidones_No_Conformes"]].copy()
        pv["Fecha"] = pv["Fecha"].dt.strftime("%d/%m/%Y")
        pv.columns = ["Fecha", "Estación", "Causa de parada", "Parada (h)", "Operativo (h)", "Procesados", "Conformes", "No conformes"]
        st.dataframe(pv, hide_index=True, width="stretch")
    else:
        caja("Todavía no hay datos activos. Elige una opción de arriba.", "warn")
    pie_siguiente("datos", habilitado=hay_datos(), tip="Carga datos para continuar.")


# =============================================================================
# PASO 3 - PARAMETROS
# =============================================================================
AYUDAS = {
    "s_p_det": "Probabilidad de que C1 detecte a tiempo (antes de la parada plena) una avería o ajuste correctivo. Supuesto de simulación: la tesis no lo mide.",
    "s_r_dur_par": "Reducción de la duración de la parada cuando C1 detecta a tiempo, en DSS parcial (solo alerta). Supuesto de simulación.",
    "s_r_dur_com": "Reducción de la duración cuando C1 detecta a tiempo y C2–C3 aportan diagnóstico y prioridad, en DSS completo. Supuesto de simulación.",
    "s_t_sin": "Tiempo medio de respuesta sin DSS. La base no lo registra (planilla manual): es un supuesto.",
    "s_t_par": "Tiempo medio de respuesta con alerta de C1 solamente. Supuesto de simulación (meta ≤ 15 min).",
    "s_t_com": "Tiempo medio de respuesta con DSS completo. Supuesto de simulación (meta ≤ 15 min).",
    "s_disp_resp": "Variación relativa del tiempo medio de respuesta entre iteraciones. Permite estimar el % de iteraciones que cumplen la meta.",
    "s_eff_c4": "Probabilidad máxima de eliminar la recurrencia de una causa de equipo ya intervenida y verificada por C4. Supuesto de simulación.",
    "s_tau_c4": "Curva de aprendizaje de C4: la eficacia crece como 1 − exp(−k/τ), con k = ocurrencias previas de la misma causa en la misma estación. Supuesto de simulación.",
    "s_propag": "Fracción de la reducción de horas de equipo que se propaga a la «espera de proceso (línea detenida)». Supuesto de simulación.",
    "s_mej_r_par": "Mejora relativa del ritmo de producción por hora operativa con DSS parcial. Supuesto de simulación.",
    "s_mej_r_com": "Mejora relativa del ritmo de producción por hora operativa con DSS completo. Supuesto de simulación.",
    "s_red_def_par": "Reducción relativa de la tasa de no conformes con DSS parcial. Supuesto de simulación.",
    "s_red_def_com": "Reducción relativa de la tasa de no conformes con DSS completo. Supuesto de simulación.",
}


def leer_supuestos():
    return {k: float(st.session_state.get("w_" + k, v)) for k, v in SUPUESTOS_DEF.items()}


def explica(texto):
    st.markdown(f'<div class="box info"><b>¿Qué es esto?</b> {texto}</div>', unsafe_allow_html=True)


def _umb_df():
    s = st.session_state
    return pd.DataFrame([{
        "Estación": EST_NOM[e], "OEE bajo (P10) %": round(u["oee_p10"] * 100, 2), "OEE de cierre (mediana) %": round(u["oee_med"] * 100, 2),
        "Rechazo alto (P90) %": round(u["rech_p90"] * 100, 2), "Rechazo de cierre (promedio) %": round(u["rech_med"] * 100, 2)}
        for e, u in sorted(s.umb.items())])


def _tab_umbrales():
    s = st.session_state
    explica("Son las líneas rojas del tablero. Si el OEE de una estación en un día baja del <b>P10</b>, o su rechazo sube del <b>P90</b>, "
            "se enciende una alerta. Los valores de «cierre» son los que debe alcanzar la estación para dar el caso por resuelto. "
            "Se calculan con tus datos y puedes cambiarlos.")
    ed = st.data_editor(_umb_df(), hide_index=True, width="stretch", key=f"w_ed_umb_{s.umb_ver}", disabled=["Estación"],
                        column_config={c: st.column_config.NumberColumn(c, min_value=0.0, max_value=100.0, step=0.05, format="%.2f")
                                       for c in ("OEE bajo (P10) %", "OEE de cierre (mediana) %", "Rechazo alto (P90) %",
                                                 "Rechazo de cierre (promedio) %")})
    for e, (_, r) in zip(sorted(s.umb), ed.iterrows()):
        s.umb[e] = dict(oee_p10=r["OEE bajo (P10) %"] / 100, oee_med=r["OEE de cierre (mediana) %"] / 100,
                        rech_p90=r["Rechazo alto (P90) %"] / 100, rech_med=r["Rechazo de cierre (promedio) %"] / 100)
    a, b, c = st.columns(3)
    a.number_input("Parada no planificada desde (minutos)", min_value=1, max_value=480, step=5, key="w_parada_min",
                   help="Una parada de equipo de este tamaño o mayor enciende una alerta (Tabla 7: 30 min).")
    b.number_input("Paradas repetidas de la misma causa (n)", min_value=2, max_value=10, step=1, key="w_rec_n",
                   help="Si la misma causa se repite este número de veces dentro de la ventana, se enciende una alerta (Tabla 7: 2 eventos).")
    c.number_input("Ventana de repetición (días)", min_value=1, max_value=30, step=1, key="w_rec_vent",
                   help="1 día equivale a un turno. Súbela (por ejemplo a 7) para detectar fallas que se repiten en la semana.")
    st.button("Restablecer valores de la tesis", key="b_reset_umb", on_click=_reset_umb,
              help="Vuelve a calcular P10/P90 con los datos cargados (con la base de la tesis son los valores de la Tabla 7).")


def _reset_umb():
    s = st.session_state
    s.umb = referencias_p10_p90(df_activo())
    s.umb_ver += 1
    s.w_parada_min, s.w_rec_n, s.w_rec_vent = 30, 2, 1


def _tab_c1():
    s = st.session_state
    explica("C1 vigila las variables de los sensores. Una variable entra en alerta cuando sale de su rango <b>varias lecturas seguidas</b> "
            "(así un pico aislado no cuenta). Aquí puedes cambiar los rangos y cuántas lecturas seguidas se piden.")
    filas = []
    for eq, vs in VARIABLES_DEF.items():
        for v in vs:
            if v["kind"] in ("range", "max"):
                lo, hi = s.rangos_ov.get((eq, v["key"]), (v["lo"], v["hi"]))
                filas.append({"Equipo": eq, "Variable": v["nombre"], "Unidad": v["unidad"],
                              "Mínimo": float(lo), "Máximo": float(hi), "Sensor": v["sensor"], "_k": v["key"], "_kind": v["kind"]})
    src = pd.DataFrame(filas)
    ed = st.data_editor(src.drop(columns=["_k", "_kind"]), hide_index=True, width="stretch", key="w_ed_rangos",
                        disabled=["Equipo", "Variable", "Unidad", "Sensor"],
                        column_config={"Mínimo": st.column_config.NumberColumn(format="%.2f", help="En las variables «máx.» (ciclo, corriente, rechazos) el mínimo es 0 y no se usa."),
                                       "Máximo": st.column_config.NumberColumn(format="%.2f")})
    ov, malos = {}, []
    for (i, r), (_, o) in zip(ed.iterrows(), src.iterrows()):
        lo, hi = (0.0 if o["_kind"] == "max" else float(r["Mínimo"])), float(r["Máximo"])
        if hi <= lo:
            malos.append(f"{r['Equipo']} · {r['Variable']}")
            continue
        base = next(v for v in VARIABLES_DEF[o["Equipo"]] if v["key"] == o["_k"])
        if (lo, hi) != (float(base["lo"]), float(base["hi"])):
            ov[(o["Equipo"], o["_k"])] = (lo, hi)
    if malos:
        caja("El máximo debe ser mayor que el mínimo; se ignoró el cambio en: " + ", ".join(malos), "warn")
    s.rangos_ov = ov
    st.number_input("Lecturas consecutivas fuera de rango para confirmar una anomalía", min_value=1, max_value=10,
                    step=1, key="w_p_lect", help="Tesis: 3 lecturas consecutivas.")
    st.caption("Las variables de calidad (OK/NG) y de estado (RUN/STOP) no tienen rango numérico. Los rangos son valores de demostración «(ejemplo)», "
               "no especificaciones validadas de los equipos reales.")
    st.button("Restablecer valores de la tesis", key="b_reset_c1", on_click=_reset_c1)


def _reset_c1():
    s = st.session_state
    s.rangos_ov = {}
    s.w_p_lect = LECT_DEF
    s.pop("w_ed_rangos", None)


def _tab_escalas():
    explica("El AMEF califica cada causa de 1 a 10 en Severidad (S), Ocurrencia (O) y Detección (D). Para usar la matriz de criticidad, "
            "esos números se pasan a tres niveles: <b>Baja, Media o Alta</b>. Aquí defines dónde empieza cada nivel.")
    a, b = st.columns(2)
    a.number_input("La escala «Media» empieza en", min_value=2, max_value=9, step=1, key="w_p_medio",
                   help="Tesis: 4 (Baja = 1–3).")
    b.number_input("La escala «Alta» empieza en", min_value=3, max_value=10, step=1, key="w_p_alto",
                   help="Tesis: 7 (Media = 4–6, Alta = 7–10).")
    m, h = int(st.session_state.w_p_medio), int(st.session_state.w_p_alto)
    if m >= h:
        caja("«Media» debe empezar antes que «Alta». Se usan los valores de la tesis hasta que lo corrijas.", "err")
        m, h = MEDIO_DEF, ALTO_DEF
    tabla_html(["Nivel", "Severidad y Ocurrencia", "Detección (D) → Detectabilidad"],
               [["Baja", f"1 – {m - 1}", f"{h} – 10 (la falla es difícil de detectar)"],
                ["Media", f"{m} – {h - 1}", f"{m} – {h - 1}"],
                ["Alta", f"{h} – 10", f"1 – {m - 1} (la falla se detecta fácil)"]])
    st.caption("La Detección se invierte: una D alta del AMEF significa que la falla es difícil de detectar, es decir, Detectabilidad Baja.")
    st.button("Restablecer valores de la tesis", key="b_reset_esc",
              on_click=lambda: st.session_state.update(w_p_medio=MEDIO_DEF, w_p_alto=ALTO_DEF))


def _matriz_df():
    s = st.session_state
    return pd.DataFrame([dict(Severidad=k[0], Frecuencia=k[1], Detectabilidad=k[2], Prioridad=p,
                              Origen="Tesis" if k in TABLA7_TESIS else "Definida por el usuario") for k, p in s.matriz.items()])


def _heat(sev):
    filas = ""
    for fr in NIVELES:
        celdas = ""
        for de in NIVELES:
            p = st.session_state.matriz[(sev, fr, de)]
            bg, fg = COLOR_PRIO[p]
            borde = "outline:2px solid #0d1117;" if (sev, fr, de) in TABLA7_TESIS else ""
            celdas += (f'<td style="background:{bg};color:{fg};text-align:center;font-weight:600;font-size:.72rem;'
                       f'padding:9px 2px;{borde}">{p}</td>')
        filas += f'<tr><td style="font-weight:600;font-size:.72rem;padding:9px 4px">{fr}</td>{celdas}</tr>'
    cab = "".join(f'<th style="text-align:center;font-size:.62rem;padding:6px 2px;letter-spacing:0">{de}</th>' for de in NIVELES)
    return (f'<div class="mini-t">Severidad {sev}</div><table class="tbl" style="table-layout:fixed"><thead><tr>'
            f'<th style="width:22%;font-size:.58rem;padding:6px 4px;letter-spacing:0">Frec.↓ Detect.→</th>{cab}</tr></thead>'
            f'<tbody>{filas}</tbody></table>')


def _reset_matriz():
    s = st.session_state
    s.matriz = {k: v["prioridad"] for k, v in MATRIZ_DEF.items()}
    s.matriz_ver += 1


def _tab_matriz():
    s = st.session_state
    explica("C3 usa esta matriz para decidir <b>qué tan urgente</b> es una causa según su Severidad, su Frecuencia y su Detectabilidad. "
            "Las 4 filas de la tesis están bloqueadas; las otras 23 las puedes cambiar. <b>Detectabilidad Baja agrava la prioridad</b> "
            "porque la falla es difícil de detectar y pasa más tiempo sin que nadie la note.")
    ed = st.data_editor(_matriz_df(), hide_index=True, width="stretch", key=f"w_ed_matriz_{s.matriz_ver}", height=420,
                        disabled=["Severidad", "Frecuencia", "Detectabilidad", "Origen"],
                        column_config={"Prioridad": st.column_config.SelectboxColumn("Prioridad", options=["Crítica", "Alta", "Media", "Baja"],
                                                                                      required=True),
                                       "Origen": st.column_config.TextColumn(help="«Tesis»: fila de la tesis (bloqueada). «Definida por el usuario»: editable.")})
    nueva = {(r.Severidad, r.Frecuencia, r.Detectabilidad): r.Prioridad for r in ed.itertuples()}
    tocadas = [k for k, p in TABLA7_TESIS.items() if nueva[k] != p]
    if tocadas:
        for k in tocadas:
            nueva[k] = TABLA7_TESIS[k]
        s.matriz = nueva
        s.matriz_ver += 1
        s.aviso_matriz = "Las filas marcadas «Tesis» están bloqueadas: se restauró su valor."
        st.rerun()
    s.matriz = nueva
    if s.get("aviso_matriz"):
        caja(s.pop("aviso_matriz"), "warn")
    st.markdown("**Vista de mapa de calor** — filas: Frecuencia · columnas: Detectabilidad (el borde negro marca las filas de la tesis)")
    cols = st.columns(3)
    for col, sev in zip(cols, NIVELES):
        col.markdown(_heat(sev), unsafe_allow_html=True)
    a, b, c = st.columns(3)
    a.button("Restablecer valores de la tesis", key="b_reset_mat", width="stretch", on_click=_reset_matriz)
    b.download_button("Descargar CSV", _matriz_df().to_csv(index=False).encode("utf-8-sig"), file_name="matriz_criticidad.csv",
                      mime="text/csv", key="dl_matriz", width="stretch")
    up = c.file_uploader("Cargar CSV", type=["csv"], key="w_up_mat", label_visibility="collapsed")
    if up is not None and s.get("up_mat_id") != up.file_id:
        s.up_mat_id = up.file_id
        try:
            dfc = pd.read_csv(up)
            dfc.columns = [str(x).strip() for x in dfc.columns]
            nuevo = {}
            for r in dfc.itertuples():
                k = (str(r.Severidad).strip(), str(r.Frecuencia).strip(), str(r.Detectabilidad).strip())
                p = str(r.Prioridad).strip()
                if k not in MATRIZ_DEF or p not in COLOR_PRIO:
                    raise ValueError(f"fila no válida: {k} → {p}")
                nuevo[k] = p
            if set(nuevo) != set(MATRIZ_DEF):
                raise ValueError("deben venir las 27 combinaciones, una vez cada una")
            for k, p in TABLA7_TESIS.items():
                nuevo[k] = p
            s.matriz = {k: nuevo[k] for k in MATRIZ_DEF}
            s.matriz_ver += 1
            s.aviso_matriz = "CSV cargado. Las filas de la tesis se mantienen como en la tesis."
            st.rerun()
        except Exception as exc:
            caja(f"No pude cargar el CSV: {exc}.", "err")


def _tab_amef():
    s = st.session_state
    explica("El AMEF es la lista de causas posibles de cada falla. Cada causa tiene tres notas del 1 al 10: <b>S</b> (qué tan grave es), "
            "<b>O</b> (qué tan seguido ocurre) y <b>D</b> (qué tan difícil es detectarla). El NPR = S × O × D ordena qué causa revisar primero. "
            "Solo las cifras de la Tabla 6 de la tesis son datos; el resto son «(ejemplo)».")
    filas = []
    for (eq, vkey, direc), cs in AMEF_DEF.items():
        v = next(x for x in VARIABLES_DEF[eq] if x["key"] == vkey)
        for c in cs:
            S, O, D = s.amef_ov.get((eq, vkey, direc, c["causa"]), (c["S"], c["O"], c["D"]))
            filas.append({"Equipo": eq, "Variable": v["nombre"], "Dirección": direc, "Causa": c["causa"], "S": int(S), "O": int(O),
                          "D": int(D), "Origen": "Tesis" if not c["ejemplo"] else "(ejemplo)", "_id": (eq, vkey, direc, c["causa"])})
    src = pd.DataFrame(filas)
    cfg = {c: st.column_config.NumberColumn(c, min_value=1, max_value=10, step=1) for c in ("S", "O", "D")}
    ed = st.data_editor(src.drop(columns=["_id"]), hide_index=True, width="stretch", key="w_ed_amef", height=420,
                        disabled=["Equipo", "Variable", "Dirección", "Causa", "Origen"], column_config=cfg)
    ov = {}
    for (_, r), (_, o) in zip(ed.iterrows(), src.iterrows()):
        base = next(c for c in AMEF_DEF[o["_id"][:3]] if c["causa"] == o["_id"][3])
        trio = (int(np.clip(r["S"], 1, 10)), int(np.clip(r["O"], 1, 10)), int(np.clip(r["D"], 1, 10)))
        if trio != (base["S"], base["O"], base["D"]):
            ov[o["_id"]] = trio
    s.amef_ov = ov
    st.caption("El NPR se recalcula solo (S × O × D). Cuando un caso se cierra, C4 sube la O en 1 y baja la D en 1 de la causa confirmada.")
    st.button("Restablecer valores de la tesis", key="b_reset_amef", on_click=_reset_amef)


def _reset_amef():
    s = st.session_state
    s.amef_ov = {}
    s.d_ajustes = {}
    s.pop("w_ed_amef", None)


def _tab_supuestos():
    explica("La validación (paso 7) simula miles de días para estimar cuánto mejoraría la línea con el DSS. Para eso necesita estos números. "
            "<b>No salen de la base ni de la tesis: son supuestos</b> y puedes cambiarlos para ver qué pasa.")

    def num(label, key, mn, mx, paso, formato="%.2f"):
        return st.number_input(label, min_value=mn, max_value=mx, step=paso, format=formato,
                               key="w_" + key, help=AYUDAS[key])
    st.markdown("**C1 · Detección anticipada**")
    c1, c2, c3 = st.columns(3)
    with c1:
        num("Detección anticipada C1 · supuesto", "s_p_det", 0.0, 1.0, 0.05)
    with c2:
        num("Reducción de duración, DSS parcial · supuesto", "s_r_dur_par", 0.0, 0.95, 0.05)
    with c3:
        num("Reducción de duración, DSS completo · supuesto", "s_r_dur_com", 0.0, 0.95, 0.05)
    st.markdown("**Tiempo de respuesta a la alerta** (meta ≤ 15 min)")
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        num("Sin DSS (min) · supuesto", "s_t_sin", 1.0, 240.0, 1.0, "%.1f")
    with c2:
        num("DSS parcial (min) · supuesto", "s_t_par", 1.0, 240.0, 1.0, "%.1f")
    with c3:
        num("DSS completo (min) · supuesto", "s_t_com", 1.0, 240.0, 1.0, "%.1f")
    with c4:
        num("Dispersión entre iteraciones · supuesto", "s_disp_resp", 0.0, 1.0, 0.05)
    st.markdown("**C4 · Eficacia y curva de aprendizaje** (solo DSS completo)")
    c1, c2, c3 = st.columns(3)
    with c1:
        num("Eficacia máxima de C4 · supuesto", "s_eff_c4", 0.0, 1.0, 0.05)
    with c2:
        num("Constante de aprendizaje τ · supuesto", "s_tau_c4", 0.5, 100.0, 0.5, "%.1f")
    with c3:
        num("Propagación a «espera de proceso» · supuesto", "s_propag", 0.0, 1.0, 0.05)
    st.markdown("**Rendimiento y calidad**")
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        num("Mejora de rendimiento, parcial · supuesto", "s_mej_r_par", 0.0, 0.5, 0.01)
    with c2:
        num("Mejora de rendimiento, completo · supuesto", "s_mej_r_com", 0.0, 0.5, 0.01)
    with c3:
        num("Reducción de no conformes, parcial · supuesto", "s_red_def_par", 0.0, 0.95, 0.05)
    with c4:
        num("Reducción de no conformes, completo · supuesto", "s_red_def_com", 0.0, 0.95, 0.05)
    st.button("Restablecer valores por defecto", key="b_reset_sup", on_click=lambda: st.session_state.update(
        {"w_" + k: v for k, v in SUPUESTOS_DEF.items()}))


def pagina_params():
    abrir_pagina("params", "Parámetros",
                 "Aquí están las reglas que usa la app: cuándo se enciende una alerta, qué rangos vigilan los sensores, "
                 "cómo se califica el riesgo y qué tan urgente es cada causa. Todo viene con los valores de la tesis; "
                 "si no estás seguro, déjalos como están y sigue.")
    tabs = st.tabs(["Umbrales de alerta", "Detección C1", "Escalas S/O/D", "Matriz de criticidad", "AMEF", "Supuestos Monte Carlo"])
    with tabs[0]:
        _tab_umbrales()
    with tabs[1]:
        _tab_c1()
    with tabs[2]:
        _tab_escalas()
    with tabs[3]:
        _tab_matriz()
    with tabs[4]:
        _tab_amef()
    with tabs[5]:
        _tab_supuestos()
    pie_siguiente("params")



# =============================================================================
# PASO 4 - TABLERO (TENDENCIAS Y OEE)
# =============================================================================
def tabla1(df, clave):
    g = tabla_estaciones(df)
    k = kpis_base(df)
    causas = df[df["Causa"] != SIN_PARADA].groupby(["Estacion_ID", "Causa"])["Parada_No_Planificada_h"].sum()

    def principal(e):
        if e not in causas.index.get_level_values(0):
            return "—"
        sr = causas.loc[e].sort_values(ascending=False)
        return f"{sr.index[0]} ({n(sr.iloc[0])} h)" if len(sr) else "—"

    def fila(nombre, f, linea=""):
        return [nombre] + [f(e) for e in (1, 2, 3)] + [linea]
    filas = [
        fila("Disponibilidad (D)", lambda e: f"{n(g.loc[e, 'D'] * 100)} %", f"{n(k['D'] * 100)} %"),
        fila("Rendimiento (R)", lambda e: f"{n(g.loc[e, 'R'] * 100)} %", f"{n(k['R'] * 100)} %"),
        fila("Calidad (Q)", lambda e: f"{n(g.loc[e, 'Q'] * 100)} %", f"{n(k['Q'] * 100)} %"),
        fila("OEE", lambda e: f"{n(g.loc[e, 'OEE'] * 100)} %", f"{n(k['OEE'] * 100)} % (promedio de estaciones)"),
        fila("Parada no planificada", lambda e: f"{n(g.loc[e, 'Pa'])} h ({int(g.loc[e, 'ev'])} eventos)", f"{n(k['Pa'])} h ({k['ev']} eventos)"),
        fila("MTBF / MTTR", lambda e: f"{n(g.loc[e, 'MTBF'])} h / {n(g.loc[e, 'MTTR'])} h", f"{n(k['MTBF'])} h / {n(k['MTTR'])} h"),
        fila("Bidones no conformes", lambda e: f"{n(g.loc[e, 'NC'], 0)} ({n(g.loc[e, 'rech'] * 100)} %)", f"{n(k['nc'], 0)} ({n(k['rechazo'] * 100)} %)"),
        fila("Causa principal de parada", principal, "—"),
    ]
    tabla_html(["Indicador", "E1 Lavado y sanitizado", "E2 Llenado y tapado", "E3 Sellado y etiquetado", "Línea"], filas)
    st.caption("MTBF = tiempo operativo / eventos de parada; MTTR = horas de parada / eventos. El % de no conformes se expresa sobre "
               "los bidones procesados por cada estación.")
    fig = go.Figure()
    for nombre, col, color in (("Disponibilidad", "D", AZUL), ("Rendimiento", "R", AMBAR), ("Calidad", "Q", VERDE), ("OEE", "OEE", INK)):
        fig.add_bar(x=[envolver(EST_NOM[e], 16) for e in (1, 2, 3)], y=(g[col] * 100).round(2).tolist(), name=nombre,
                    marker_color=color, texttemplate="%{y:.1f} %", textposition="outside")
    fig.update_yaxes(range=[0, 118], title="%")
    st.plotly_chart(estilo_fig(fig, 380, "Factores del OEE por estación"), width="stretch", key=f"g_t1_{clave}")


def grafico_pareto(df, clave, altura=470):
    p = tabla_pareto(df)
    if p.empty:
        caja("No hay paradas no planificadas en este periodo.", "info")
        return
    fig = make_subplots(specs=[[{"secondary_y": True}]])
    x = [envolver(c, 12) for c in p["Causa"]]
    fig.add_bar(x=x, y=p["Horas"].round(2), name="Horas de parada", marker_color=AZUL, texttemplate="%{y:.1f} h",
                textposition="outside", secondary_y=False)
    fig.add_scatter(x=x, y=(p["Acum"] * 100).round(1), name="% acumulado", mode="lines+markers", line=dict(color=INK),
                    secondary_y=True)
    fig.update_yaxes(title="Horas", secondary_y=False, range=[0, float(p["Horas"].max()) * 1.2])
    fig.update_yaxes(title="% acumulado", secondary_y=True, range=[0, 105], tickvals=[0, 25, 50, 75, 100],
                     ticktext=["0 %", "25 %", "50 %", "75 %", "100 %"], showgrid=False)
    st.plotly_chart(estilo_fig(fig, altura, "Pareto de horas de parada no planificada por causa"), width="stretch", key=f"g_par_{clave}")
    t = p.copy()
    t["Horas"] = t["Horas"].map(lambda v: n(v))
    t["Pct"] = (p["Pct"] * 100).map(lambda v: f"{n(v)} %")
    t["Acum"] = (p["Acum"] * 100).map(lambda v: f"{n(v)} %")
    t.columns = ["Causa", "Horas de parada", "N.º de eventos", "% del total", "% acumulado"]
    st.dataframe(t, hide_index=True, width="stretch")


def serie_oee(df, freq):
    d = df.copy()
    if freq == "Diario":
        d["_p"] = d["Fecha"]
    elif freq == "Semanal":
        d["_p"] = d["Fecha"] - pd.to_timedelta(d["Fecha"].dt.weekday, unit="D")
    else:
        d["_p"] = d["Fecha"].dt.to_period("M").dt.start_time
    g = d.groupby(["_p", "Estacion_ID"]).agg(Tp=("Tiempo_Planificado_h", "sum"), To=("Tiempo_Operativo_h", "sum"),
                                              Teor=("Teorica", "sum"), Proc=("Bidones_Procesados", "sum"),
                                              Conf=("Bidones_Conformes", "sum")).reset_index()
    g["OEE"] = (g.To / g.Tp.replace(0, np.nan)) * (g.Proc / g.Teor.replace(0, np.nan)) * (g.Conf / g.Proc.replace(0, np.nan))
    piv = g.pivot(index="_p", columns="Estacion_ID", values="OEE")
    piv["Línea"] = piv.mean(axis=1)
    return piv


def pestana_tendencias(df):
    freq = st.radio("Ver por", ["Diario", "Semanal", "Mensual"], horizontal=True, key="w_tend_freq",
                    help="Diario muestra mucho ruido; semanal o mensual muestran mejor la tendencia.")
    piv = serie_oee(df, freq)
    fig = go.Figure()
    colores = {1: AZUL, 2: AMBAR, 3: VERDE}
    for e in (1, 2, 3):
        fig.add_scatter(x=piv.index, y=(piv[e] * 100).round(2), mode="lines", name=EST_CORTO[e],
                        line=dict(color=colores[e], width=1.4 if freq == "Diario" else 2), opacity=0.7 if freq == "Diario" else 1)
    etiquetas = freq != "Diario"
    fig.add_scatter(x=piv.index, y=(piv["Línea"] * 100).round(2), mode="lines+markers+text" if etiquetas else "lines",
                    name="Línea (promedio)", line=dict(color=INK, width=3),
                    text=[f"{v:.1f}" for v in piv["Línea"] * 100] if etiquetas else None, textposition="top center")
    if not etiquetas:
        ult = piv["Línea"].dropna()
        if len(ult):
            fig.add_scatter(x=[ult.index[-1]], y=[round(ult.iloc[-1] * 100, 2)], mode="markers+text", showlegend=False,
                            text=[f"{ult.iloc[-1] * 100:.1f}"], textposition="top center", marker=dict(color=INK, size=8))
    fig.add_scatter(x=[piv.index.min(), piv.index.max()], y=[METAS["OEE"][1] * 100] * 2, mode="lines", name=f"Meta {META_TXT['OEE']}",
                    line=dict(color=VERDE, dash="dash", width=2))
    fig.update_yaxes(title="OEE (%)")
    fig.update_xaxes(tickformat="%b-%y" if freq == "Mensual" else "%d/%m")
    st.plotly_chart(estilo_fig(fig, 430, f"OEE de la línea y de cada estación ({freq.lower()})"), width="stretch", key="g_tend")


def pestana_rechazo(df, umb):
    g = tabla_estaciones(df)
    fig = go.Figure()
    fig.add_bar(x=[envolver(EST_NOM[e], 16) for e in (1, 2, 3)], y=(g["rech"] * 100).round(2), name="Rechazo del periodo",
                marker_color=INK, texttemplate="%{y:.2f} %", textposition="inside", insidetextanchor="end",
                textfont=dict(color="#fff"))
    fig.add_scatter(x=[envolver(EST_NOM[e], 16) for e in (1, 2, 3)], y=[round(umb[e]["rech_p90"] * 100, 2) for e in (1, 2, 3)],
                    mode="markers", name="Umbral P90 (alerta)", marker=dict(symbol="line-ew", size=70, line=dict(width=3, color=ROJO)))
    fig.update_yaxes(title="% de bidones procesados", range=[0, max(3.5, float(umb[3]["rech_p90"] * 100 * 1.3), float(umb[2]["rech_p90"] * 100 * 1.3))])
    st.plotly_chart(estilo_fig(fig, 360, "Índice de rechazo por estación y su umbral P90"), width="stretch", key="g_rech_bar")
    d = df.copy()
    d["rech"] = d["Bidones_No_Conformes"] / d["Bidones_Procesados"].replace(0, np.nan)
    piv = d.pivot(index="Fecha", columns="Estacion_ID", values="rech")
    fig = go.Figure()
    colores = {1: AZUL, 2: AMBAR, 3: VERDE}
    for e in (1, 2, 3):
        fig.add_scatter(x=piv.index, y=(piv[e] * 100).round(2), mode="lines", name=EST_CORTO[e], line=dict(color=colores[e], width=1.3))
        fig.add_scatter(x=[piv.index.min(), piv.index.max()], y=[umb[e]["rech_p90"] * 100] * 2, mode="lines",
                        name=f"P90 {EST_CORTO[e]}", line=dict(color=colores[e], width=1.6, dash="dash"))
    fig.update_yaxes(title="Rechazo diario (%)")
    fig.update_xaxes(tickformat="%d/%m")
    st.plotly_chart(estilo_fig(fig, 380, "Rechazo diario de cada estación (línea punteada = su P90)"), width="stretch", key="g_rech_dia")
    filas = []
    for e in (1, 2, 3):
        x = piv[e].dropna()
        sobre = int((x > umb[e]["rech_p90"]).sum())
        filas.append([EST_NOM[e], f"{n(umb[e]['rech_p90'] * 100)} %", f"{sobre} de {len(x)} días", f"{n(g.loc[e, 'rech'] * 100)} %"])
    tabla_html(["Estación", "Umbral P90", "Días sobre el umbral", "Rechazo del periodo"], filas)


def pagina_tablero():
    s = st.session_state
    abrir_pagina("tablero", "Tablero (tendencias y OEE)",
                 "Aquí solo ves resultados: cómo va la línea frente a las metas. No hay que decidir nada todavía. "
                 "Cuando termines de mirar, ve al paso 5 para decidir qué hacer con las alertas.")
    df_all = df_activo()
    opciones = ["Todos"] + [int(m) for m in sorted(df_all["Mes"].unique())]
    if s.get("w_tab_mes") not in opciones:
        s.w_tab_mes = "Todos"
    mes = st.selectbox("Periodo", opciones, key="w_tab_mes", format_func=lambda m: "Todo el periodo" if m == "Todos" else mes_lbl(m))
    df = df_all if mes == "Todos" else df_all[df_all["Mes"] == mes]
    tabs = st.tabs(["Resumen", "Tendencias", "Por estación", "Pareto de paradas", "Índice de rechazo"])
    with tabs[0]:
        k = kpis_base(df)
        g = tabla_estaciones(df)
        alertas = alertas_en(df, df_all, s.umb, leer_par())
        comp = (mes == "Todos")
        cols = st.columns(4)
        kpi(cols[0], "OEE de línea", f"{n(k['OEE'] * 100)} %", f"Meta {META_TXT['OEE']}", "ok" if cumple("OEE", k["OEE"]) else "bad")
        kpi(cols[1], "Índice de rechazo", f"{n(k['rechazo'] * 100)} %", f"Meta {META_TXT['rechazo']}", "ok" if cumple("rechazo", k["rechazo"]) else "bad")
        kpi(cols[2], "Horas de parada", f"{n(k['Pa'])} h", f"Meta {META_TXT['Pa']}" if comp else "La meta es para todo el periodo",
            ("ok" if cumple("Pa", k["Pa"]) else "bad") if comp else None)
        kpi(cols[3], "Alertas activas", n(len(alertas), 0), "en el periodo elegido (ver paso 5)")
        peor = int(g["OEE"].idxmin())
        mas_paradas = int(g["Pa"].idxmax())
        brecha = (METAS["OEE"][1] - k["OEE"]) * 100
        if brecha > 0:
            msg = (f"La línea está <b>{n(brecha, 1)} puntos</b> debajo de la meta de OEE ({META_TXT['OEE']}). La estación con menor OEE es "
                   f"<b>{EST_CORTO[peor]}</b> ({n(g.loc[peor, 'OEE'] * 100)} %) y la que más horas de parada acumula es <b>{EST_CORTO[mas_paradas]}</b> "
                   f"({n(g.loc[mas_paradas, 'Pa'])} h).")
            tipo = "warn"
        else:
            msg = (f"La línea <b>cumple la meta</b> de OEE ({META_TXT['OEE']}). La estación con menor OEE es <b>{EST_CORTO[peor]}</b> "
                   f"({n(g.loc[peor, 'OEE'] * 100)} %).")
            tipo = "ok"
        st.write("")
        caja(msg, tipo)
    with tabs[1]:
        pestana_tendencias(df)
    with tabs[2]:
        tabla1(df, "tab")
    with tabs[3]:
        grafico_pareto(df, "tab")
    with tabs[4]:
        pestana_rechazo(df, s.umb)
    pie_siguiente("tablero", etiqueta="Ver decisiones para las alertas (paso 5) ›")


# =============================================================================
# PASO 5 - DECISION (PARTE POST)
# =============================================================================
ROLES = ["Operador", "Mantenimiento", "Jefe de Producción"]
DIAS_ES = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
PASOS_WZ = ["Qué pasó", "Por qué pasó", "Qué tan urgente es", "Qué decido"]


def _causas_wz():
    s = st.session_state
    wz = s.wz
    desc = set()
    if wz.get("caso_id"):
        caso = next((c for c in s.casos if c["id"] == wz["caso_id"]), None)
        if caso:
            desc = {tuple(x) for x in caso["descartadas"]}
    return causas_para_alerta(wz["alerta"], s.d_ajustes, desc)


def caso_duplicado(al_id, causa):
    """Caso ABIERTO de la misma alerta (o variable de sensor) y la misma causa, si existe."""
    for c in st.session_state.casos:
        if (c["estado"] == "Abierto" and c["alerta"]["id"] == al_id and c["causa"].get("ident") == causa.get("ident")
                and c["causa"]["causa"] == causa["causa"]):
            return c
    return None


def ir_a_caso(caso_id):
    """Abre el paso 6 con ese caso ya seleccionado."""
    s = st.session_state
    s.pagina = "ejecucion"
    s.w_ej_caso = caso_id
    s.wz = None


def texto_caso_creado(caso, verbo="creado"):
    a = caso["alerta"]
    return (f"✓ Caso {caso['id']} {verbo}: {EST_NOM[a['est']]} · {caso['causa']['causa']} · prioridad {caso['prioridad']}")


def wz_mover(delta):
    s = st.session_state
    s.wz["paso"] = int(np.clip(s.wz["paso"] + delta, 0, 3))
    s.wz.pop("caso_creado", None)
    s.wz.pop("duplicado", None)


def wz_cancelar():
    st.session_state.wz = None


def wz_confirmar():
    s = st.session_state
    wz = s.wz
    al = wz["alerta"]
    causas = _causas_wz()
    elegida = next((c for c in causas if c["ident"] == wz.get("causa_ident")), None)
    if al["tipo"] == "insumos":
        elegida = dict(causa=al["causa_parada"], efecto="Falta de insumos (fuera del alcance del prototipo)", S=0, O=0, D=0, npr=0,
                       equipo="—", vkey="—", vnombre="—", direccion="—", ident=None, ejemplo=False)
        c3 = None
    else:
        if elegida is None:
            return
        c3 = evaluar_c3(elegida)
    if wz.get("caso_id"):
        caso = next(c for c in s.casos if c["id"] == wz["caso_id"])
        info = decision_info(al, elegida, c3)
        caso.update(causa=dict(elegida), prioridad=c3["prioridad"] if c3 else "No aplica",
                    tipo_tpm=c3["tipo"] if c3 else "Fuera del alcance (Logística)", regla=c3["regla"] if c3 else "", estado="Abierto",
                    plazo=info["plazo"], plazo_ref=info["plazo_ref"], accion=info["accion"])
        verbo = "actualizado"
    else:
        dup = caso_duplicado(al["id"], elegida)
        if dup:                       # segundo clic o caso repetido: no se crea otro
            if wz.get("caso_creado") != dup["id"]:
                wz["caso_creado"], wz["duplicado"] = dup["id"], True
                st.toast(f"Este caso ya está abierto ({dup['id']})", icon="⚠️")
            return
        caso = nuevo_caso(al, elegida, c3, origen="sensor" if al["tipo"] == "sensor" else "datos")
        if al["tipo"] == "sensor":
            caso["sensor"] = wz["sensor"]
        s.casos.append(caso)
        verbo = "creado"
    wz["caso_creado"], wz["duplicado"], wz["msg"] = caso["id"], False, texto_caso_creado(caso, verbo)
    st.toast(wz["msg"], icon="✅")


def _tarjeta_causas(causas):
    filas = []
    for k, c in enumerate(causas):
        col = NPR_COL[color_npr(c["npr"])]
        vig = [x for x in causas if not x["descartada"]]
        marca = "★ " if (vig and c is vig[0]) else ""
        nombre = marca + c["causa"] + (" <i>(ejemplo)</i>" if c["ejemplo"] else "")
        if c["descartada"]:
            nombre = f"<s>{nombre}</s> <b>(descartada)</b>"
        if c["ajustada"]:
            nombre += " ⚑"
        filas.append([nombre, f"{c['equipo']} · {c['vnombre']}", c["efecto"], c["S"], c["O"], c["D"],
                      f'<span class="npr" style="background:{col[0]};color:{col[1]}">{c["npr"]}</span>'])
    tabla_html(["Causa probable", "Equipo · variable", "Efecto de la falla", "S", "O", "D", "NPR"], filas)


def _wizard():
    s = st.session_state
    wz = s.wz
    al = wz["alerta"]
    paso = wz["paso"]
    st.markdown(f"##### Paso {paso + 1} de 4 · {PASOS_WZ[paso]}")
    st.progress((paso + 1) / 4)
    causas = _causas_wz()
    vigentes = [c for c in causas if not c["descartada"]]
    puede_seguir = True
    dup = None
    if paso == 0:
        extra = ""
        if al["tipo"] == "oee":
            nombres = {"D": "D (paradas)", "R": "R (esperas o velocidad)", "Q": "Q (rechazos)"}
            extra = f"<br><b>Factor que más cayó:</b> {nombres.get(al.get('factor', 'D'))}"
        st.markdown(f'<div class="decision"><div class="t">{al["senal"]}</div><div class="m"><b>Fecha:</b> {pd.Timestamp(al["fecha"]):%d/%m/%Y}'
                    f'<br><b>Valor observado:</b> {al["valor"]}<br><b>Umbral:</b> {al["umbral"]}{extra}</div></div>', unsafe_allow_html=True)
        st.caption("Esto es lo que encendió la alerta. En el siguiente paso verás las causas posibles.")
    elif paso == 1:
        if al["tipo"] == "insumos":
            caja(f"La causa ya se conoce: <b>{al['causa_parada']}</b>. Es un problema de abastecimiento y está <b>fuera del alcance</b> del "
                 "prototipo: se registra y se avisa a Logística.", "info")
        elif not causas:
            caja("No hay causas del AMEF asociadas a esta alerta.", "warn")
            puede_seguir = False
        else:
            st.caption("Las causas están ordenadas por NPR (S × O × D): la ★ es la que conviene revisar primero. Elegir una causa no la confirma: "
                       "Mantenimiento debe inspeccionar y dejar evidencia.")
            _tarjeta_causas(causas)
            if not vigentes:
                caja("Ya se revisaron todas las causas del AMEF. Escala al jefe de Mantenimiento para un análisis de causa raíz.", "warn")
                puede_seguir = False
            else:
                opciones = {f"{'★ ' if i == 0 else ''}{c['causa']}{' (ejemplo)' if c['ejemplo'] else ''} · {c['equipo']} — NPR {c['npr']}": c
                            for i, c in enumerate(vigentes)}
                key = f"w_wz_causa_{al['id']}_{len(s.casos)}"
                elegida = opciones[st.radio("¿Qué causa vas a revisar?", list(opciones), key=key)]
                wz["causa_ident"] = elegida["ident"]
            st.caption("La relación entre la alerta y los equipos del AMEF es un supuesto de demostración «(ejemplo)».")
    elif paso == 2:
        if al["tipo"] == "insumos":
            caja("No aplica: la prioridad de mantenimiento no se calcula para faltas de insumos (fuera del alcance).", "info")
        else:
            c = next((x for x in vigentes if x["ident"] == wz.get("causa_ident")), vigentes[0] if vigentes else None)
            if c is None:
                caja("Elige primero una causa en el paso anterior.", "warn")
                puede_seguir = False
            else:
                wz["causa_ident"] = c["ident"]
                r = evaluar_c3(c)
                bg, fg = COLOR_PRIO[r["prioridad"]]
                a, b = st.columns([1, 2.2], gap="large")
                a.markdown(f'<div class="eyebrow">Prioridad</div><div class="prio" style="background:{bg};color:{fg}">{r["prioridad"].upper()}</div>',
                           unsafe_allow_html=True)
                b.markdown(f"**Causa:** {c['causa']} (NPR {c['npr']})  \n**Tipo de mantenimiento TPM:** {r['tipo']} — {r['accion']}  \n"
                           f"**Plazo según la prioridad:** {plazo_prio(r['prioridad'], r['tipo'])}")
                tabla_html(["Dimensión", "Valor del AMEF", "Nivel en la matriz"],
                           [["Severidad", f"S = {c['S']}", r["sev"]], ["Frecuencia", f"O = {c['O']}", r["frec"]],
                            ["Detectabilidad (D invertida)", f"D = {c['D']}", r["det"]]])
                st.caption(f"Regla aplicada: {r['regla']}. La matriz se configura en el paso 3.")
    else:
        c = next((x for x in vigentes if x["ident"] == wz.get("causa_ident")), None)
        if al["tipo"] != "insumos" and c is None:
            caja("Elige primero una causa (paso 2).", "warn")
            puede_seguir = False
        else:
            if al["tipo"] == "insumos":
                info = decision_info(al, None, None)
                causa_txt = f"Falta de insumos: {al['causa_parada']}"
            else:
                info = decision_info(al, c, evaluar_c3(c))
                causa_txt = c["causa"]
            ref = f' <span style="color:#6b7280;font-size:.85rem">{info["plazo_ref"]}</span>' if info["plazo_ref"] else ""
            etiqueta_plazo = "Plazo (Tabla 7)" if al["tipo"] in ("rechazo", "oee", "insumos") else "Plazo"
            st.markdown(
                f'<div class="decision"><div class="t">Decisión</div><div class="m"><b>Qué hacer ahora:</b> {info["accion"]}<br>'
                f'<b>Causa a revisar:</b> {causa_txt}<br><b>Prioridad:</b> {info["prio_txt"]}<br>'
                f'<b>Responsable:</b> {al["resp"]}<br><b>{etiqueta_plazo}:</b> {info["plazo"]}{ref}<br>'
                f'<b>Cómo sabré que funcionó:</b> {al["cierre"]}</div></div>', unsafe_allow_html=True)
            if not wz.get("caso_id"):
                cand = c if al["tipo"] != "insumos" else dict(causa=al["causa_parada"], ident=None)
                dup = caso_duplicado(al["id"], cand)
    creado = wz.get("caso_creado") if paso == 3 else None
    if paso == 3 and (creado or dup):
        id_ = creado or dup["id"]
        if creado and not wz.get("duplicado") and wz.get("msg"):
            caja(wz["msg"], "ok")
        else:
            caja(f"Este caso ya está abierto (<b>{id_}</b>).", "warn")
        st.button("Ir a 6. Ejecución y resultante ›", key="b_wz_ir6", type="primary", width="stretch", on_click=ir_a_caso, args=(id_,))
    a, b, c_ = st.columns([1, 1, 1.3])
    a.button("‹ Anterior", key="b_wz_prev", width="stretch", disabled=paso == 0, on_click=wz_mover, args=(-1,))
    b.button("Cerrar" if (creado or dup) else "Cancelar", key="b_wz_cancel", width="stretch", on_click=wz_cancelar)
    if paso < 3:
        c_.button("Siguiente ›", key="b_wz_next", type="primary", width="stretch", disabled=not puede_seguir, on_click=wz_mover, args=(1,))
    elif creado or dup:
        c_.button(f"Caso {creado or dup['id']} ya creado", key="b_wz_ok", width="stretch", disabled=True)
    else:
        c_.button("Confirmar decisión", key="b_wz_ok", type="primary", width="stretch", disabled=not puede_seguir, on_click=wz_confirmar)


def _lista_alertas():
    s = st.session_state
    df_all = df_activo()
    tipo = st.radio("Periodo a revisar", ["Un día", "Una semana", "Un mes", "Todo el periodo"], horizontal=True, key="w_dec_tipo")
    if tipo == "Un día":
        dias = sorted(df_all["Fecha"].unique())
        if "dia_tipico" not in s:
            s.dia_tipico = dia_tipico(df_all, s.umb, leer_par())
        if s.get("w_dec_dia") not in dias:
            s.w_dec_dia = s.dia_tipico
        dia = st.selectbox("Día", dias, key="w_dec_dia", format_func=lambda d: f"{pd.Timestamp(d):%d/%m/%Y} ({DIAS_ES[pd.Timestamp(d).weekday()]})",
                           help="Se abre en un día típico: sin cortes de servicio, con 1–2 alertas y OEE cercano a la mediana.")
        dfp = df_all[df_all["Fecha"] == dia]
    elif tipo == "Una semana":
        sem = (df_all["Fecha"] - pd.to_timedelta(df_all["Fecha"].dt.weekday, unit="D"))
        semanas = sorted(sem.unique())
        if s.get("w_dec_sem") not in semanas:
            s.w_dec_sem = semanas[len(semanas) // 2]
        w = st.selectbox("Semana (inicia el lunes)", semanas, key="w_dec_sem", format_func=lambda d: pd.Timestamp(d).strftime("%d/%m/%Y"))
        dfp = df_all[sem == w]
    elif tipo == "Un mes":
        meses = [int(m) for m in sorted(df_all["Mes"].unique())]
        if s.get("w_dec_mes") not in meses:
            s.w_dec_mes = meses[0]
        m = st.selectbox("Mes", meses, key="w_dec_mes", format_func=mes_lbl)
        dfp = df_all[df_all["Mes"] == m]
    else:
        dfp = df_all
    alertas = alertas_en(dfp, df_all, s.umb, leer_par())
    rol = st.selectbox("Ver alertas para el rol", ["Todos"] + ROLES, key="w_dec_rol",
                       help="Resalta las alertas que le tocan a cada rol según la Tabla 7.")
    if rol != "Todos":
        alertas = [a for a in alertas if rol in a["roles"]]
    if not alertas:
        caja("No hay alertas en este periodo: el operador continúa con el monitoreo y el mantenimiento autónomo programado.", "ok")
        return
    rojas = sum(a["color"] == "rojo" for a in alertas)
    st.markdown(f"**{len(alertas)} alerta(s)** · {rojas} grave(s) · ordenadas de la más grave a la menos grave")
    todas = len(alertas) > 12 and st.checkbox(f"Mostrar las {len(alertas)} alertas (por defecto solo las 12 más graves)", key="w_dec_todas")
    vistas = alertas if (todas or len(alertas) <= 12) else alertas[:12]
    casos_por_alerta = {c["alerta"]["id"]: c["id"] for c in s.casos}
    filas = [[TIPO_EMOJI[a["color"]], f"{pd.Timestamp(a['fecha']):%d/%m/%Y}", EST_CORTO[a["est"]], a["senal"].split(" · ")[0], a["valor"], a["umbral"],
              casos_por_alerta.get(a["id"], "—")] for a in vistas]
    tabla_html(["", "Fecha", "Est.", "Señal", "Valor", "Umbral", "Caso"], filas)
    st.caption("🔴 grave · 🟠 atender · 🟡 informativa (fuera del alcance)")
    por_id = {a["id"]: a for a in vistas}
    ids = list(por_id)
    if s.get("w_dec_alerta") not in ids:
        s.w_dec_alerta = ids[0]
    elegido = por_id[st.selectbox("¿Qué alerta vas a atender?", ids, key="w_dec_alerta",
                                  format_func=lambda i: f"{TIPO_EMOJI[por_id[i]['color']]} "
                                                         f"{pd.Timestamp(por_id[i]['fecha']):%d/%m} · {por_id[i]['senal']} · {por_id[i]['valor']}")]
    abierto = next((c for c in s.casos if c["alerta"]["id"] == elegido["id"] and c["estado"] in ("Abierto", "Cerrado")), None)
    reabierto = next((c for c in s.casos if c["alerta"]["id"] == elegido["id"] and c["estado"] == "Reabierto"), None)
    if abierto:
        caja(f"Esta alerta ya tiene el caso <b>{abierto['id']}</b> ({abierto['estado']}).", "info")
    if reabierto:
        caja(f"Esta alerta tiene el caso <b>{reabierto['id']}</b> reabierto: sigue con la siguiente causa.", "warn")
        st.button(f"Continuar el caso {reabierto['id']} (reabierto) ›", key="b_continuar", type="primary", width="stretch",
                  on_click=cb_reabrir, args=(reabierto["id"],))
    else:
        st.button("Atender esta alerta ›", key="b_atender", type="primary", width="stretch", disabled=abierto is not None,
                  on_click=lambda: st.session_state.update(wz=dict(alerta=elegido, paso=0, causa_ident=None, caso_id=None)))


# ---------------------------------------------------------------- sensores (C1)
def reiniciar_sensores():
    s = st.session_state
    s.d_seed = s.get("d_seed", 7) + 1
    rng = np.random.default_rng(s.d_seed)
    s.d_datos, s.d_fallas = generar_datos(ESC_C1[s.w_d_esc], rng)
    s.d_fallas_act = {e: dict(f) for e, f in s.d_fallas.items()}
    s.d_bandas = {}
    s.w_d_lect = N_LECTURAS


def _fig_variable(v, serie, fuera, i, banda, titulo):
    x = np.arange(len(serie)) * INTERVALO_S / 60.0
    fig = go.Figure()
    if v["kind"] in ("okng", "runstop"):
        malo = "NG" if v["kind"] == "okng" else "STOP"
        bueno = "OK" if v["kind"] == "okng" else "RUN"
        y = (serie == malo).astype(int).values
        fig.add_scatter(x=x[:i + 1], y=y[:i + 1], mode="lines", line=dict(shape="hv", color=INK, width=2), name="Estado")
        fig.add_hrect(y0=-0.2, y1=0.5, fillcolor="#12a150", opacity=0.08, line_width=0)
        fig.update_yaxes(tickvals=[0, 1], ticktext=[bueno, malo], range=[-0.2, 1.3])
        m = fuera.values[:i + 1]
        fig.add_scatter(x=x[:i + 1][m], y=y[:i + 1][m], mode="markers", marker=dict(color=ROJO, size=7), name="Fuera de rango")
    else:
        y = serie.astype(float).values
        if v["kind"] == "range":
            fig.add_hrect(y0=v["lo"], y1=v["hi"], fillcolor=AZUL, opacity=0.10, line_width=0)
            fig.add_hline(y=v["lo"], line_dash="dash", line_color=ROJO, annotation_text=f"mín {v['lo']:g}", annotation_position="bottom left")
        fig.add_hline(y=v["hi"], line_dash="dash", line_color=ROJO, annotation_text=f"máx {v['hi']:g}", annotation_position="top left")
        if banda:
            if v["kind"] == "range" and banda[0] > 0:
                fig.add_hline(y=banda[0], line_dash="dashdot", line_color="#ef6c00")
            fig.add_hline(y=banda[1], line_dash="dashdot", line_color="#ef6c00", annotation_text=f"Banda de alerta C4 ({banda[1]:g})",
                          annotation_position="bottom right")
        fig.add_scatter(x=x[:i + 1], y=y[:i + 1], mode="lines", line=dict(color=INK, width=1.8), name=v["nombre"])
        m = fuera.values[:i + 1]
        fig.add_scatter(x=x[:i + 1][m], y=y[:i + 1][m], mode="markers", marker=dict(color=ROJO, size=6), name="Fuera de rango")
        fig.add_scatter(x=[x[i]], y=[y[i]], mode="markers+text", text=[f"{y[i]:.2f}"], textposition="top center",
                        marker=dict(color=AZUL, size=9), showlegend=False)
        fig.update_yaxes(title=v["unidad"])
    fig.add_vline(x=x[i], line_dash="dot", line_color="#9ca3af")
    fig.update_xaxes(title="Tiempo simulado (min)", range=[0, x[-1]])
    return estilo_fig(fig, 360, titulo)


def _crear_caso_sensor(ev, direccion, causa):
    s = st.session_state
    v = ev["var"]
    t7 = T7["sensor"]
    num = ev["valor_obs"] if not isinstance(ev["valor_obs"], str) else np.nan
    al = dict(id=f"S|{ev['equipo']}|{v['key']}|{ev['lectura']}", tipo="sensor", fecha=ev["marca"], est=int(ev["estacion"][1]),
              causa_parada="", senal=f"Sensor fuera de rango · {ev['equipo']} · {v['nombre']}", valor=formato_valor(v, ev["valor_obs"]),
              umbral=formato_umbral(v, ev["banda"]), valor_num=num, umbral_num=None, gravedad=1.0, color="rojo",
              sensor_ref=(ev["equipo"], v["key"], direccion),
              decision=t7["decision"], resp=t7["resp"], plazo=t7["plazo"], roles=t7["roles"],
              cierre="Variable en rango en las 3 últimas de 12 lecturas")
    c = dict(causa)
    dup = caso_duplicado(al["id"], c)
    if dup:                          # segundo clic o caso repetido: no se crea otro
        st.toast(f"Este caso ya está abierto ({dup['id']})", icon="⚠️")
        return
    c3 = evaluar_c3(c)
    caso = nuevo_caso(al, c, c3, origen="sensor")
    caso["sensor"] = dict(equipo=ev["equipo"], key=v["key"], direccion=direccion, valor_obs=ev["valor_obs"])
    s.casos.append(caso)
    s.msg_sensor = (caso["id"], texto_caso_creado(caso))
    st.toast(s.msg_sensor[1], icon="✅")


def _sensores():
    s = st.session_state
    st.markdown("Esta simulación genera lecturas de los sensores cada 5 s. Una variable entra en alerta cuando sale de su rango "
                f"{LECTURAS_CONSEC} lecturas seguidas. Desde aquí puedes **crear un caso** a partir de una variable en alerta.")
    if "d_datos" not in s:
        reiniciar_sensores()
    a, b = st.columns([3, 1])
    a.selectbox("Escenario simulado", list(ESC_C1), key="w_d_esc", on_change=reiniciar_sensores)
    b.write("")
    b.button("Generar nuevos datos", key="b_d_gen", width="stretch", on_click=reiniciar_sensores)
    datos, bandas = s.d_datos, s.d_bandas
    n_l = len(next(iter(datos.values())))
    i = st.slider("Lectura (1 lectura = 5 s simulados)", 1, n_l, key="w_d_lect") - 1
    res = analizar(datos, bandas)
    oee = calcular_oee(datos)
    en_anom = [e for e in ("E1", "E2", "E3") if any(bool(res[eq]["alarma"].iloc[i]) for eq in VARIABLES if EQUIPOS[eq] == e)]
    vo = float(oee.iloc[i])
    if en_anom and vo >= OEE_REF:
        caja(f"<b>OEE agregado: {vo:.1f} %</b> (parece aceptable) — pero C1 detecta anomalía en {', '.join(en_anom)}. "
             "<b>El OEE agregado no revela dónde se origina el problema.</b>", "warn")
    elif en_anom:
        caja(f"<b>OEE agregado: {vo:.1f} %</b> — anomalía detectada en {', '.join(en_anom)}.", "err")
    else:
        caja(f"<b>OEE agregado: {vo:.1f} %</b> — sin anomalías confirmadas en las 3 estaciones.", "ok")
    cols = st.columns(3)
    for col, est in zip(cols, ("E1", "E2", "E3")):
        eqs = [eq for eq in VARIABLES if EQUIPOS[eq] == est]
        alarma_est = any(bool(res[eq]["alarma"].iloc[i]) for eq in eqs)
        html = (f'<div class="sem"><div class="st"><span class="dotc {"r" if alarma_est else "g"}"></span>{EST_NOM[int(est[1])]}'
                f'<span style="margin-left:auto;font-size:.72rem;color:{ROJO if alarma_est else VERDE}">{"ANOMALÍA" if alarma_est else "Normal"}</span></div>')
        for eq in eqs:
            malas = [v["nombre"] for v in VARIABLES[eq] if bool(res[eq]["sost"][v["key"]].iloc[i])]
            html += (f'<div class="eq"><span class="dotc {"r" if malas else "g"}"></span><b>{eq}</b>'
                     f'<span style="color:#6b7280">{" · " + ", ".join(malas) if malas else ""}</span></div>')
        col.markdown(html + "</div>", unsafe_allow_html=True)
    st.write("")
    c1, c2 = st.columns(2)
    equipo = c1.selectbox("Equipo", list(VARIABLES), key="w_d_eq", format_func=lambda e: f"{e} ({EQUIPOS[e]})")
    nombres = {v["nombre"]: v for v in VARIABLES[equipo]}
    vsel = nombres[c2.selectbox("Variable", list(nombres), key=f"w_d_var_{equipo}")]
    st.plotly_chart(_fig_variable(vsel, datos[equipo][vsel["key"]], res[equipo]["fuera"][vsel["key"]], i, bandas.get((equipo, vsel["key"])),
                                  f"{equipo} — {vsel['nombre']} ({vsel['unidad']})"), width="stretch", key="g_sen_var")
    with st.expander("Ver el OEE agregado de la línea"):
        x = np.arange(len(oee)) * INTERVALO_S / 60.0
        fo = go.Figure()
        fo.add_scatter(x=x[:i + 1], y=oee.values[:i + 1], mode="lines", line=dict(color=AZUL, width=2), name="OEE agregado")
        fo.add_hline(y=OEE_REF, line_dash="dash", line_color=VERDE, annotation_text=f"Referencia {OEE_REF:.0f} %", annotation_position="top left")
        fo.add_scatter(x=[x[i]], y=[oee.values[i]], mode="markers+text", text=[f"{oee.values[i]:.1f} %"], textposition="bottom center",
                       marker=dict(color=INK, size=9), showlegend=False)
        fo.update_yaxes(range=[60, 104], title="OEE (%)")
        fo.update_xaxes(title="Tiempo simulado (min)", range=[0, x[-1]])
        st.plotly_chart(estilo_fig(fo, 330, "OEE agregado de la línea (demostrativo)"), width="stretch", key="g_sen_oee")
    with st.expander("Verdad de la simulación (para verificar el detector)"):
        if s.d_fallas:
            for eq, f in s.d_fallas.items():
                vs = ", ".join(next(v["nombre"] for v in VARIABLES[eq] if v["key"] == k) for k, _ in f["vars"])
                st.write(f"- {eq} ({vs}) desde la lectura {f['inicio']} [t = {mmss(f['inicio'])}]")
        else:
            st.write("Ninguna: operación normal, solo picos aislados de ruido que el filtro ignora.")
        st.caption("Los rangos y valores nominales son parámetros de demostración «(ejemplo)».")
    st.markdown("##### Crear un caso desde un sensor")
    eventos = eventos_c1(datos, res, bandas, i)
    if not eventos:
        caja("Ninguna variable está en alerta en esta lectura. Elige el escenario (b), (c) o (d) y mueve la lectura hacia el final.", "info")
        return
    etiquetas = {f"{e['equipo']} — {e['var']['nombre']} ({e['estacion']})": e for e in eventos}
    ev = etiquetas[st.selectbox("Variable en alerta", list(etiquetas), key="w_d_evento")]
    v = ev["var"]
    direccion, causas = diagnosticar(ev, s.d_ajustes, set())
    for c in causas:
        c.update(equipo=ev["equipo"], vkey=v["key"], vnombre=v["nombre"], direccion=direccion, ident=(ev["equipo"], v["key"], direccion, c["causa"]))
    if not causas:
        caja("No hay causas del AMEF para esta combinación.", "warn")
        return
    _tarjeta_causas([dict(c, descartada=False) for c in causas])
    opciones = {f"{'★ ' if k == 0 else ''}{c['causa']} — NPR {c['npr']}": c for k, c in enumerate(causas)}
    elegida = opciones[st.radio("¿Qué causa vas a revisar?", list(opciones), key=f"w_d_rad_{ev['equipo']}_{v['key']}")]
    r = evaluar_c3(elegida)
    bg, fg = COLOR_PRIO[r["prioridad"]]
    st.markdown(f'Prioridad: <span class="npr" style="background:{bg};color:{fg}">{r["prioridad"].upper()}</span> · mantenimiento '
                f'<b>{r["tipo"].lower()}</b> · {plazo_prio(r["prioridad"], r["tipo"])}', unsafe_allow_html=True)
    dup = caso_duplicado(f"S|{ev['equipo']}|{v['key']}|{ev['lectura']}", elegida)
    if dup:
        ms = s.get("msg_sensor")
        if ms and ms[0] == dup["id"]:
            caja(ms[1], "ok")
        else:
            caja(f"Este caso ya está abierto (<b>{dup['id']}</b>).", "warn")
        st.button(f"Caso {dup['id']} ya creado", key="b_caso_sensor", width="stretch", disabled=True)
        st.button("Ir a 6. Ejecución y resultante ›", key="b_ir6_sensor", type="primary", width="stretch", on_click=ir_a_caso, args=(dup["id"],))
    else:
        st.button("Confirmar decisión (crear caso desde el sensor)", key="b_caso_sensor", type="primary", width="stretch",
                  on_click=_crear_caso_sensor, args=(ev, direccion, elegida))


def pagina_decision():
    s = st.session_state
    abrir_pagina("decision", "Decisión (parte post)",
                 "Ya viste el OEE y el rechazo. Aquí decides qué hacer: elige una alerta y sigue 4 pasos cortos (qué pasó, por qué, qué tan urgente "
                 "y qué decido). Al confirmar se crea un caso que seguirás en el paso 6.")
    msg = s.pop("msg_dec", None)
    if msg:
        caja(msg[1], msg[0])
    t1, t2 = st.tabs(["Alertas de los datos", "Detección por sensores (C1)"])
    with t1:
        if s.get("wz"):
            _wizard()
        else:
            _lista_alertas()
    with t2:
        _sensores()
    if not (s.get("wz") and s.wz.get("paso") == 3):      # en el paso 4 de 4 solo queda «Ir a 6…» / «Confirmar decisión»
        pie_siguiente("decision", tip="Confirma una decisión para pasar a la ejecución.")


# =============================================================================
# PASO 6 - EJECUCION Y RESULTANTE
# =============================================================================
ESTADO_COL = {"Abierto": ("#eff6ff", "#1d4ed8"), "Cerrado": ("#ecfdf3", "#067647"), "Reabierto": ("#fef3f2", "#b42318")}


def cb_verificar():
    s = st.session_state
    caso = next((c for c in s.casos if c["id"] == s.get("w_ej_caso")), None)
    if caso is None:
        return
    resp, accion = str(s.get("w_ej_resp", "")).strip(), str(s.get("w_ej_accion", "")).strip()
    if not resp or not accion:
        s.res_ej = dict(caso_id=caso["id"], error="Escribe quién hizo la acción y qué se hizo antes de verificar.")
        return
    if caso["origen"] == "sensor":
        res = verificar_sensor(caso, s.get("w_ej_efecto", "Efectiva"))
    else:
        res = verificar_caso(caso, df_activo(), s.umb)
    if res["sin_datos"]:
        s.res_ej = dict(caso_id=caso["id"], res=res, lineas=[], estado=caso["estado"])
        return
    caso["intervenciones"].append(dict(fecha=s.get("w_ej_fecha"), responsable=resp, accion=accion, esperado=res["esperado_txt"],
                                       observado=res["observado_txt"], ok=res["ok"]))
    c = caso["causa"]
    if res["ok"]:
        caso["estado"] = "Cerrado"
        if caso["origen"] == "sensor":
            lineas = res["lineas"]
        elif caso["alerta"]["tipo"] == "insumos":
            lineas = ["C1/C2/C3 sin cambios: la causa está fuera del alcance del prototipo."]
        else:
            lineas = retro_kpi(caso, res, s.umb)
    else:
        caso["estado"] = "Reabierto"
        if c.get("ident"):
            caso["descartadas"].append(c["ident"])
        desc = {tuple(x) for x in caso["descartadas"]}
        sig = [x for x in causas_para_alerta(caso["alerta"], s.d_ajustes, desc) if not x["descartada"]] if caso["origen"] != "sensor" else []
        if caso["origen"] == "sensor":
            ev = caso["sensor"]
            dirc = ev["direccion"]
            alt = [x for x in causas_vigentes(ev["equipo"], ev["key"], dirc, s.d_ajustes) if (ev["equipo"], ev["key"], dirc, x["causa"]) not in desc]
            alt.sort(key=lambda x: -x["npr"])
            lineas = ([f"C2: «{c['causa']}» queda descartada; la siguiente por NPR es «{alt[0]['causa']}»."] if alt
                      else ["C2: se revisaron todas las causas del AMEF: escala al jefe de Mantenimiento (análisis de causa raíz)."])
        elif c.get("ident"):
            lineas = ([f"C2: «{c['causa']}» queda descartada; la siguiente por NPR es «{sig[0]['causa']}»."] if sig
                      else ["C2: se revisaron todas las causas del AMEF: escala al jefe de Mantenimiento (análisis de causa raíz)."])
        else:
            lineas = ["La falta de insumos persiste: se escala a Logística (fuera del alcance del prototipo)."]
        if caso["origen"] == "sensor" and res.get("parcial"):
            lineas.insert(0, "La respuesta fue parcial: según la Tabla 7 el caso permanece abierto.")
    caso["retro"] = lineas
    caso["resultado"] = dict(ok=res["ok"], esperado=res["esperado_txt"], observado=res["observado_txt"])
    s.res_ej = dict(caso_id=caso["id"], res=res, lineas=lineas, estado=caso["estado"])
    s.w_ej_accion = ""


def cb_reabrir(caso_id):
    s = st.session_state
    caso = next(c for c in s.casos if c["id"] == caso_id)
    s.wz = dict(alerta=caso["alerta"], paso=1, causa_ident=None, caso_id=caso_id, sensor=caso.get("sensor"))
    s.pagina = "decision"


def _grafico_antes_despues(res):
    if res.get("antes") is None or res.get("despues") is None or not (np.isfinite(res["antes"]) and np.isfinite(res["despues"])):
        st.caption("No hay un valor numérico comparable para dibujar el antes y el después.")
        return
    tu = res["tipo_u"]
    esc = 100 if tu == "pct" else 1
    fig = go.Figure()
    fig.add_bar(x=[envolver(res["etiqueta_antes"], 18), envolver(res["etiqueta_despues"], 18)], y=[res["antes"] * esc, res["despues"] * esc],
                marker_color=[GRIS, VERDE if res["ok"] else ROJO], texttemplate="%{y:.2f}" + (" %" if tu == "pct" else ""),
                textposition="inside", insidetextanchor="end", textfont=dict(color="#fff"), name="Valor")
    if res.get("esperado") is not None and np.isfinite(res["esperado"]):
        fig.add_hline(y=res["esperado"] * esc, line_dash="dash", line_color=INK,
                      annotation_text=f"Esperado {res['cond']} {fmt_u(tu, res['esperado'])}", annotation_position="top left")
    if res.get("rango"):
        fig.add_hrect(y0=res["rango"][0], y1=res["rango"][1], fillcolor=VERDE, opacity=0.10, line_width=0)
    fig.update_yaxes(title={"pct": "%", "h": "horas", "int": "paradas"}.get(tu, res.get("unidad", "")))
    fig.update_layout(showlegend=False)
    st.plotly_chart(estilo_fig(fig, 340, "Antes y después"), width="stretch", key="g_antes_despues")


def cb_anular_ok(caso_id):
    s = st.session_state
    s.casos = [c for c in s.casos if c["id"] != caso_id]
    s.anular_pend = None
    if s.get("res_ej") and s.res_ej.get("caso_id") == caso_id:
        s.pop("res_ej")
    if s.get("wz") and s.wz.get("caso_id") == caso_id:
        s.wz = None
    s.msg_dec = ("info", f"Caso {caso_id} anulado.")
    if not s.casos:
        s.pagina = "decision"


def cb_reiniciar_casos():
    """Borra todos los casos; conserva los datos cargados y los parámetros."""
    s = st.session_state
    s.casos, s.caso_seq, s.wz, s.anular_pend = [], 0, None, None
    for k in ("res_ej", "msg_sensor"):
        s.pop(k, None)
    s.msg_dec = ("info", "Se borraron todos los casos. Los datos y los parámetros se conservaron.")
    if s.pagina == "ejecucion":
        s.pagina = "decision"


def pagina_ejecucion():
    s = st.session_state
    abrir_pagina("ejecucion", "Ejecución y resultante",
                 "Aquí registras qué hiciste y compruebas, con los datos del periodo siguiente, si la acción funcionó. Si funcionó, el caso se cierra y "
                 "el sistema aprende; si no, el caso se reabre con la siguiente causa y vuelves al paso 5.")
    msg = s.pop("msg_dec", None)
    if msg:
        caja(msg[1], msg[0])
    filas = []
    for c in s.casos:
        bg, fg = ESTADO_COL[c["estado"]]
        filas.append([c["id"], c["alerta"]["senal"], c["causa"]["causa"], c["prioridad"], c.get("accion", c["alerta"]["decision"]),
                      c["plazo"], f'<span class="npr" style="background:{bg};color:{fg}">{c["estado"]}</span>'])
    tabla_html(["ID", "Alerta", "Causa", "Prioridad", "Qué hacer", "Plazo", "Estado"], filas)
    with st.expander("Anular un caso creado por error"):
        ids_todos = [c["id"] for c in s.casos]
        if s.get("w_anu_id") not in ids_todos:
            s.w_anu_id = ids_todos[0]
        sel_anu = st.selectbox("Caso a anular", ids_todos, key="w_anu_id",
                               format_func=lambda i: f"{i} · {next(c for c in s.casos if c['id'] == i)['alerta']['senal']} · "
                                                     f"{next(c for c in s.casos if c['id'] == i)['estado']}")
        if s.get("anular_pend") == sel_anu:
            caja(f"¿Seguro que quieres anular el caso <b>{sel_anu}</b>? Se elimina de la lista y no se puede deshacer.", "warn")
            c1, c2 = st.columns(2)
            c1.button(f"Sí, anular {sel_anu}", key="b_anular_si", type="primary", width="stretch", on_click=cb_anular_ok, args=(sel_anu,))
            c2.button("No, conservarlo", key="b_anular_no", width="stretch", on_click=lambda: st.session_state.update(anular_pend=None))
        else:
            st.button("Anular caso", key="b_anular", width="stretch", on_click=lambda: st.session_state.update(anular_pend=sel_anu))
        st.caption("Anular borra el caso; no deshace los ajustes de C1/C2/C3 que ya haya aplicado un caso cerrado.")
    abiertos = [c for c in s.casos if c["estado"] == "Abierto"]
    if not abiertos:
        reab = [c for c in s.casos if c["estado"] == "Reabierto"]
        st.write("")
        if reab:
            caja("No hay casos abiertos para verificar. Los casos reabiertos esperan una nueva decisión en el paso 5.", "info")
        else:
            caja("Todos los casos están cerrados.", "ok")
    else:
        st.write("")
        st.markdown("##### Registra la acción y verifica")
        ids = [c["id"] for c in abiertos]
        if s.get("w_ej_caso") not in ids:
            s.w_ej_caso = ids[0]
        caso = next(c for c in abiertos if c["id"] == st.selectbox("Caso a verificar", ids, key="w_ej_caso",
                                                                    format_func=lambda i: f"{i} · {next(c for c in abiertos if c['id'] == i)['alerta']['senal']}"))
        ref = f' <span style="color:#6b7280;font-size:.85rem">{caso["plazo_ref"]}</span>' if caso.get("plazo_ref") else ""
        st.markdown(f'<div class="decision" style="padding:12px 16px"><div class="m"><b>Qué hacer:</b> {caso.get("accion", caso["alerta"]["decision"])}<br>'
                    f'<b>Prioridad:</b> {caso["prioridad"]} · <b>Responsable:</b> {caso["responsable"]}<br>'
                    f'<b>Plazo:</b> {caso["plazo"]}{ref}</div></div>', unsafe_allow_html=True)
        a, b = st.columns(2)
        a.date_input("Fecha de la acción", key="w_ej_fecha")
        b.text_input("Responsable", key="w_ej_resp", placeholder="Nombre de quien hizo la acción")
        st.text_input("Acción realizada", key="w_ej_accion", placeholder="Ej.: cambio del empaque de la tapa y ajuste del torque")
        if caso["origen"] == "sensor":
            st.radio("Efecto simulado de la intervención (selector de prueba)", EFECTOS, key="w_ej_efecto", horizontal=True,
                     help="Este caso viene de sensores: define cómo se comportan las 12 lecturas posteriores (efectiva, parcial o no efectiva).")
        else:
            st.caption("Se compara con los datos del periodo siguiente a la alerta, según la Tabla 7. No hace falta simular nada.")
        st.button("Verificar resultante", key="b_verificar", type="primary", width="stretch", on_click=cb_verificar)
    reab = [c for c in s.casos if c["estado"] == "Reabierto"]
    if reab:
        st.write("")
        st.markdown("##### Casos reabiertos")
        ids_r = [c["id"] for c in reab]
        if s.get("w_ej_reab") not in ids_r:
            s.w_ej_reab = ids_r[0]
        sel = st.selectbox("Caso reabierto", ids_r, key="w_ej_reab",
                           format_func=lambda i: f"{i} · {next(c for c in reab if c['id'] == i)['alerta']['senal']}")
        st.button("Volver al paso 5 con este caso ›", key="b_reabrir_sel", width="stretch", on_click=cb_reabrir, args=(sel,))
    r = s.get("res_ej")
    if r:
        caso = next((c for c in s.casos if c["id"] == r["caso_id"]), None)
        st.divider()
        if r.get("error"):
            caja(r["error"], "err")
        elif r["res"]["sin_datos"]:
            caja(r["res"]["mensaje"], "warn")
        elif caso:
            res = r["res"]
            st.markdown(f"##### Resultado del caso {caso['id']}")
            a, b = st.columns(2)
            a.markdown(f'<div class="ev"><div class="lab">Esperado</div><div class="big">{res["esperado_txt"]}</div></div>', unsafe_allow_html=True)
            b.markdown(f'<div class="ev"><div class="lab">Observado ({res["ventana"]})</div><div class="big">{res["observado_txt"]}</div></div>',
                       unsafe_allow_html=True)
            if res["ok"]:
                st.markdown('<div class="resultado ok">CERRADO ✓</div>', unsafe_allow_html=True)
            else:
                st.markdown('<div class="resultado bad">REABIERTO ✗</div>', unsafe_allow_html=True)
            _grafico_antes_despues(res)
            if r["lineas"]:
                st.markdown("**Qué cambió en el sistema:**" if res["ok"] else "**Qué pasa ahora:**")
                for ln in r["lineas"]:
                    st.write("• " + ln)
            if not res["ok"]:
                st.button("Volver al paso 5 con este caso ›", key="b_reabrir", type="primary", width="stretch",
                          on_click=cb_reabrir, args=(caso["id"],))
    pie_siguiente("ejecucion")



# =============================================================================
# PASO 7 - VALIDACION (MONTE CARLO)
# =============================================================================
def kpis_base_filtrado(df, est, mes, causa):
    """Línea base real con los mismos filtros (misma fórmula que la simulación)."""
    fb = filtrar_fact(fact_base(df), est, mes, causa)
    if fb.empty:
        return {}
    L, _ = kpis_por_iteracion(fb)
    r = L.iloc[0]
    return {k: (float(r[k]) if k in r and pd.notna(r[k]) else np.nan) for k in
            ("OEE", "D", "R", "Q", "Pa", "ev", "MTBF", "MTTR", "rechazo", "conformes")}


def html_progreso(escenarios, estado):
    filas = []
    for sc in escenarios:
        hecho, total = estado.get(sc, (0, 1))
        if hecho >= total:
            badge = '<span class="badge c">Completado</span>'
        elif hecho > 0:
            badge = '<span class="badge e">En ejecución</span>'
        else:
            badge = '<span class="badge p">Pendiente</span>'
        pct = 100 * hecho / max(total, 1)
        filas.append(f'<div class="prog-row"><div><b>{ESCENARIOS_SIM[sc]}</b></div><div>{badge}</div>'
                     f'<div>{hecho} / {total}</div><div class="bar"><div style="width:{pct:.1f}%"></div></div></div>')
    return "".join(filas)


def _val_config():
    s = st.session_state
    st.markdown("Define cuántas veces repetir el experimento y qué escenarios comparar. Solo las causas de equipo y la espera de proceso cambian "
                "con el DSS; insumos, personal y cortes de servicio quedan sin cambio (fuera del alcance).")
    a, b = st.columns(2, gap="large")
    a.number_input("Iteraciones Monte Carlo", min_value=5, max_value=1000, step=5, key="w_cfg_iter",
                   help="Cada iteración remuestrea (con reemplazo) los días de cada mes de tus datos. Más iteraciones = resultado más estable.")
    b.number_input("Semilla de reproducibilidad", min_value=0, max_value=10_000_000, step=1, key="w_cfg_seed",
                   help="Con la misma semilla y los mismos parámetros se obtienen los mismos resultados.")
    st.multiselect("Escenarios a comparar", ESCENARIOS_SIM, key="w_cfg_scen",
                   help="Sin DSS reproduce la línea base; DSS parcial activa solo C1 (detección); DSS completo activa C1–C4.")
    P = leer_supuestos()
    with st.expander("Ver los supuestos que usa la simulación"):
        tabla_html(["Supuesto", "DSS parcial", "DSS completo"], [
            ["Detección anticipada C1 (probabilidad)", n(P["s_p_det"]), n(P["s_p_det"])],
            ["Reducción de duración al detectar a tiempo", n(P["s_r_dur_par"]), n(P["s_r_dur_com"])],
            ["Tiempo de respuesta medio (min)", n(P["s_t_par"], 1), n(P["s_t_com"], 1)],
            ["Mejora de rendimiento", n(P["s_mej_r_par"]), n(P["s_mej_r_com"])],
            ["Reducción de no conformes", n(P["s_red_def_par"]), n(P["s_red_def_com"])],
            ["Eficacia máxima de C4 / aprendizaje τ", "—", f"{n(P['s_eff_c4'])} / {n(P['s_tau_c4'], 1)}"],
            ["Propagación a la espera de proceso", n(P["s_propag"]), n(P["s_propag"])]])
        st.caption("Son supuestos de simulación, no datos de la base ni de la tesis. Se editan en el paso 3 (pestaña «Supuestos Monte Carlo»).")
        st.button("Editar supuestos (paso 3)", key="b_edit_sup", on_click=ir, args=("params",))


def _val_sim():
    s = st.session_state
    escs = sorted(ESCENARIOS_SIM.index(e) for e in s.w_cfg_scen)
    if not escs:
        caja("Elige al menos un escenario en la pestaña Configuración.", "err")
        return
    n_it, semilla = int(s.w_cfg_iter), int(s.w_cfg_seed)
    caja(f"Escenarios: <b>{len(escs)}</b> | Iteraciones por escenario: <b>{n_it}</b> | Total de corridas: <b>{len(escs) * n_it}</b> | "
         f"Semilla: <b>{semilla}</b>", "info")
    st.markdown("##### Progreso por escenario")
    ph = st.empty()
    res = s.get("v_res")
    ph.markdown(html_progreso(escs, {sc: (res["n_it"], res["n_it"]) for sc in res["escenarios"]} if res else {}), unsafe_allow_html=True)
    if st.button("Ejecutar experimento", key="b_run_mc", type="primary", width="stretch"):
        estado = {sc: (0, n_it) for sc in escs}

        def progreso(sc, hecho, total):
            estado[sc] = (hecho, total)
            ph.markdown(html_progreso(escs, estado), unsafe_allow_html=True)
        s.v_res = ejecutar_experimento(df_activo(), leer_supuestos(), n_it, semilla, escs, progreso)
        s.pop("v_xlsx", None)
        s.pop("xlsx_total", None)
        caja("Experimento completado. Ve a la pestaña «3. Resultados».", "ok")
        st.button("Ver resultados ›", key="b_ver_res", on_click=lambda: st.session_state.update(w_v_sub="3. Resultados"))
    st.caption("Comparación pareada: dentro de cada iteración los escenarios usan los mismos días muestreados.")


def _filtros(res, df):
    st.markdown('<div class="eyebrow">Filtros · actualizan tarjetas, gráficos y tablas</div>', unsafe_allow_html=True)
    causas = sorted(res["fact"]["causa"].unique())
    meses = [int(m) for m in sorted(res["fact"]["mes"].unique())]
    c1, c2, c3, c4 = st.columns(4)
    opciones_esc = [ESCENARIOS_SIM[sc] for sc in res["escenarios"]]
    if st.session_state.get("w_f_esc") not in opciones_esc:
        st.session_state.w_f_esc = opciones_esc[-1]
    esc = c1.selectbox("Escenario", opciones_esc, key="w_f_esc")
    est = c2.selectbox("Estación", ["Todas", "E1", "E2", "E3"], key="w_f_est", format_func=lambda e: e if e == "Todas" else EST_NOM[int(e[1])])
    if st.session_state.get("w_f_mes") not in ["Todos"] + meses:
        st.session_state.w_f_mes = "Todos"
    mes = c3.selectbox("Mes", ["Todos"] + meses, key="w_f_mes", format_func=lambda m: "Todos" if m == "Todos" else mes_lbl(m))
    if st.session_state.get("w_f_causa") not in ["Todas"] + causas:
        st.session_state.w_f_causa = "Todas"
    causa = c4.selectbox("Causa", ["Todas"] + causas, key="w_f_causa")
    return ESCENARIOS_SIM.index(esc), est, mes, causa


def _tabla_comparacion(L, base, escenarios):
    claves = ["OEE", "D", "R", "Q", "rechazo", "Pa", "MTBF", "MTTR", "t_resp"]
    S = resumen_escenarios(L, claves)
    filas, filas_pct = [], []
    for k in claves:
        fila = [ETIQ[k] if k != "Pa" else "Parada no planificada", META_TXT[k], fmt(TIPO[k], base.get(k)) if k != "t_resp" else "No registrado"]
        fila_p = [fila[0], META_TXT[k]]
        for sc in escenarios:
            r = S[(S.scen == sc) & (S.clave == k)]
            if r.empty:
                fila.append("—")
                fila_p.append("—")
                continue
            m = float(r.media.iloc[0])
            fila.append(f"{fmt(TIPO[k], m)} {'✓' if cumple(k, m) else '✗'}")
            fila_p.append(f"{n(float(r.pct_meta.iloc[0]), 0)} %")
        filas.append(fila)
        filas_pct.append(fila_p)
    return (["Indicador", "Meta (Tabla 2)", "Línea base"] + [ESCENARIOS_SIM[sc] for sc in escenarios], filas,
            ["Indicador", "Meta (Tabla 2)"] + [ESCENARIOS_SIM[sc] for sc in escenarios], filas_pct)


def _val_resultados():
    s = st.session_state
    res, df = s.v_res, df_activo()
    st.caption("Los valores son la media de las iteraciones; los rangos son los percentiles 5 y 95.")
    scen_sel, est, mes, causa = _filtros(res, df)
    escs = res["escenarios"]
    filtrado = (est, mes, causa) != ("Todas", "Todos", "Todas")
    fact = filtrar_fact(res["fact"], est, mes, causa)
    if fact.empty:
        caja("Sin datos para esa combinación de filtros.", "warn")
        return
    L, E = kpis_por_iteracion(fact)
    base = kpis_base_filtrado(df, est, mes, causa)
    claves = [k for k, _, _ in INDICADORES]
    S = resumen_escenarios(L, claves)
    tabs = st.tabs(["Indicadores", "Comparación de escenarios", "Evolución mensual del OEE", "Desempeño por estación",
                    "Variabilidad", "Paradas por causa", "Detalle operativo"])
    with tabs[0]:
        st.caption(f"Escenario representativo: **{ESCENARIOS_SIM[scen_sel]}** (cámbialo con el filtro Escenario).")
        if scen_sel == 0:
            r = S[(S.scen == 0) & (S.clave == "OEE")]
            if not r.empty and base.get("OEE") is not None:
                caja(f"El escenario sin DSS reproduce la línea base: OEE simulado <b>{fmt('pct', float(r.media.iloc[0]))}</b> frente a "
                     f"<b>{fmt('pct', base['OEE'])}</b> de la base.", "info")
        principales = [("OEE", "OEE de línea"), ("Pa", "Horas de parada"), ("rechazo", "Índice de rechazo"), ("t_resp", "Tiempo de respuesta")]
        cols = st.columns(4)
        for col, (k, e) in zip(cols, principales):
            r = S[(S.scen == scen_sel) & (S.clave == k)]
            if r.empty:
                kpi(col, e, "—")
                continue
            m = float(r.media.iloc[0])
            sub = f"Meta {META_TXT[k]} · p5–p95: {fmt(TIPO[k], float(r.p5.iloc[0]))} – {fmt(TIPO[k], float(r.p95.iloc[0]))}"
            kpi(col, e, fmt(TIPO[k], m), sub, "ok" if cumple(k, m) else "bad")
        r = S[(S.scen == scen_sel) & (S.clave == "OEE")]
        r2 = S[(S.scen == scen_sel) & (S.clave == "t_resp")]
        if not r.empty:
            st.caption(f"Iteraciones con OEE ≥ meta: **{n(float(r.pct_meta.iloc[0]), 0)} %**"
                       + (f" · Iteraciones con tiempo de respuesta ≤ meta: **{n(float(r2.pct_meta.iloc[0]), 0)} %**" if not r2.empty else ""))
        with st.expander("Ver todos los indicadores"):
            filas = []
            for k, e, t in INDICADORES:
                rr = S[(S.scen == scen_sel) & (S.clave == k)]
                if rr.empty:
                    continue
                m = float(rr.media.iloc[0])
                c = cumple(k, m)
                filas.append([e, fmt(t, m), f"{fmt(t, float(rr.p5.iloc[0]))} – {fmt(t, float(rr.p95.iloc[0]))}", META_TXT.get(k, "—"),
                              ("✓" if c else "✗") if c is not None else "—",
                              f"{n(float(rr.pct_meta.iloc[0]), 0)} %" if k in METAS else "—"])
            tabla_html(["Indicador", "Media", "Rango p5–p95", "Meta", "Cumple", "% de iteraciones que cumplen"], filas)
    with tabs[1]:
        if filtrado:
            caja("Las metas de la Tabla 2 aplican a la línea completa; con filtros activos la comparación es referencial.", "warn")
        r = S[S.clave == "OEE"].sort_values("scen")
        fig = go.Figure(go.Bar(
            x=[envolver(ESCENARIOS_SIM[sc], 16) for sc in r.scen], y=(r.media * 100).round(2), marker_color=[COLOR_ESC[sc] for sc in r.scen],
            texttemplate="%{y:.2f} %", textposition="inside", insidetextanchor="middle", textfont=dict(color="#fff"), name="OEE (media, con rango p5–p95)",
            error_y=dict(type="data", symmetric=False, array=((r.p95 - r.media) * 100).tolist(), arrayminus=((r.media - r.p5) * 100).tolist(), color="#6b7280")))
        cats = [envolver(ESCENARIOS_SIM[sc], 16) for sc in r.scen]
        if base.get("OEE") is not None and np.isfinite(base["OEE"]):
            fig.add_scatter(x=cats, y=[base["OEE"] * 100] * len(cats), mode="lines", name=f"Línea base {n(base['OEE'] * 100)} %",
                            line=dict(color=INK, dash="dot", width=2))
        fig.add_scatter(x=cats, y=[METAS["OEE"][1] * 100] * len(cats), mode="lines", name=f"Meta {META_TXT['OEE']}",
                        line=dict(color=VERDE, dash="dash", width=2))
        fig.update_yaxes(title="OEE (%)", range=[0, max(75, float(r.p95.max() * 100) + 10)])
        st.plotly_chart(estilo_fig(fig, 400, "OEE por escenario (media y rango p5–p95)"), width="stretch", key="g_cmp_oee")
        cab, filas, cab2, filas_p = _tabla_comparacion(L, base, escs)
        st.markdown("**Comparación contra las metas de la Tabla 2** (✓ cumple, ✗ no cumple)")
        tabla_html(cab, filas)
        with st.expander("Porcentaje de iteraciones que cumplen la meta"):
            tabla_html(cab2, filas_p)
    with tabs[2]:
        fm = fact.groupby(["it", "scen", "mes", "est"], as_index=False)[COLS_SUMA].sum()
        fm["OEE"] = (fm.To / fm.Tp.replace(0, np.nan)) * (fm.Proc / fm.Teor.replace(0, np.nan)) * (fm.Conf / fm.Proc.replace(0, np.nan))
        lm = fm.groupby(["it", "scen", "mes"], as_index=False)["OEE"].mean()
        fig = go.Figure()
        for sc in escs:
            g = lm[lm.scen == sc].groupby("mes")["OEE"]
            mu, p5, p95 = g.mean() * 100, g.quantile(0.05) * 100, g.quantile(0.95) * 100
            x = [mes_lbl(m) for m in mu.index]
            col = COLOR_ESC[sc]
            rgb = tuple(int(col[i:i + 2], 16) for i in (1, 3, 5))
            fig.add_scatter(x=x, y=p95.round(2), mode="lines", line=dict(width=0), showlegend=False, hoverinfo="skip")
            fig.add_scatter(x=x, y=p5.round(2), mode="lines", line=dict(width=0), fill="tonexty", fillcolor=f"rgba({rgb[0]},{rgb[1]},{rgb[2]},0.15)",
                            showlegend=False, hoverinfo="skip")
            fig.add_scatter(x=x, y=mu.round(2), mode="lines+markers+text", name=ESCENARIOS_SIM[sc], line=dict(color=col, width=2.4),
                            marker=dict(size=7), text=[f"{v:.1f}" for v in mu], textposition="top center", textfont=dict(size=10, color=col))
        fig.update_yaxes(title="OEE (%)")
        st.plotly_chart(estilo_fig(fig, 430, "Evolución mensual del OEE (media y banda p5–p95)"), width="stretch", key="g_mensual")
        piv = lm.groupby(["scen", "mes"])["OEE"].mean().unstack(0) * 100
        piv.index = [mes_lbl(m) for m in piv.index]
        piv.columns = [ESCENARIOS_SIM[sc] for sc in piv.columns]
        st.dataframe(piv.round(2), width="stretch")
    with tabs[3]:
        ee = E.groupby(["scen", "est"], as_index=False)[["D", "R", "Q", "OEE", "Pa", "Proc", "NC", "ev"]].mean()
        fig = go.Figure()
        for sc in escs:
            g = ee[ee.scen == sc]
            fig.add_bar(x=[envolver(EST_NOM[e], 16) for e in g.est], y=(g.OEE * 100).round(2), name=ESCENARIOS_SIM[sc], marker_color=COLOR_ESC[sc],
                        texttemplate="%{y:.1f} %", textposition="outside")
        fig.update_yaxes(title="OEE (%)", range=[0, 85])
        st.plotly_chart(estilo_fig(fig, 400, "OEE por estación y escenario"), width="stretch", key="g_est")
        t = ee.copy()
        t["Escenario"] = t.scen.map(lambda sc: ESCENARIOS_SIM[sc])
        t["Estación"] = t.est.map(EST_NOM)
        for c in ("D", "R", "Q", "OEE"):
            t[c] = (t[c] * 100).map(lambda x: f"{n(x)} %")
        t["Rechazo"] = (ee.NC / ee.Proc * 100).map(lambda x: f"{n(x)} %")
        t["Pa"] = ee.Pa.map(lambda x: f"{n(x)} h")
        t["ev"] = ee.ev.map(lambda x: n(x, 1))
        t = t[["Escenario", "Estación", "D", "R", "Q", "OEE", "Pa", "ev", "Rechazo"]]
        t.columns = ["Escenario", "Estación", "Disponibilidad", "Rendimiento", "Calidad", "OEE", "Horas de parada", "Eventos", "Rechazo (sobre procesados)"]
        st.dataframe(t, hide_index=True, width="stretch")
    with tabs[4]:
        fig = go.Figure()
        for sc in escs:
            y = (L[L.scen == sc]["OEE"] * 100).round(3)
            nombre = envolver(ESCENARIOS_SIM[sc], 16)
            fig.add_box(y=y, name=nombre, marker_color=COLOR_ESC[sc], boxpoints="all", jitter=0.4, pointpos=0, boxmean=True,
                        line=dict(width=1.6), marker=dict(size=4, opacity=0.55), showlegend=False)
        fig.update_yaxes(title="OEE (%)")
        st.plotly_chart(estilo_fig(fig, 420, "Distribución del OEE entre iteraciones (la línea punteada es la media)"), width="stretch", key="g_box")
        filas = []
        for sc in escs:
            g = L[L.scen == sc]
            fila = [ESCENARIOS_SIM[sc]]
            for k in ("OEE", "D", "R", "Q", "Pa", "conformes"):
                x = g[k].dropna()
                cv = x.std(ddof=1) / x.mean() * 100 if len(x) > 1 and x.mean() else np.nan
                fila.append(f"{n(cv, 2)} %")
            filas.append(fila)
        st.markdown("**Coeficiente de variación entre iteraciones** (desviación estándar / media; más bajo = más estable)")
        tabla_html(["Escenario", "OEE", "Disponibilidad", "Rendimiento", "Calidad", "Horas de parada", "Bidones conformes"], filas)
    with tabs[5]:
        pc = fact[fact.causa != SIN_PARADA].groupby(["scen", "it", "causa"], as_index=False)["Pa"].sum()
        pm = (pc.groupby(["scen", "causa"])["Pa"].sum() / res["n_it"]).reset_index()
        if pm.empty:
            caja("Sin paradas para esos filtros.", "info")
        else:
            ref = scen_sel if scen_sel in escs else escs[0]
            orden = (pm[pm.scen == ref].sort_values("Pa", ascending=False)["causa"].tolist()
                     + [c for c in pm.causa.unique() if c not in pm[pm.scen == ref].causa.tolist()])
            fig = go.Figure()
            for sc in escs:
                g = pm[pm.scen == sc].set_index("causa").reindex(orden)
                fig.add_bar(x=[envolver(c, 12) for c in orden], y=g["Pa"].fillna(0).round(2), name=ESCENARIOS_SIM[sc], marker_color=COLOR_ESC[sc],
                            texttemplate="%{y:.1f}", textposition="outside")
            fig.update_yaxes(title="Horas de parada (media por iteración)")
            st.plotly_chart(estilo_fig(fig, 470, "Paradas por causa y escenario"), width="stretch", key="g_causas")
            with st.expander(f"Pareto · {ESCENARIOS_SIM[scen_sel]}"):
                g = pm[pm.scen == scen_sel].sort_values("Pa", ascending=False)
                if not g.empty:
                    g = g.assign(pct=g.Pa / g.Pa.sum())
                    fp = make_subplots(specs=[[{"secondary_y": True}]])
                    xx = [envolver(c, 12) for c in g.causa]
                    fp.add_bar(x=xx, y=g.Pa.round(2), marker_color=COLOR_ESC[scen_sel], name="Horas", texttemplate="%{y:.1f} h",
                               textposition="outside", secondary_y=False)
                    fp.add_scatter(x=xx, y=(g.pct.cumsum() * 100).round(1), mode="lines+markers", name="% acumulado", line=dict(color=INK), secondary_y=True)
                    fp.update_yaxes(range=[0, g.Pa.max() * 1.25], secondary_y=False, title="Horas")
                    fp.update_yaxes(range=[0, 105], secondary_y=True, title="% acumulado", tickvals=[0, 25, 50, 75, 100],
                                    ticktext=["0 %", "25 %", "50 %", "75 %", "100 %"], showgrid=False)
                    st.plotly_chart(estilo_fig(fp, 450, f"Pareto · {ESCENARIOS_SIM[scen_sel]}"), width="stretch", key="g_pareto_sel")
            resumen = pm.pivot(index="causa", columns="scen", values="Pa").fillna(0)
            resumen.columns = [ESCENARIOS_SIM[sc] for sc in resumen.columns]
            st.dataframe(resumen.round(2).sort_values(resumen.columns[0], ascending=False), width="stretch")
    with tabs[6]:
        d = res["detalle"]
        d = d[d.scen == scen_sel]
        if est != "Todas":
            d = d[d["Estación"] == est]
        if mes != "Todos":
            d = d[d.mes == int(mes)]
        if causa != "Todas":
            d = d[d.Causa_DSS == causa]
        st.caption("Detalle de la iteración 1 (días remuestreados) del escenario elegido: parada y producción antes y después del efecto del DSS.")
        mostrar = d.drop(columns=["scen", "mes"]).copy()
        mostrar["Fecha"] = pd.to_datetime(mostrar["Fecha"]).dt.strftime("%Y-%m-%d")
        for c in ("Parada_DSS_h", "Operativo_DSS_h"):
            mostrar[c] = mostrar[c].round(2)
        st.dataframe(mostrar, hide_index=True, width="stretch", height=430)
        tot = d[["Parada_base_h", "Parada_DSS_h", "Conformes"]].sum()
        c = st.columns(3)
        kpi(c[0], "Parada base (h)", n(tot.Parada_base_h), "filas mostradas")
        kpi(c[1], "Parada con DSS (h)", n(tot.Parada_DSS_h), "filas mostradas")
        kpi(c[2], "Bidones conformes", n(tot.Conformes, 0), "filas mostradas")


def pagina_valid():
    s = st.session_state
    abrir_pagina("valid", "Validación (Monte Carlo)",
                 "Esto responde: <b>¿cuánto mejoraría la línea con el DSS?</b> El experimento repite miles de veces los días de tus datos, una vez sin DSS "
                 "y otra con DSS, y compara los resultados contra las metas. Es opcional: no hace falta para decidir sobre las alertas.")
    sub = st.radio("Sección", ["1. Configuración", "2. Simulación", "3. Resultados"], horizontal=True, key="w_v_sub", label_visibility="collapsed")
    if sub.startswith("1"):
        _val_config()
    elif sub.startswith("2"):
        _val_sim()
    elif s.get("v_res") is None:
        caja("Todavía no hay resultados. Ejecuta el experimento en la sección «2. Simulación».", "info")
    else:
        _val_resultados()
    pie_siguiente("valid")


# =============================================================================
# PASO 8 - EXPORTACION
# =============================================================================
def frames_mc(res, df):
    L, E = kpis_por_iteracion(res["fact"])
    claves = [k for k, _, _ in INDICADORES]
    S = resumen_escenarios(L, claves)
    S["Escenario"] = S.scen.map(lambda sc: ESCENARIOS_SIM[sc])
    S["Indicador"] = S.clave.map(ETIQ)
    S["Meta"] = S.clave.map(lambda k: META_TXT.get(k, ""))
    resumen = S[["Escenario", "Indicador", "Meta", "media", "sd", "p5", "p95", "pct_meta"]]
    resumen.columns = ["Escenario", "Indicador", "Meta (Tabla 2)", "Media", "Desv. estándar", "p5", "p95", "% iteraciones que cumplen"]
    est = E.groupby(["scen", "est"], as_index=False)[["D", "R", "Q", "OEE", "Pa", "ev", "Proc", "NC"]].mean()
    est["Escenario"] = est.scen.map(lambda sc: ESCENARIOS_SIM[sc])
    est["Estación"] = est.est.map(EST_NOM)
    est["Rechazo"] = est.NC / est.Proc
    est = est[["Escenario", "Estación", "D", "R", "Q", "OEE", "Pa", "ev", "Rechazo"]]
    est.columns = ["Escenario", "Estación", "Disponibilidad", "Rendimiento", "Calidad", "OEE", "Horas de parada", "Eventos", "Índice de rechazo"]
    fm = res["fact"].groupby(["it", "scen", "mes", "est"], as_index=False)[COLS_SUMA].sum()
    fm["OEE"] = (fm.To / fm.Tp.replace(0, np.nan)) * (fm.Proc / fm.Teor.replace(0, np.nan)) * (fm.Conf / fm.Proc.replace(0, np.nan))
    lm = fm.groupby(["it", "scen", "mes"], as_index=False)["OEE"].mean()
    men = lm.groupby(["scen", "mes"])["OEE"].agg(media="mean", p5=lambda x: x.quantile(0.05), p95=lambda x: x.quantile(0.95)).reset_index()
    men["Escenario"] = men.scen.map(lambda sc: ESCENARIOS_SIM[sc])
    men["Mes"] = men.mes.map(mes_lbl)
    men = men[["Escenario", "Mes", "media", "p5", "p95"]]
    men.columns = ["Escenario", "Mes", "OEE media", "OEE p5", "OEE p95"]
    it = L.copy()
    it["Escenario"] = it.scen.map(lambda sc: ESCENARIOS_SIM[sc])
    it = it[["Escenario", "it"] + claves].rename(columns={"it": "Iteración", **ETIQ})
    it["Iteración"] += 1
    return dict(resumen=resumen, estaciones=est, mensual=men, iteraciones=it)


def df_kpis_metas(df):
    k = kpis_base(df)
    filas = []
    for key, (v0, u, cond) in META_UI.items():
        nombre = KPI_INFO[key][0]
        val = k.get(key)
        meta = METAS[key][1] * (100 if u == "%" else 1)
        if key == "t_resp":
            filas.append([nombre, "No registrado", cond, meta, u, "—"])
        else:
            c = cumple(key, val)
            filas.append([nombre, round(val * (100 if u == "%" else 1), 4), cond, meta, u, "Sí" if c else "No"])
    return pd.DataFrame(filas, columns=["KPI", "Línea base", "Condición", "Meta", "Unidad", "Cumple la meta"])


def df_alertas(df):
    s = st.session_state
    al = alertas_en(df, df, s.umb, leer_par())
    return pd.DataFrame([dict(ID=a["id"], Fecha=pd.Timestamp(a["fecha"]).strftime("%Y-%m-%d"), Estación=EST_NOM[a["est"]], Tipo=T7[a["tipo"]]["nombre"],
                              Señal=a["senal"], Valor=a["valor"], Umbral=a["umbral"], Gravedad=a["color"], Decisión=a["decision"],
                              Responsable=a["resp"], Plazo=a["plazo"], Resultante_esperada=a["cierre"]) for a in al])


def df_casos():
    s = st.session_state
    filas = []
    for c in s.casos:
        a, ca, r = c["alerta"], c["causa"], c.get("resultado") or {}
        filas.append(dict(ID=c["id"], Origen=c["origen"], Fecha_alerta=pd.Timestamp(a["fecha"]).strftime("%Y-%m-%d"), Estación=EST_NOM[a["est"]],
                          Alerta=a["senal"], Valor=a["valor"], Umbral=a["umbral"], Causa=ca["causa"], Equipo=ca.get("equipo", "—"),
                          S=ca.get("S"), O=ca.get("O"), D=ca.get("D"), NPR=ca.get("npr"), Prioridad=c["prioridad"], Tipo_TPM=c["tipo_tpm"],
                          Responsable=c["responsable"], Qué_hacer=c.get("accion", a["decision"]), Plazo=c["plazo"], Estado=c["estado"], Esperado=r.get("esperado", ""),
                          Observado=r.get("observado", ""), Retroalimentación=" | ".join(c.get("retro", [])),
                          Acciones=" | ".join(f"{i['fecha']}: {i['responsable']} — {i['accion']}" for i in c["intervenciones"])))
    return pd.DataFrame(filas)


def df_amef():
    s = st.session_state
    filas = []
    for (eq, vkey, direc), cs in AMEF.items():
        v = next(x for x in VARIABLES[eq] if x["key"] == vkey)
        for c in causas_vigentes(eq, vkey, direc, s.d_ajustes):
            filas.append(dict(Equipo=eq, Variable=v["nombre"], Dirección=direc, Causa=c["causa"], Efecto=c["efecto"], S=c["S"], O=c["O"], D=c["D"],
                              NPR=c["npr"], Origen="Tesis" if not c["ejemplo"] else "(ejemplo)", Ajustada_por_C4="Sí" if c["ajustada"] else "No"))
    return pd.DataFrame(filas)


def df_parametros():
    s = st.session_state
    filas = []
    for e, u in sorted(s.umb.items()):
        filas.append(("Umbrales de alerta", f"{EST_CORTO[e]} · OEE bajo (P10) %", round(u["oee_p10"] * 100, 3)))
        filas.append(("Umbrales de alerta", f"{EST_CORTO[e]} · OEE de cierre (mediana) %", round(u["oee_med"] * 100, 3)))
        filas.append(("Umbrales de alerta", f"{EST_CORTO[e]} · Rechazo alto (P90) %", round(u["rech_p90"] * 100, 3)))
        filas.append(("Umbrales de alerta", f"{EST_CORTO[e]} · Rechazo de cierre (promedio) %", round(u["rech_med"] * 100, 3)))
    par = leer_par()
    filas += [("Reglas de alerta", "Parada no planificada desde (min)", par["parada_min"]), ("Reglas de alerta", "Repetición: n eventos", par["rec_n"]),
              ("Reglas de alerta", "Repetición: ventana (días)", par["rec_ventana"]),
              ("Detección C1", "Lecturas consecutivas", LECTURAS_CONSEC), ("Escalas S/O/D", "La escala «Media» empieza en", UMBRAL_MEDIO),
              ("Escalas S/O/D", "La escala «Alta» empieza en", UMBRAL_ALTO)]
    for eq, vs in VARIABLES.items():
        for v in vs:
            if v["kind"] in ("range", "max"):
                filas.append(("Rangos C1", f"{eq} · {v['nombre']} ({v['unidad']})", f"{v['lo']:g} – {v['hi']:g}"))
    for k, (v0, u, cond) in META_UI.items():
        filas.append(("Metas", KPI_INFO[k][0], META_TXT[k]))
    for k, v in leer_supuestos().items():
        filas.append(("Supuestos Monte Carlo (supuesto, no es dato)", k, v))
    return pd.DataFrame(filas, columns=["Grupo", "Parámetro", "Valor"])


def construir_excel_total():
    s = st.session_state
    d = s.datos
    df = d["df"]
    g = tabla_estaciones(df)
    t1 = pd.DataFrame({EST_NOM[e]: dict(Disponibilidad=g.loc[e, "D"], Rendimiento=g.loc[e, "R"], Calidad=g.loc[e, "Q"], OEE=g.loc[e, "OEE"],
                                        Parada_h=g.loc[e, "Pa"], Eventos=g.loc[e, "ev"], MTBF_h=g.loc[e, "MTBF"], MTTR_h=g.loc[e, "MTTR"],
                                        No_conformes=g.loc[e, "NC"], Indice_rechazo=g.loc[e, "rech"]) for e in g.index}).T.reset_index()
    t1 = t1.rename(columns={"index": "Estación"})
    info = pd.DataFrame([("Origen", d["origen"]), ("Archivo", d["archivo"]), ("Detalle", d["detalle"]), ("Periodo", periodo_txt(d)),
                         ("Días laborables", d["dias"]), ("Registros", len(df))], columns=["Dato", "Valor"])
    datos = df[["Fecha", "Estacion_ID", "Tiempo_Planificado_h", "Parada_No_Planificada_h", "Tiempo_Operativo_h", "Causa",
                "Capacidad_Nominal_bidones_h", "Bidones_Procesados", "Bidones_Conformes", "Bidones_No_Conformes"]].copy()
    datos["Fecha"] = datos["Fecha"].dt.strftime("%Y-%m-%d")
    hojas = [("KPIs y metas", df_kpis_metas(df)), ("Datos activos", info), ("Tabla 1", t1), ("Alertas", df_alertas(df)), ("Casos", df_casos()),
             ("Matriz de criticidad", _matriz_df()), ("AMEF", df_amef()), ("Parámetros", df_parametros()), ("Datos", datos)]
    if s.get("v_res") is not None:
        fr = frames_mc(s.v_res, df)
        hojas += [("MC Resumen", fr["resumen"]), ("MC Estaciones", fr["estaciones"]), ("MC Mensual", fr["mensual"]), ("MC Iteraciones", fr["iteraciones"])]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for nombre, tabla in hojas:
            (tabla if len(tabla) else pd.DataFrame({"Sin registros": []})).to_excel(xw, sheet_name=nombre, index=False)
        for ws in xw.book.worksheets:
            for col in ws.columns:
                ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col[:60])
                ws.column_dimensions[col[0].column_letter].width = min(max(12, ancho + 2), 70)
    return buf.getvalue(), [h[0] for h in hojas]


def pagina_export():
    s = st.session_state
    abrir_pagina("export", "Exportación",
                 "Descarga todo lo que hiciste en un solo Excel: KPIs, datos, tablas, alertas, casos, matriz, AMEF, parámetros y, si corriste la "
                 "validación, sus resultados. También hay CSV sueltos.")
    firma = (s.datos["archivo"], len(s.casos), tuple(c["estado"] for c in s.casos), str(s.matriz), str(s.amef_ov), str(s.d_ajustes),
             str(s.metas), id(s.get("v_res")), str(sorted(s.umb.items())))
    if s.get("xlsx_total_firma") != firma:
        s.xlsx_total = construir_excel_total()
        s.xlsx_total_firma = firma
    data, hojas = s.xlsx_total
    st.download_button("Descargar Excel completo", data, file_name="agua_bora_resultados.xlsx", key="dl_xlsx_total", type="primary",
                       width="stretch", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.markdown("**Hojas incluidas:** " + " · ".join(hojas))
    if s.get("v_res") is None:
        st.caption("Las hojas «MC …» aparecen cuando ejecutas la validación (paso 7).")
    st.markdown("**CSV sueltos**")
    df = df_activo()
    a, b = st.columns(2)
    c, d = st.columns(2)
    a.download_button("Casos (CSV)", df_casos().to_csv(index=False).encode("utf-8-sig"), file_name="casos.csv", mime="text/csv", key="dl_c_casos", width="stretch")
    b.download_button("Alertas (CSV)", df_alertas(df).to_csv(index=False).encode("utf-8-sig"), file_name="alertas.csv", mime="text/csv", key="dl_c_alertas", width="stretch")
    c.download_button("KPIs y metas (CSV)", df_kpis_metas(df).to_csv(index=False).encode("utf-8-sig"), file_name="kpis_y_metas.csv", mime="text/csv", key="dl_c_kpis", width="stretch")
    d.download_button("Matriz de criticidad (CSV)", _matriz_df().to_csv(index=False).encode("utf-8-sig"), file_name="matriz_criticidad.csv", mime="text/csv", key="dl_c_matriz", width="stretch")
    pie_siguiente("export")


# =============================================================================
# PRINCIPAL
# =============================================================================
PAGINAS = dict(inicio=pagina_inicio, kpis=pagina_kpis, datos=pagina_datos, params=pagina_params, tablero=pagina_tablero,
               decision=pagina_decision, ejecucion=pagina_ejecucion, valid=pagina_valid, export=pagina_export)


def inicializar_estado():
    s = st.session_state
    s.setdefault("pagina", "inicio")
    s.setdefault("visto", set())
    s.setdefault("casos", [])
    s.setdefault("caso_seq", 0)
    s.setdefault("wz", None)
    s.setdefault("metas", {})
    s.setdefault("matriz", {k: v["prioridad"] for k, v in MATRIZ_DEF.items()})
    s.setdefault("matriz_ver", 0)
    s.setdefault("rangos_ov", {})
    s.setdefault("amef_ov", {})
    s.setdefault("d_ajustes", {})
    s.setdefault("d_seed", 7)
    s.setdefault("umb", {})
    s.setdefault("umb_ver", 0)
    por_defecto = {
        "w_parada_min": 30, "w_rec_n": 2, "w_rec_vent": 1, "w_p_lect": LECT_DEF, "w_p_medio": MEDIO_DEF, "w_p_alto": ALTO_DEF,
        "w_gen_ini": pd.Timestamp("2026-07-01").date(), "w_gen_fin": pd.Timestamp("2026-09-30").date(), "w_gen_sem": 42,
        "w_cfg_iter": 60, "w_cfg_seed": 42, "w_cfg_scen": list(ESCENARIOS_SIM), "w_v_sub": "1. Configuración",
        "w_d_esc": "(d) Cuello de botella oculto (OEE aceptable)", "w_tend_freq": "Semanal", "w_dec_tipo": "Un día",
        "w_ej_fecha": pd.Timestamp.today().date(), "w_ej_resp": "", "w_ej_accion": "", "w_ej_efecto": EFECTOS[0],
    }
    por_defecto.update({"w_" + k: v for k, v in SUPUESTOS_DEF.items()})
    for k, v in por_defecto.items():
        s.setdefault(k, v)
    # Streamlit borra el estado de los widgets que no se dibujan en una ejecución; reasignar la clave lo conserva
    for k in list(s.keys()):
        if isinstance(k, str) and k.startswith("w_") and not k.startswith(("w_up", "w_ed_")):
            s[k] = s[k]


def inyectar_html(html):
    """Ejecuta un script en la página (st.iframe en Streamlit nuevo; components.html en versiones anteriores)."""
    if hasattr(st, "iframe"):
        st.iframe(html, height=1)
    else:
        import streamlit.components.v1 as components
        components.html(html, height=0)


def main():
    st.set_page_config(page_title="Agua Bora · DSS", page_icon="💧", layout="wide")
    st.markdown(CSS, unsafe_allow_html=True)
    # Evita la traducción automática de Chrome (rompe el DOM de React y los textos)
    inyectar_html("""<script>
    (function () {
      var d = window.parent.document, h = d.documentElement;
      h.setAttribute("lang", "es"); h.setAttribute("translate", "no"); h.classList.add("notranslate");
      if (!d.querySelector('meta[name="google"]')) {
        var m = d.createElement("meta"); m.name = "google"; m.content = "notranslate"; d.head.appendChild(m);
      }
    })();
    </script>""")
    inicializar_estado()
    aplicar_config()
    s = st.session_state
    msg = bloqueo(s.pagina)
    if msg:
        abrir_pagina("inicio", NOMBRE[s.pagina], "Este paso todavía no está disponible.")
        caja("🔒 " + msg, "warn")
        s.button_destino = "datos" if not hay_datos() else "kpis"
        st.button("Ir al paso que falta", key="b_ir_falta", type="primary", on_click=ir, args=(s.button_destino,))
    else:
        PAGINAS[s.pagina]()
    barra_lateral()
    st.markdown(f'<div class="foot">Agua Bora · Prototipo TPM-IIoT · datos simulados con fines académicos · {VERSION}</div>', unsafe_allow_html=True)


main()
